"""Format-preserving openpyxl read/write for Wallet Editor Dropbox registry."""

from __future__ import annotations

from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from integrations.wallet_editor_registry_lifecycle import (
    ALL_RESULTS_COLUMNS,
    HOLD_COLUMNS,
    LEGACY_VALUE_COLUMN,
    MISSING_OTLEZKA_DATE_TEXT,
    OPERATION_DATE_COLUMN,
    OPERATION_DATE_NUMBER_FORMAT,
    OTLEZKA_COLUMNS,
    RED_FILL,
    RUNS_COLUMNS,
    SHEET_ALL_RESULTS,
    SHEET_HOLD,
    SHEET_OTLEZKA,
    SHEET_RUNS,
    _cell_str,
    migrate_legacy_all_results,
    migrate_legacy_runs,
    normalize_all_results,
    normalize_sheet,
    is_legacy_all_results,
    parse_operation_date,
)

CARD_TEXT_FORMAT = "@"
TEXT_COLUMNS = frozenset({"card"})


def copy_cell_style(src_cell: Cell, dst_cell: Cell) -> None:
    """Copy visual cell style without changing value."""
    dst_cell.font = copy(src_cell.font)
    dst_cell.border = copy(src_cell.border)
    dst_cell.fill = copy(src_cell.fill)
    dst_cell.number_format = src_cell.number_format
    dst_cell.protection = copy(src_cell.protection)
    dst_cell.alignment = copy(src_cell.alignment)


def copy_row_style(
    ws: Worksheet,
    src_row: int,
    dst_row: int,
    column_indices: list[int],
) -> None:
    for col_idx in column_indices:
        copy_cell_style(ws.cell(row=src_row, column=col_idx), ws.cell(row=dst_row, column=col_idx))


def _column_indices(header_map: dict[str, int], columns: list[str]) -> list[int]:
    return [header_map[col] for col in columns if col in header_map]


def _count_data_rows(ws: Worksheet, headers: dict[str, int]) -> int:
    count = 0
    for row_idx in range(2, ws.max_row + 1):
        if _row_has_data(ws, row_idx, headers):
            count += 1
    return count


def _style_template_row(ws: Worksheet, headers: dict[str, int], previous_data_rows: int) -> int | None:
    if previous_data_rows > 0:
        last_row: int | None = None
        for row_idx in range(2, ws.max_row + 1):
            if _row_has_data(ws, row_idx, headers):
                last_row = row_idx
        return last_row
    if ws.max_row >= 2:
        return 2
    return None


def _copy_header_style_from_neighbor(ws: Worksheet, col_idx: int) -> None:
    if col_idx <= 1:
        return
    copy_cell_style(ws.cell(row=1, column=col_idx - 1), ws.cell(row=1, column=col_idx))


def _copy_header_style_from_right_neighbor(ws: Worksheet, col_idx: int) -> None:
    if col_idx >= ws.max_column:
        return
    copy_cell_style(ws.cell(row=1, column=col_idx + 1), ws.cell(row=1, column=col_idx))


def _ensure_operation_date_column(ws: Worksheet) -> None:
    if "Дата операции" in _header_map(ws):
        return
    ws.insert_cols(1)
    ws.cell(row=1, column=1, value="Дата операции")
    _copy_header_style_from_right_neighbor(ws, 1)


def card_as_text(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _cell_str(value)


def _header_map(ws: Worksheet) -> dict[str, int]:
    result: dict[str, int] = {}
    for col_idx, cell in enumerate(ws[1], 1):
        if cell.value is not None:
            key = str(cell.value).strip()
            if key:
                result[key] = col_idx
    return result


def _row_has_data(ws: Worksheet, row_idx: int, headers: dict[str, int]) -> bool:
    for col_idx in headers.values():
        if _cell_str(ws.cell(row=row_idx, column=col_idx).value):
            return True
    return False


def _read_cell(ws: Worksheet, row_idx: int, col_name: str, headers: dict[str, int]) -> str:
    col_idx = headers.get(col_name)
    if not col_idx:
        return ""
    raw = ws.cell(row=row_idx, column=col_idx).value
    if col_name in TEXT_COLUMNS:
        return card_as_text(raw)
    return _cell_str(raw)


def _set_cell_value(
    ws: Worksheet,
    row_idx: int,
    col_idx: int,
    col_name: str,
    value: object,
) -> None:
    cell = ws.cell(row=row_idx, column=col_idx)
    if col_name in TEXT_COLUMNS:
        cell.value = card_as_text(value)
        cell.number_format = CARD_TEXT_FORMAT
    elif col_name == OPERATION_DATE_COLUMN:
        excel_date = parse_operation_date(value)
        if excel_date is None:
            cell.value = None
        else:
            cell.value = excel_date
            cell.number_format = OPERATION_DATE_NUMBER_FORMAT
    else:
        cell.value = value if value != "" else None


def _remove_column_if_present(ws: Worksheet, col_name: str) -> None:
    headers = _header_map(ws)
    idx = headers.get(col_name)
    if idx is not None:
        ws.delete_cols(idx)


def read_all_results_ws(ws: Worksheet | None) -> pd.DataFrame:
    if ws is None or ws.max_row < 1:
        return normalize_all_results(pd.DataFrame())
    headers = _header_map(ws)
    if not headers:
        return normalize_all_results(pd.DataFrame())

    rows: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        if not _row_has_data(ws, row_idx, headers):
            continue
        record = {col: _read_cell(ws, row_idx, col, headers) for col in ALL_RESULTS_COLUMNS}
        if LEGACY_VALUE_COLUMN in headers:
            record[LEGACY_VALUE_COLUMN] = _read_cell(ws, row_idx, LEGACY_VALUE_COLUMN, headers)
        rows.append(record)

    if not rows:
        return normalize_all_results(pd.DataFrame())
    df = pd.DataFrame(rows)
    if is_legacy_all_results(df):
        df = migrate_legacy_all_results(df)
    return normalize_all_results(df)


def read_runs_ws(ws: Worksheet | None) -> pd.DataFrame:
    if ws is None or ws.max_row < 2:
        return migrate_legacy_runs(pd.DataFrame())
    headers = _header_map(ws)
    rows: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        if not _row_has_data(ws, row_idx, headers):
            continue
        record = {col: _read_cell(ws, row_idx, col, headers) for col in RUNS_COLUMNS}
        rows.append(record)
    if not rows:
        return migrate_legacy_runs(pd.DataFrame())
    return migrate_legacy_runs(pd.DataFrame(rows))


def read_user_sheet_ws(ws: Worksheet | None, columns: list[str]) -> pd.DataFrame:
    if ws is None or ws.max_row < 2:
        return normalize_sheet(pd.DataFrame(), columns)
    headers = _header_map(ws)
    rows: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        if not _row_has_data(ws, row_idx, headers):
            continue
        record: dict[str, str] = {}
        for col in columns:
            record[col] = _read_cell(ws, row_idx, col, headers)
        rows.append(record)
    if not rows:
        return normalize_sheet(pd.DataFrame(), columns)
    return normalize_sheet(pd.DataFrame(rows), columns)


def load_registry_frames(local_path: Path, download_status: str) -> tuple[
    pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, bool, bool
]:
    if download_status == "not_found":
        return (
            normalize_all_results(pd.DataFrame()),
            migrate_legacy_runs(pd.DataFrame()),
            normalize_sheet(pd.DataFrame(), HOLD_COLUMNS),
            normalize_sheet(pd.DataFrame(), OTLEZKA_COLUMNS),
            False,
            False,
        )

    wb = load_workbook(local_path, data_only=False)
    hold_exists = SHEET_HOLD in wb.sheetnames
    otlezka_exists = SHEET_OTLEZKA in wb.sheetnames

    all_df = read_all_results_ws(wb[SHEET_ALL_RESULTS] if SHEET_ALL_RESULTS in wb.sheetnames else None)
    runs_df = read_runs_ws(wb[SHEET_RUNS] if SHEET_RUNS in wb.sheetnames else None)
    hold_df = (
        read_user_sheet_ws(wb[SHEET_HOLD], HOLD_COLUMNS)
        if hold_exists
        else normalize_sheet(pd.DataFrame(), HOLD_COLUMNS)
    )
    otlezka_df = (
        read_user_sheet_ws(wb[SHEET_OTLEZKA], OTLEZKA_COLUMNS)
        if otlezka_exists
        else normalize_sheet(pd.DataFrame(), OTLEZKA_COLUMNS)
    )
    wb.close()
    return all_df, runs_df, hold_df, otlezka_df, hold_exists, otlezka_exists


def _ensure_sheet(wb: Workbook, name: str) -> Worksheet:
    if name not in wb.sheetnames:
        wb.create_sheet(name)
    return wb[name]


def _write_headers(ws: Worksheet, columns: list[str]) -> dict[str, int]:
    for col_idx, name in enumerate(columns, 1):
        ws.cell(row=1, column=col_idx, value=name)
    return {name: idx for idx, name in enumerate(columns, 1)}


def _sync_all_results_sheet(ws: Worksheet, df: pd.DataFrame) -> None:
    df = normalize_all_results(df)
    _remove_column_if_present(ws, LEGACY_VALUE_COLUMN)

    existing_headers = _header_map(ws)
    if not existing_headers:
        header_map = _write_headers(ws, ALL_RESULTS_COLUMNS)
    else:
        _ensure_operation_date_column(ws)
        header_map = _header_map(ws)
        for col_name in ALL_RESULTS_COLUMNS:
            if col_name not in header_map:
                next_col = ws.max_column + 1
                ws.cell(row=1, column=next_col, value=col_name)
                _copy_header_style_from_neighbor(ws, next_col)
                header_map[col_name] = next_col
        header_map = {k: header_map[k] for k in ALL_RESULTS_COLUMNS if k in header_map}

    data_rows = len(df)
    previous_data_rows = _count_data_rows(ws, header_map) if header_map else 0
    template_row = _style_template_row(ws, header_map, previous_data_rows)
    style_columns = _column_indices(header_map, ALL_RESULTS_COLUMNS)

    for row_offset in range(data_rows):
        row_idx = row_offset + 2
        if (
            row_offset >= previous_data_rows
            and template_row is not None
            and row_idx != template_row
        ):
            copy_row_style(ws, template_row, row_idx, style_columns)
        for col_name in ALL_RESULTS_COLUMNS:
            col_idx = header_map.get(col_name)
            if not col_idx:
                continue
            value = df.iloc[row_offset].get(col_name, "")
            _set_cell_value(ws, row_idx, col_idx, col_name, value)

    last_data_row = data_rows + 1
    if ws.max_row > last_data_row:
        for row_idx in range(last_data_row + 1, ws.max_row + 1):
            for col_name in ALL_RESULTS_COLUMNS:
                col_idx = header_map.get(col_name)
                if col_idx:
                    ws.cell(row=row_idx, column=col_idx).value = None

    _apply_missing_otlezka_red_fill_ws(ws, header_map)


def _sync_runs_sheet(ws: Worksheet, runs_df: pd.DataFrame) -> None:
    runs_df = migrate_legacy_runs(runs_df)
    header_map = _header_map(ws)
    if not header_map:
        header_map = _write_headers(ws, RUNS_COLUMNS)
    else:
        for col_name in RUNS_COLUMNS:
            if col_name not in header_map:
                next_col = ws.max_column + 1
                ws.cell(row=1, column=next_col, value=col_name)
                _copy_header_style_from_neighbor(ws, next_col)
                header_map[col_name] = next_col

    previous_data_rows = _count_data_rows(ws, header_map) if header_map else 0
    template_row = _style_template_row(ws, header_map, previous_data_rows)
    style_columns = _column_indices(header_map, RUNS_COLUMNS)

    for row_offset in range(len(runs_df)):
        row_idx = row_offset + 2
        if (
            row_offset >= previous_data_rows
            and template_row is not None
            and row_idx != template_row
        ):
            copy_row_style(ws, template_row, row_idx, style_columns)
        for col_name in RUNS_COLUMNS:
            col_idx = header_map.get(col_name)
            if not col_idx:
                continue
            _set_cell_value(ws, row_idx, col_idx, col_name, runs_df.iloc[row_offset].get(col_name, ""))


def _create_user_sheet_headers(ws: Worksheet, columns: list[str]) -> None:
    _write_headers(ws, columns)


def _apply_missing_otlezka_red_fill_ws(ws: Worksheet, header_map: dict[str, int]) -> None:
    col_idx = header_map.get("Дата включения")
    if not col_idx:
        return
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        if _cell_str(cell.value) == MISSING_OTLEZKA_DATE_TEXT:
            cell.fill = RED_FILL


def save_registry_workbook(
    local_path: Path,
    *,
    all_results: pd.DataFrame,
    runs: pd.DataFrame,
    hold_exists: bool,
    otlezka_exists: bool,
    is_new_file: bool,
) -> None:
    local_path.parent.mkdir(parents=True, exist_ok=True)

    if is_new_file:
        wb = Workbook()
        default = wb.active
        if default is not None:
            wb.remove(default)
        ws_all = wb.create_sheet(SHEET_ALL_RESULTS)
        _write_headers(ws_all, ALL_RESULTS_COLUMNS)
        ws_runs = wb.create_sheet(SHEET_RUNS)
        _write_headers(ws_runs, RUNS_COLUMNS)
        if not hold_exists:
            _create_user_sheet_headers(wb.create_sheet(SHEET_HOLD), HOLD_COLUMNS)
        if not otlezka_exists:
            _create_user_sheet_headers(wb.create_sheet(SHEET_OTLEZKA), OTLEZKA_COLUMNS)
    else:
        wb = load_workbook(local_path)

    ws_all = _ensure_sheet(wb, SHEET_ALL_RESULTS)
    _sync_all_results_sheet(ws_all, all_results)

    ws_runs = _ensure_sheet(wb, SHEET_RUNS)
    _sync_runs_sheet(ws_runs, runs)

    if not hold_exists and SHEET_HOLD not in wb.sheetnames:
        _create_user_sheet_headers(wb.create_sheet(SHEET_HOLD), HOLD_COLUMNS)
    if not otlezka_exists and SHEET_OTLEZKA not in wb.sheetnames:
        _create_user_sheet_headers(wb.create_sheet(SHEET_OTLEZKA), OTLEZKA_COLUMNS)

    wb.save(local_path)
    wb.close()


def create_styled_registry_workbook(path: Path) -> None:
    """Test helper: workbook with column width, freeze panes, header fill."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = SHEET_ALL_RESULTS
    for col_idx, name in enumerate(ALL_RESULTS_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")
    ws.column_dimensions["A"].width = 22.5
    ws.column_dimensions["G"].width = 18.0
    ws.freeze_panes = "A2"
    wb.create_sheet(SHEET_RUNS)
    _write_headers(wb[SHEET_RUNS], RUNS_COLUMNS)
    wb.save(path)
    wb.close()
