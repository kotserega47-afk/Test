"""Export-only workbook formatting for WalletEditor registry (read-only report)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Iterable
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.datetime_utils import EXCEL_DATETIME_FORMAT, EXCEL_DATE_FORMAT, parse_msk_datetime
from integrations.wallet_editor_registry_lifecycle import (
    ACTION_ADD_PARTNER,
    ACTION_REMOVE_PARTNER,
    ALL_RESULTS_COLUMNS,
    HOLD_COLUMNS,
    OPERATION_DATE_COLUMN,
    OTLEZKA_COLUMNS,
    RUNS_COLUMNS,
    parse_operation_date,
)

EXPORT_SHEET_RESULTS = "Результаты"
EXPORT_SHEET_RUNS = "Запуски"
EXPORT_SHEET_HOLD = "Hold"
EXPORT_SHEET_OTLEZKA = "Отлёжка"
EXPORT_SHEET_STATS = "Статистика"
EXPORT_SHEET_README = "README"

EXPORT_DATETIME_DISPLAY = "%d.%m.%Y %H:%M:%S"
EXPORT_DATE_DISPLAY = "%d.%m.%Y"
HOLD_ADDED_DATE_COLUMN = "Дата добавления"
STATS_PARTNER_HEADER = "Партнёр"
STATS_TOTAL_HEADER = "Всего"
STATS_REMOVE_TITLE = "Remove Partner"
STATS_ADD_TITLE = "Add Partner"
STATS_TOTAL_ROW_LABEL = "Итого"
STATS_FIRST_DATE_COL = 3

MAX_COLUMN_WIDTH = 50
MSK = ZoneInfo("Europe/Moscow")

ALL_RESULTS_DATETIME_COLUMNS = (
    OPERATION_DATE_COLUMN,
    "Дата отключения",
    "Дата включения",
)
RUNS_DATETIME_COLUMNS = ("started_at", "finished_at")
HOLD_DATETIME_COLUMNS = (HOLD_ADDED_DATE_COLUMN,)

_THIN = Side(style="thin")
_ALL_BORDERS = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BOLD = Font(bold=True)


def sort_all_results_for_export(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    work = df.copy()
    work["_sort_op"] = work[OPERATION_DATE_COLUMN].map(
        lambda value: parse_operation_date(value) or date.min
    )
    work["_sort_idx"] = range(len(work))
    work = work.sort_values(["_sort_op", "_sort_idx"], ascending=[True, True], kind="mergesort")
    return work.drop(columns=["_sort_op", "_sort_idx"]).reset_index(drop=True)


def sort_runs_for_export(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    work = df.copy()
    work["_sort_started"] = work["started_at"].map(_parse_sort_datetime)
    work["_sort_idx"] = range(len(work))
    work = work.sort_values(["_sort_started", "_sort_idx"], ascending=[True, True], kind="mergesort")
    return work.drop(columns=["_sort_started", "_sort_idx"]).reset_index(drop=True)


def sort_hold_for_export(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty or HOLD_ADDED_DATE_COLUMN not in df.columns:
        return df.reset_index(drop=True)
    work = df.copy()
    work["_sort_added"] = work[HOLD_ADDED_DATE_COLUMN].map(
        lambda value: parse_operation_date(value) or date.min
    )
    work["_sort_idx"] = range(len(work))
    work = work.sort_values(["_sort_added", "_sort_idx"], ascending=[True, True], kind="mergesort")
    return work.drop(columns=["_sort_added", "_sort_idx"]).reset_index(drop=True)


def format_export_datetime_value(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, datetime):
        return value.strftime(EXPORT_DATETIME_DISPLAY)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).strftime(EXPORT_DATETIME_DISPLAY)

    text = str(value).strip()
    if not text or text.lower() in {"nan", "nat", "none"}:
        return ""

    parsed_dt = _try_parse_datetime(text)
    if parsed_dt is not None:
        return parsed_dt.strftime(EXPORT_DATETIME_DISPLAY)

    parsed_date = parse_operation_date(text)
    if parsed_date is not None:
        if " " in text or ":" in text:
            for fmt in (EXCEL_DATETIME_FORMAT, "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
                try:
                    return datetime.strptime(text.replace("Z", ""), fmt).strftime(
                        EXPORT_DATETIME_DISPLAY
                    )
                except ValueError:
                    continue
        return datetime.combine(parsed_date, datetime.min.time()).strftime(EXPORT_DATETIME_DISPLAY)

    return text


def format_dataframe_datetimes(df: pd.DataFrame, columns: Iterable[str]) -> pd.DataFrame:
    if df.empty:
        return df
    work = df.copy()
    for col in columns:
        if col in work.columns:
            work[col] = work[col].map(format_export_datetime_value)
    return work


def prepare_export_frames(
    all_results: pd.DataFrame,
    runs: pd.DataFrame,
    hold: pd.DataFrame,
    otlezka: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    all_results = format_dataframe_datetimes(
        sort_all_results_for_export(all_results),
        ALL_RESULTS_DATETIME_COLUMNS,
    )
    runs = format_dataframe_datetimes(sort_runs_for_export(runs), RUNS_DATETIME_COLUMNS)
    hold = format_dataframe_datetimes(sort_hold_for_export(hold), HOLD_DATETIME_COLUMNS)
    return all_results, runs, hold, otlezka.reset_index(drop=True)


def build_statistics_sheet_rows(all_results: pd.DataFrame) -> list[list[object]]:
    rows: list[list[object]] = []
    rows.extend(_build_statistics_block(all_results, ACTION_REMOVE_PARTNER, STATS_REMOVE_TITLE))
    rows.append([])
    rows.extend(_build_statistics_block(all_results, ACTION_ADD_PARTNER, STATS_ADD_TITLE))
    return rows


def write_registry_export_workbook(
    local_path: Path,
    *,
    all_results: pd.DataFrame,
    runs: pd.DataFrame,
    hold: pd.DataFrame,
    otlezka: pd.DataFrame,
    readme_lines: list[str],
) -> None:
    all_results, runs, hold, otlezka = prepare_export_frames(all_results, runs, hold, otlezka)

    local_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)

    _write_data_sheet(wb.create_sheet(EXPORT_SHEET_RESULTS), ALL_RESULTS_COLUMNS, all_results)
    _write_data_sheet(wb.create_sheet(EXPORT_SHEET_RUNS), RUNS_COLUMNS, runs)
    _write_data_sheet(wb.create_sheet(EXPORT_SHEET_HOLD), HOLD_COLUMNS, hold)
    _write_data_sheet(wb.create_sheet(EXPORT_SHEET_OTLEZKA), OTLEZKA_COLUMNS, otlezka)
    _write_statistics_sheet(wb.create_sheet(EXPORT_SHEET_STATS), all_results)
    _write_readme_sheet(wb.create_sheet(EXPORT_SHEET_README), readme_lines)

    wb.save(local_path)
    wb.close()


def _parse_sort_datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    parsed = _try_parse_datetime(str(value).strip()) if value is not None else None
    if parsed is not None:
        return parsed.replace(tzinfo=None)
    parsed_date = parse_operation_date(value)
    if parsed_date is not None:
        return datetime.combine(parsed_date, datetime.min.time())
    return datetime.min


def _try_parse_datetime(text: str) -> datetime | None:
    if not text:
        return None
    try:
        return parse_msk_datetime(text).replace(tzinfo=None)
    except ValueError:
        pass
    iso_candidate = text.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(iso_candidate)
        if parsed.tzinfo is not None:
            parsed = parsed.astimezone(MSK)
        return parsed.replace(tzinfo=None)
    except ValueError:
        return None


def _operation_date_key(value: object) -> date | None:
    return parse_operation_date(value)


def _build_statistics_block(
    all_results: pd.DataFrame,
    action: str,
    title: str,
) -> list[list[object]]:
    if all_results.empty or "action" not in all_results.columns:
        return [[title], [STATS_PARTNER_HEADER], [STATS_TOTAL_ROW_LABEL]]

    subset = all_results[
        all_results["action"].astype(str).str.strip().str.lower() == action.lower()
    ]
    if subset.empty:
        return [[title], [STATS_PARTNER_HEADER, "", STATS_TOTAL_HEADER], [STATS_TOTAL_ROW_LABEL, "", 0]]

    dated = subset.copy()
    dated["_op_date"] = dated[OPERATION_DATE_COLUMN].map(_operation_date_key)
    dated = dated[dated["_op_date"].notna()]
    if dated.empty:
        return [[title], [STATS_PARTNER_HEADER, "", STATS_TOTAL_HEADER], [STATS_TOTAL_ROW_LABEL, "", 0]]

    dates = sorted(dated["_op_date"].unique(), reverse=True)
    partners = sorted(
        {str(p).strip() for p in dated["partner"].fillna("").astype(str) if str(p).strip()},
        key=lambda value: value.casefold(),
    )

    header = [STATS_PARTNER_HEADER, ""]
    header.extend(d.strftime(EXPORT_DATE_DISPLAY) for d in dates)
    header.append(STATS_TOTAL_HEADER)

    matrix: dict[str, dict[date, int]] = {partner: dict.fromkeys(dates, 0) for partner in partners}
    for _, row in dated.iterrows():
        partner = str(row.get("partner", "")).strip()
        op_date = row["_op_date"]
        if partner in matrix and op_date in matrix[partner]:
            matrix[partner][op_date] += 1

    rows: list[list[object]] = [[title], header]
    column_totals = dict.fromkeys(dates, 0)
    grand_total = 0

    for partner in partners:
        counts = matrix[partner]
        row_total = sum(counts.values())
        grand_total += row_total
        line = [partner, ""]
        for op_date in dates:
            count = counts[op_date]
            column_totals[op_date] += count
            line.append(count if count else "")
        line.append(row_total if row_total else "")
        rows.append(line)

    total_row = [STATS_TOTAL_ROW_LABEL, ""]
    for op_date in dates:
        total = column_totals[op_date]
        total_row.append(total if total else "")
    total_row.append(grand_total if grand_total else "")
    rows.append(total_row)
    return rows


def _write_data_sheet(ws: Worksheet, columns: list[str], df: pd.DataFrame) -> None:
    for col_idx, name in enumerate(columns, 1):
        ws.cell(row=1, column=col_idx, value=name)

    max_row = len(df) + 1
    max_col = len(columns)
    for row_offset in range(len(df)):
        row_idx = row_offset + 2
        for col_idx, col_name in enumerate(columns, 1):
            value = df.iloc[row_offset].get(col_name, "")
            if value is None or (isinstance(value, float) and pd.isna(value)):
                value = ""
            ws.cell(row=row_idx, column=col_idx, value=value)

    _apply_sheet_formatting(
        ws,
        max_row=max_row,
        max_col=max_col,
        autofilter=ws.title
        in {EXPORT_SHEET_RESULTS, EXPORT_SHEET_RUNS, EXPORT_SHEET_HOLD, EXPORT_SHEET_OTLEZKA},
    )


def _write_statistics_sheet(ws: Worksheet, all_results: pd.DataFrame) -> None:
    rows = build_statistics_sheet_rows(all_results)
    max_row = 0
    max_col = 0
    for row_idx, row in enumerate(rows, 1):
        if not row:
            continue
        max_row = row_idx
        for col_idx, value in enumerate(row, 1):
            max_col = max(max_col, col_idx)
            ws.cell(row=row_idx, column=col_idx, value=value)

    if max_row == 0:
        return

    for row_idx in range(1, max_row + 1):
        first = ws.cell(row=row_idx, column=1).value
        if first in {STATS_REMOVE_TITLE, STATS_ADD_TITLE}:
            for col_idx in range(1, max_col + 1):
                ws.cell(row=row_idx, column=col_idx).font = _BOLD

    header_rows = [
        row_idx
        for row_idx in range(1, max_row + 1)
        if ws.cell(row=row_idx, column=1).value == STATS_PARTNER_HEADER
    ]
    for row_idx in header_rows:
        for col_idx in range(1, max_col + 1):
            ws.cell(row=row_idx, column=col_idx).font = _BOLD

    for row_idx in range(1, max_row + 1):
        if ws.cell(row=row_idx, column=1).value == STATS_TOTAL_ROW_LABEL:
            for col_idx in range(1, max_col + 1):
                ws.cell(row=row_idx, column=col_idx).font = _BOLD

    _apply_sheet_formatting(ws, max_row=max_row, max_col=max_col, autofilter=False)


def _write_readme_sheet(ws: Worksheet, lines: list[str]) -> None:
    max_row = len(lines)
    for row_idx, line in enumerate(lines, 1):
        ws.cell(row=row_idx, column=1, value=line)
    if max_row:
        _apply_sheet_formatting(ws, max_row=max_row, max_col=1, autofilter=False)


def _apply_sheet_formatting(
    ws: Worksheet,
    *,
    max_row: int,
    max_col: int,
    autofilter: bool,
) -> None:
    if max_row < 1 or max_col < 1:
        return

    for col_idx in range(1, max_col + 1):
        ws.cell(row=1, column=col_idx).font = _BOLD

    for row_idx in range(1, max_row + 1):
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value not in (None, ""):
                cell.border = _ALL_BORDERS
                if cell.alignment is None or cell.alignment.horizontal is None:
                    cell.alignment = Alignment(vertical="center")

    _apply_column_widths(ws, max_col=max_col)
    ws.freeze_panes = "A2" if max_row >= 2 else None
    if autofilter and max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"


def _apply_column_widths(ws: Worksheet, *, max_col: int) -> None:
    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 8), MAX_COLUMN_WIDTH)
