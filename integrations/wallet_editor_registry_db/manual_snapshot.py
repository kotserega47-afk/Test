"""Normalized snapshot of operator sheets hold / Отлёжка for manual sync."""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from integrations.wallet_editor_registry_lifecycle import (
    HOLD_COLUMNS,
    OTLEZKA_COLUMNS,
    SHEET_HOLD,
    SHEET_OTLEZKA,
    _cell_str,
    _normalize_key,
)
from integrations.wallet_editor_registry_xlsx import read_user_sheet_ws

SNAPSHOT_VERSION = 1

HOLD_REQUIRED_COLUMNS = frozenset({"card", "partner"})
OTLEZKA_REQUIRED_COLUMNS = frozenset({"partner", "Полные дни"})


class ManualSyncValidationError(ValueError):
    """Structural validation failure — sync must not mutate PG."""


@dataclass(frozen=True, slots=True)
class HoldSnapshotRow:
    card: str
    partner: str
    card_norm: str
    partner_norm: str
    added_at: str | None
    comment: str | None
    source_row_index: int


@dataclass(frozen=True, slots=True)
class OtlezkaSnapshotRow:
    partner: str
    partner_norm: str
    full_days: int
    comment: str | None
    source_row_index: int


@dataclass(frozen=True, slots=True)
class ManualSnapshot:
    hold_rows: tuple[HoldSnapshotRow, ...]
    otlezka_rows: tuple[OtlezkaSnapshotRow, ...]
    warnings: tuple[str, ...] = ()
    hold_sheet_exists: bool = True
    otlezka_sheet_exists: bool = True


def _missing_columns(headers: set[str], required: frozenset[str]) -> list[str]:
    return sorted(required - headers)


def _parse_full_days(raw: object) -> int:
    if raw is None or (isinstance(raw, float) and str(raw) == "nan"):
        raise ManualSyncValidationError("Полные дни must be a non-negative integer")
    try:
        days = int(float(raw))
    except (TypeError, ValueError) as exc:
        raise ManualSyncValidationError("Полные дни must be a non-negative integer") from exc
    if days < 0:
        raise ManualSyncValidationError("Полные дни must be >= 0")
    return days


def parse_manual_workbook(local_path: Path, download_status: str) -> ManualSnapshot:
    """Parse hold/Отлёжка only; ignore all other sheets."""
    if download_status == "not_found":
        raise ManualSyncValidationError("manual workbook not found in Dropbox")
    if download_status == "error":
        raise ManualSyncValidationError("manual workbook download failed")

    warnings: list[str] = []
    wb = load_workbook(local_path, data_only=True)
    try:
        hold_exists = SHEET_HOLD in wb.sheetnames
        otlezka_exists = SHEET_OTLEZKA in wb.sheetnames

        if not hold_exists:
            raise ManualSyncValidationError(f"missing required sheet: {SHEET_HOLD}")

        if not otlezka_exists:
            warnings.append(f"missing sheet {SHEET_OTLEZKA}: treated as empty (warning only)")

        hold_ws = wb[SHEET_HOLD]
        hold_headers = {
            str(c.value).strip()
            for c in hold_ws[1]
            if c.value is not None and str(c.value).strip()
        }
        missing_hold = _missing_columns(hold_headers, HOLD_REQUIRED_COLUMNS)
        if missing_hold:
            raise ManualSyncValidationError(
                f"hold sheet missing required columns: {', '.join(missing_hold)}"
            )

        hold_df = read_user_sheet_ws(hold_ws, list(HOLD_COLUMNS))
        hold_rows = _build_hold_rows(hold_df)

        otlezka_rows: tuple[OtlezkaSnapshotRow, ...] = ()
        if otlezka_exists:
            otlezka_ws = wb[SHEET_OTLEZKA]
            otlezka_headers = {
                str(c.value).strip()
                for c in otlezka_ws[1]
                if c.value is not None and str(c.value).strip()
            }
            missing_ot = _missing_columns(otlezka_headers, OTLEZKA_REQUIRED_COLUMNS)
            if missing_ot:
                raise ManualSyncValidationError(
                    f"Отлёжка sheet missing required columns: {', '.join(missing_ot)}"
                )
            otlezka_df = read_user_sheet_ws(otlezka_ws, list(OTLEZKA_COLUMNS))
            otlezka_rows = _build_otlezka_rows(otlezka_df)

        return ManualSnapshot(
            hold_rows=hold_rows,
            otlezka_rows=otlezka_rows,
            warnings=tuple(warnings),
            hold_sheet_exists=hold_exists,
            otlezka_sheet_exists=otlezka_exists,
        )
    finally:
        wb.close()


def _build_hold_rows(hold_df) -> tuple[HoldSnapshotRow, ...]:
    rows: list[HoldSnapshotRow] = []
    seen: set[tuple[str, str]] = set()
    for index, row in hold_df.iterrows():
        card = _cell_str(row.get("card", ""))
        partner = _cell_str(row.get("partner", ""))
        if not card or not partner:
            continue
        card_norm = _normalize_key(card)
        partner_norm = _normalize_key(partner)
        key = (card_norm, partner_norm)
        if key in seen:
            raise ManualSyncValidationError(
                f"duplicate hold key card={card!r} partner={partner!r}"
            )
        seen.add(key)
        added_at = _cell_str(row.get("Дата добавления", "")) or None
        comment = _cell_str(row.get("comment", "")) or None
        rows.append(
            HoldSnapshotRow(
                card=card,
                partner=partner,
                card_norm=card_norm,
                partner_norm=partner_norm,
                added_at=added_at,
                comment=comment,
                source_row_index=int(index) + 2,
            )
        )
    return tuple(rows)


def _build_otlezka_rows(otlezka_df) -> tuple[OtlezkaSnapshotRow, ...]:
    rows: list[OtlezkaSnapshotRow] = []
    seen: set[str] = set()
    for index, row in otlezka_df.iterrows():
        partner = _cell_str(row.get("partner", ""))
        if not partner:
            continue
        partner_norm = _normalize_key(partner)
        if partner_norm in seen:
            raise ManualSyncValidationError(f"duplicate Отлёжка partner={partner!r}")
        seen.add(partner_norm)
        full_days = _parse_full_days(row.get("Полные дни", ""))
        comment = _cell_str(row.get("comment", "")) or None
        rows.append(
            OtlezkaSnapshotRow(
                partner=partner,
                partner_norm=partner_norm,
                full_days=full_days,
                comment=comment,
                source_row_index=int(index) + 2,
            )
        )
    return tuple(rows)


def build_canonical_payload(snapshot: ManualSnapshot) -> dict:
    hold_entries = [
        {
            "c": row.card_norm,
            "p": row.partner_norm,
            "a": row.added_at,
            "m": row.comment,
        }
        for row in sorted(snapshot.hold_rows, key=lambda r: (r.card_norm, r.partner_norm))
    ]
    otlezka_entries = [
        {
            "p": row.partner_norm,
            "d": row.full_days,
            "m": row.comment,
        }
        for row in sorted(snapshot.otlezka_rows, key=lambda r: r.partner_norm)
    ]
    return {
        "v": SNAPSHOT_VERSION,
        "hold": hold_entries,
        "otlezka": otlezka_entries,
    }


def compute_snapshot_hash(snapshot: ManualSnapshot) -> str:
    payload = build_canonical_payload(snapshot)
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def snapshot_hash_short(snapshot_hash: str, *, length: int = 12) -> str:
    return snapshot_hash[:length] if snapshot_hash else ""
