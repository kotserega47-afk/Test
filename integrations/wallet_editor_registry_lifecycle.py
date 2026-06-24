"""Lifecycle fields, hold/Отлёжка rules for Wallet Editor Dropbox registry."""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from core.datetime_utils import EXCEL_DATETIME_FORMAT, EXCEL_DATE_FORMAT, now_msk
from utils.loggers import get_logger
from utils.log_profiles import LOG_PROFILES

icon, _log_name = LOG_PROFILES["DROPBOX"]
log = get_logger(_log_name, icon)

ACTION_REMOVE_PARTNER = "remove_partner"
ACTION_ADD_PARTNER = "add_partner"
HOLD_MARK = "HOLD"
MISSING_OTLEZKA_DATE_TEXT = "Нет даты отлёжки"
MISSING_OTLEZKA_STATUS = "НЕТ ДАТЫ ОТЛЁЖКИ"

STATUS_VKLUCHENO = "ВКЛЮЧЕНО"
STATUS_PROPUSHENO = "ПРОПУЩЕНО"
STATUS_OSHIBKA = "ОШИБКА"
STATUS_HOLD = "HOLD"
STATUS_OZHIDAET = "ОЖИДАЕТ"
STATUS_K_VKLUCHENIYU = "К ВКЛЮЧЕНИЮ"
STATUS_PROSROCHENO = "ПРОСРОЧЕНО"

SHEET_ALL_RESULTS = "all_results"
SHEET_RUNS = "runs"
SHEET_HOLD = "hold"
SHEET_OTLEZKA = "Отлёжка"

LEGACY_VALUE_COLUMN = "value"
OPERATION_DATE_COLUMN = "Дата операции"
DISABLE_DATE_COLUMN = "Дата отключения"
OPERATION_DATE_NUMBER_FORMAT = "dd.mm.yyyy"

ALL_RESULTS_COLUMNS = [
    OPERATION_DATE_COLUMN,
    DISABLE_DATE_COLUMN,
    "Дата включения",
    "Статус включения",
    "Включено",
    "Комментарий включения",
    "card",
    "partner",
    "action",
    "status",
    "comment",
    "hold",
]

RUNS_COLUMNS = [
    "started_at",
    "finished_at",
    "input_rows",
    "success_rows",
    "failed_rows",
    "skipped_rows",
    "output_file",
]

HOLD_COLUMNS = ["Дата добавления", "card", "partner", "comment"]
OTLEZKA_COLUMNS = ["partner", "Полные дни", "comment"]

LEGACY_ALL_RESULTS_COLUMNS = [
    "run_id",
    "run_started_at",
    "run_finished_at",
    "operator_profile",
    "source",
    "input_file",
    "output_file",
    "telegram_chat_id",
    "telegram_user_id",
    "row_index",
    "Дата отключения",
    "card",
    "action",
    "value",
    "status",
    "comment",
]

LEGACY_RUNS_COLUMNS = [
    "run_id",
    "started_at",
    "finished_at",
    "source",
    "operator_profile",
    "input_rows",
    "success_rows",
    "failed_rows",
    "skipped_rows",
    "output_file",
    "telegram_chat_id",
    "telegram_user_id",
]

RESULT_SOURCE_COLUMNS = [
    OPERATION_DATE_COLUMN,
    DISABLE_DATE_COLUMN,
    "card",
    "action",
    "value",
    "status",
    "comment",
]


def result_row_dates(action: str, status: str, processed_at: datetime) -> tuple[str, str]:
    """Return (operation_date, disable_date) for a registry result row."""
    operation_date = processed_at.strftime(EXCEL_DATE_FORMAT)
    action_norm = (action or "").strip().lower()
    status_norm = (status or "").strip().upper()
    if action_norm == ACTION_REMOVE_PARTNER and status_norm == "OK":
        return operation_date, processed_at.strftime(EXCEL_DATETIME_FORMAT)
    return operation_date, ""


def parse_operation_date(value: object) -> date | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text or text.lower() in {"nan", "nat", "none"}:
        return None
    for fmt in (EXCEL_DATETIME_FORMAT, EXCEL_DATE_FORMAT):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None

RED_FILL = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")

WARN_MESSAGE_TEMPLATE = (
    "⚠️ Не настроена отлёжка\n\n"
    "Партнёр:\n{partner}\n\n"
    "Файл:\nwallet_editor.xlsx\n\n"
    "Лист:\nОтлёжка\n\n"
    "Дата включения не рассчитана."
)


def _state_dir() -> Path:
    import os

    return Path(os.getenv("STATE_DIR", "/data/state"))


def wallet_editor_results_dir() -> Path:
    return _state_dir() / "wallet_editor" / "results"


def wallet_editor_outbox_dir() -> Path:
    return _state_dir() / "wallet_editor" / "outbox"


def durable_result_path(run_id: str) -> Path:
    return wallet_editor_results_dir() / f"{run_id}.xlsx"


def warned_partners_path() -> Path:
    return _state_dir() / "wallet_editor" / "missing_hold_days_warned.json"


def processed_run_ids_path() -> Path:
    return _state_dir() / "wallet_editor" / "registry_processed_run_ids.json"


def outbox_index_path() -> Path:
    return wallet_editor_outbox_dir() / "index.json"


OUTBOX_STATUS_PENDING = "pending"
OUTBOX_STATUS_SYNCING = "syncing"
OUTBOX_STATUS_SYNCED = "synced"
OUTBOX_STATUS_FAILED = "failed"

OUTBOX_ACTIVE_STATUSES = frozenset(
    {OUTBOX_STATUS_PENDING, OUTBOX_STATUS_SYNCING, OUTBOX_STATUS_FAILED}
)


def _empty_sheet(columns: list[str]) -> pd.DataFrame:
    return pd.DataFrame(columns=columns)


def _cell_str(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, datetime):
        if value.tzinfo is None:
            return value.strftime(EXCEL_DATETIME_FORMAT)
        return value.strftime(EXCEL_DATETIME_FORMAT)
    if isinstance(value, date):
        return value.strftime(EXCEL_DATE_FORMAT)
    return str(value).strip()


def _normalize_key(value: object) -> str:
    return _cell_str(value).casefold()


def partner_from_row(action: str, value: str) -> str:
    action_norm = (action or "").strip().lower()
    if action_norm in {ACTION_REMOVE_PARTNER, ACTION_ADD_PARTNER}:
        return (value or "").strip()
    return ""


def parse_disable_datetime(value: object) -> datetime | None:
    text = _cell_str(value)
    if not text:
        return None
    for fmt in (EXCEL_DATETIME_FORMAT, EXCEL_DATE_FORMAT):
        try:
            dt = datetime.strptime(text, fmt)
            return dt
        except ValueError:
            continue
    return None


def load_hold_pairs(hold_df: pd.DataFrame) -> set[tuple[str, str]]:
    if hold_df is None or hold_df.empty:
        return set()
    pairs: set[tuple[str, str]] = set()
    for _, row in hold_df.iterrows():
        card = _cell_str(row.get("card", ""))
        partner = _cell_str(row.get("partner", ""))
        if card and partner:
            pairs.add((_normalize_key(card), _normalize_key(partner)))
    return pairs


def load_otlezka_days(otlezka_df: pd.DataFrame) -> dict[str, int]:
    if otlezka_df is None or otlezka_df.empty:
        return {}
    result: dict[str, int] = {}
    for _, row in otlezka_df.iterrows():
        partner = _cell_str(row.get("partner", ""))
        if not partner:
            continue
        raw_days = row.get("Полные дни", "")
        if raw_days is None or (isinstance(raw_days, float) and pd.isna(raw_days)):
            continue
        try:
            days = int(float(raw_days))
        except (TypeError, ValueError):
            continue
        result[_normalize_key(partner)] = days
    return result


def is_legacy_all_results(df: pd.DataFrame) -> bool:
    return "run_id" in df.columns


def migrate_legacy_all_results(df: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    for _, row in df.iterrows():
        action = _cell_str(row.get("action", ""))
        value = _cell_str(row.get("value", ""))
        entry = {col: "" for col in ALL_RESULTS_COLUMNS}
        entry[OPERATION_DATE_COLUMN] = _cell_str(row.get(OPERATION_DATE_COLUMN, ""))
        entry[DISABLE_DATE_COLUMN] = _cell_str(row.get(DISABLE_DATE_COLUMN, ""))
        from integrations.wallet_editor_registry_xlsx import card_as_text

        entry["card"] = card_as_text(row.get("card", ""))
        entry["action"] = action
        entry["status"] = _cell_str(row.get("status", ""))
        entry["comment"] = _cell_str(row.get("comment", ""))
        entry["partner"] = partner_from_row(action, value)
        rows.append(entry)
    if not rows:
        return _empty_sheet(ALL_RESULTS_COLUMNS)
    return pd.DataFrame(rows, columns=ALL_RESULTS_COLUMNS)


def _migrate_partner_from_legacy_value(out: pd.DataFrame) -> pd.DataFrame:
    if LEGACY_VALUE_COLUMN not in out.columns:
        return out
    for idx in out.index:
        partner = _cell_str(out.at[idx, "partner"])
        if partner:
            continue
        action = _cell_str(out.at[idx, "action"])
        legacy_value = _cell_str(out.at[idx, LEGACY_VALUE_COLUMN])
        migrated = partner_from_row(action, legacy_value)
        if migrated:
            out.at[idx, "partner"] = migrated
    return out.drop(columns=[LEGACY_VALUE_COLUMN], errors="ignore")


def normalize_all_results(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return _empty_sheet(ALL_RESULTS_COLUMNS)
    if is_legacy_all_results(df):
        df = migrate_legacy_all_results(df)
    out = df.copy()
    for col in ALL_RESULTS_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    out = _migrate_partner_from_legacy_value(out)
    for col in ALL_RESULTS_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    return out[ALL_RESULTS_COLUMNS]


def normalize_sheet(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    if df is None or df.empty:
        return _empty_sheet(columns)
    out = df.copy()
    for col in columns:
        if col not in out.columns:
            out[col] = ""
    return out[columns]


def rows_from_result_excel(result_df: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    for _, row in result_df.iterrows():
        action = _cell_str(row.get("action", ""))
        value = _cell_str(row.get("value", ""))
        entry = {col: "" for col in ALL_RESULTS_COLUMNS}
        for col in RESULT_SOURCE_COLUMNS:
            if col == LEGACY_VALUE_COLUMN:
                continue
            entry[col] = _cell_str(row[col]) if col in result_df.columns else ""
        entry["partner"] = partner_from_row(action, value)
        rows.append(entry)
    if not rows:
        return _empty_sheet(ALL_RESULTS_COLUMNS)
    return pd.DataFrame(rows, columns=ALL_RESULTS_COLUMNS)


def lifecycle_status(
    *,
    vklyucheno: str,
    hold: str,
    reenable_date: str,
    missing_otlezka: bool,
    today: date,
) -> str:
    v = (vklyucheno or "").strip().upper()
    if v == "OK":
        return STATUS_VKLUCHENO
    if v == "SKIP":
        return STATUS_PROPUSHENO
    if v == "FAIL":
        return STATUS_OSHIBKA
    if (hold or "").strip().upper() == HOLD_MARK:
        return STATUS_HOLD
    if missing_otlezka:
        return MISSING_OTLEZKA_STATUS
    if not reenable_date or reenable_date == MISSING_OTLEZKA_DATE_TEXT:
        return ""
    try:
        rd = datetime.strptime(reenable_date, EXCEL_DATE_FORMAT).date()
    except ValueError:
        return ""
    if rd > today:
        return STATUS_OZHIDAET
    if rd == today:
        return STATUS_K_VKLUCHENIYU
    if rd < today:
        return STATUS_PROSROCHENO
    return ""


def recalculate_all_results(
    all_results: pd.DataFrame,
    hold_df: pd.DataFrame,
    otlezka_df: pd.DataFrame,
    *,
    today: date | None = None,
) -> tuple[pd.DataFrame, set[str]]:
    """Recalculate lifecycle fields for every row. Returns (df, partners_missing_otlezka)."""
    today = today or now_msk().date()
    df = normalize_all_results(all_results)
    hold_pairs = load_hold_pairs(hold_df)
    otlezka_days = load_otlezka_days(otlezka_df)
    missing_partners: set[str] = set()

    for idx in df.index:
        action = _cell_str(df.at[idx, "action"])
        status = _cell_str(df.at[idx, "status"])
        card = _cell_str(df.at[idx, "card"])
        partner = _cell_str(df.at[idx, "partner"])
        if not partner:
            legacy_value = ""
            if LEGACY_VALUE_COLUMN in df.columns:
                legacy_value = _cell_str(df.at[idx, LEGACY_VALUE_COLUMN])
            partner = partner_from_row(action, legacy_value)
        df.at[idx, "partner"] = partner

        vklyucheno = _cell_str(df.at[idx, "Включено"])
        komment = _cell_str(df.at[idx, "Комментарий включения"])

        on_hold = bool(card and partner and (_normalize_key(card), _normalize_key(partner)) in hold_pairs)
        if on_hold:
            df.at[idx, "hold"] = HOLD_MARK
            df.at[idx, "Дата включения"] = ""
            df.at[idx, "Статус включения"] = lifecycle_status(
                vklyucheno=vklyucheno,
                hold=HOLD_MARK,
                reenable_date="",
                missing_otlezka=False,
                today=today,
            )
            continue

        df.at[idx, "hold"] = ""

        eligible = (
            status.upper() == "OK"
            and action.lower() == ACTION_REMOVE_PARTNER
            and bool(_cell_str(df.at[idx, "Дата отключения"]))
            and bool(partner)
        )
        if not eligible:
            df.at[idx, "Дата включения"] = ""
            df.at[idx, "Статус включения"] = (
                lifecycle_status(
                    vklyucheno=vklyucheno,
                    hold="",
                    reenable_date="",
                    missing_otlezka=False,
                    today=today,
                )
                if vklyucheno
                else ""
            )
            continue

        disable_dt = parse_disable_datetime(df.at[idx, "Дата отключения"])
        partner_key = _normalize_key(partner)
        if partner_key not in otlezka_days:
            df.at[idx, "Дата включения"] = MISSING_OTLEZKA_DATE_TEXT
            missing_partners.add(partner)
            df.at[idx, "Статус включения"] = lifecycle_status(
                vklyucheno=vklyucheno,
                hold="",
                reenable_date=MISSING_OTLEZKA_DATE_TEXT,
                missing_otlezka=True,
                today=today,
            )
            continue

        if disable_dt is None:
            df.at[idx, "Дата включения"] = ""
            df.at[idx, "Статус включения"] = lifecycle_status(
                vklyucheno=vklyucheno,
                hold="",
                reenable_date="",
                missing_otlezka=False,
                today=today,
            )
            continue

        reenable = (disable_dt.date() + timedelta(days=otlezka_days[partner_key])).strftime(EXCEL_DATE_FORMAT)
        df.at[idx, "Дата включения"] = reenable
        df.at[idx, "Статус включения"] = lifecycle_status(
            vklyucheno=vklyucheno,
            hold="",
            reenable_date=reenable,
            missing_otlezka=False,
            today=today,
        )

    return df, missing_partners


def load_warned_partners() -> set[str]:
    path = warned_partners_path()
    if not path.exists():
        return set()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        items = data.get("warned_partners", [])
        return {_normalize_key(p) for p in items if _cell_str(p)}
    except Exception:
        return set()


def save_warned_partners(partners: Iterable[str]) -> None:
    path = warned_partners_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"warned_partners": sorted({_cell_str(p) for p in partners if _cell_str(p)})}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def sync_warned_partners_after_otlezka(
    otlezka_df: pd.DataFrame,
    warned: set[str],
) -> set[str]:
    configured = {_normalize_key(p) for p in load_otlezka_days(otlezka_df)}
    updated = {p for p in warned if p not in configured}
    save_warned_partners({_cell_str(p) for p in updated})
    return updated


def partners_to_warn(missing: set[str], warned: set[str]) -> list[str]:
    result: list[str] = []
    for partner in sorted(missing, key=lambda p: p.casefold()):
        key = _normalize_key(partner)
        if key and key not in warned:
            result.append(partner)
    return result


class ProcessedRunIdsCorruptedError(RuntimeError):
    """registry_processed_run_ids.json exists but cannot be parsed safely."""


@dataclass(frozen=True, slots=True)
class ProcessedRunIdsLoadResult:
    run_ids: frozenset[str]
    corrupted: bool
    error: str | None = None


def load_processed_run_ids_result() -> ProcessedRunIdsLoadResult:
    """Load processed run ids; corrupted=True when file exists but is unreadable."""
    path = processed_run_ids_path()
    if not path.exists():
        return ProcessedRunIdsLoadResult(frozenset(), False, None)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError("processed_run_ids root must be a JSON object")
        raw_ids = data.get("run_ids", [])
        if not isinstance(raw_ids, list):
            raise ValueError("processed_run_ids.run_ids must be a JSON array")
        run_ids = frozenset(str(x) for x in raw_ids if str(x).strip())
        return ProcessedRunIdsLoadResult(run_ids, False, None)
    except Exception as exc:
        log.exception("[WalletEditorRegistry] processed_run_ids corrupted path=%s", path)
        return ProcessedRunIdsLoadResult(frozenset(), True, str(exc))


def is_processed_run_ids_corrupted() -> bool:
    return load_processed_run_ids_result().corrupted


def processed_run_ids_corruption_error() -> str | None:
    return load_processed_run_ids_result().error


def load_processed_run_ids() -> set[str]:
    result = load_processed_run_ids_result()
    if result.corrupted:
        raise ProcessedRunIdsCorruptedError(
            result.error or "registry_processed_run_ids.json is corrupted"
        )
    return set(result.run_ids)


def save_processed_run_ids(run_ids: set[str]) -> None:
    path = processed_run_ids_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps({"run_ids": sorted(run_ids)}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def result_row_fingerprint(row: pd.Series | dict[str, object]) -> str:
    """Stable fingerprint for dedup / repair detection."""
    if isinstance(row, dict):
        get = row.get
    else:
        get = row.get

    action = _cell_str(get("action", "")).lower()
    value = _cell_str(get("value", ""))
    partner = _cell_str(get("partner", "")) or partner_from_row(action, value)
    parts = [
        _normalize_key(get("card", "")),
        _normalize_key(partner),
        action,
        _cell_str(get(DISABLE_DATE_COLUMN, "")),
        _cell_str(get("status", "")).upper(),
        _cell_str(get("comment", "")),
    ]
    return "|".join(parts)


def fingerprints_from_result_excel(result_path: str) -> list[str]:
    from integrations.wallet_editor_registry_xlsx import card_as_text

    df = pd.read_excel(
        result_path,
        engine="openpyxl",
        converters={"card": card_as_text},
    )
    return [result_row_fingerprint(df.loc[idx]) for idx in df.index]


def registry_row_fingerprints(all_results: pd.DataFrame) -> set[str]:
    df = normalize_all_results(all_results)
    if df.empty:
        return set()
    return {result_row_fingerprint(df.loc[idx]) for idx in df.index}


def missing_result_fingerprints(
    result_path: str,
    all_results: pd.DataFrame,
) -> list[str]:
    """Fingerprints from result file that are absent in registry all_results."""
    present = registry_row_fingerprints(all_results)
    return missing_result_fingerprints_in_set(result_path, present)


def missing_result_fingerprints_in_set(
    result_path: str,
    present: set[str] | frozenset[str],
) -> list[str]:
    """Fingerprints from result file that are absent in the provided fingerprint set."""
    return [fp for fp in fingerprints_from_result_excel(result_path) if fp not in present]


def filter_rows_not_in_registry(
    new_rows: pd.DataFrame,
    all_results: pd.DataFrame,
) -> pd.DataFrame:
    """Return only rows whose fingerprint is not already in all_results."""
    if new_rows is None or new_rows.empty:
        return _empty_sheet(ALL_RESULTS_COLUMNS)
    present = registry_row_fingerprints(all_results)
    keep_idx: list[int] = []
    for idx in new_rows.index:
        if result_row_fingerprint(new_rows.loc[idx]) not in present:
            keep_idx.append(int(idx))
    if not keep_idx:
        return _empty_sheet(ALL_RESULTS_COLUMNS)
    return new_rows.loc[keep_idx].reset_index(drop=True)


def run_id_already_processed(
    run_id: str,
    runs_df: pd.DataFrame,
    *,
    result_path: str | None = None,
    all_results_df: pd.DataFrame | None = None,
) -> bool:
    """
  Return True when append should be skipped as duplicate.

  If run_id is in processed_run_ids but result rows are missing from registry
  (restored workbook trap), returns False so repair re-append can proceed.
    """
    if is_processed_run_ids_corrupted():
        raise ProcessedRunIdsCorruptedError(
            processed_run_ids_corruption_error()
            or "registry_processed_run_ids.json is corrupted"
        )

    in_processed = run_id in load_processed_run_ids()
    in_legacy_runs = False
    if runs_df is not None and not runs_df.empty and "run_id" in runs_df.columns:
        in_legacy_runs = run_id in runs_df["run_id"].astype(str).tolist()

    if not in_processed and not in_legacy_runs:
        return False

    if result_path and all_results_df is not None:
        missing = missing_result_fingerprints(result_path, all_results_df)
        if missing:
            return False

    return True


def mark_run_processed(run_id: str) -> None:
    if is_processed_run_ids_corrupted():
        log.error(
            "[WalletEditorRegistry] skip mark_run_processed run_id=%s: processed_run_ids corrupted",
            run_id,
        )
        return
    ids = load_processed_run_ids()
    ids.add(run_id)
    save_processed_run_ids(ids)


def unmark_run_processed(run_id: str) -> None:
    if is_processed_run_ids_corrupted():
        log.error(
            "[WalletEditorRegistry] skip unmark_run_processed run_id=%s: processed_run_ids corrupted",
            run_id,
        )
        return
    ids = load_processed_run_ids()
    if run_id in ids:
        ids.discard(run_id)
        save_processed_run_ids(ids)


def build_runs_row(
    *,
    started_at: datetime,
    finished_at: datetime,
    input_rows: int,
    stats_ok: int,
    stats_fail: int,
    stats_skip: int,
    output_file: str,
) -> pd.DataFrame:
    from core.datetime_utils import ensure_aware_msk

    row = {
        "started_at": ensure_aware_msk(started_at).strftime(EXCEL_DATETIME_FORMAT),
        "finished_at": ensure_aware_msk(finished_at).strftime(EXCEL_DATETIME_FORMAT),
        "input_rows": input_rows,
        "success_rows": stats_ok,
        "failed_rows": stats_fail,
        "skipped_rows": stats_skip,
        "output_file": output_file,
    }
    return pd.DataFrame([row], columns=RUNS_COLUMNS)


def migrate_legacy_runs(runs_df: pd.DataFrame) -> pd.DataFrame:
    if runs_df is None or runs_df.empty:
        return _empty_sheet(RUNS_COLUMNS)
    out = runs_df.copy()
    for col in RUNS_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    return out[RUNS_COLUMNS]


def apply_missing_otlezka_red_fill(local_path: Path) -> None:
    wb = load_workbook(local_path)
    if SHEET_ALL_RESULTS not in wb.sheetnames:
        wb.save(local_path)
        return
    ws = wb[SHEET_ALL_RESULTS]
    headers = [cell.value for cell in ws[1]]
    if "Дата включения" not in headers:
        wb.save(local_path)
        return
    col_idx = headers.index("Дата включения") + 1
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        if _cell_str(cell.value) == MISSING_OTLEZKA_DATE_TEXT:
            cell.fill = RED_FILL
    wb.save(local_path)
