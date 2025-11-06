# analyzers/payout.py
import os
import yaml
import gc
import psutil
from utils.logger import logger
import pandas as pd
from utils.excel_utils import style_worksheet, write_df_to_sheet
from openpyxl import Workbook
from datetime import datetime

# путь к конфигу
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(BASE_DIR, "..", "config", "payout_config.yaml")

# загрузка конфига
with open(CONFIG_PATH, "r", encoding="utf-8") as f:
    CONFIG = yaml.safe_load(f) or {}


def log_memory(step: str):
    """Логирование использования памяти"""
    process = psutil.Process(os.getpid())
    mb = process.memory_info().rss / 1024 / 1024
    logger.info(f"🧠 {step}: {mb:.1f} MB")


def run(payout_file: str, card_files: list, *args, **kwargs):
    logger.info(f"[payout] 🚀 Начало анализа: {os.path.basename(payout_file)}")
    log_memory("Начало payout анализа")

    # Проверяем, что есть файл направлений (cd)
    if not card_files:
        logger.warning("[payout] ⚠️ Не найден cd-файл (направления карт).")
        return {}

    cd_file = card_files[0]

    # === 1️⃣ ОПТИМИЗИРОВАННАЯ Загрузка payout-файла ===
    try:
        # Оптимизированные типы данных и только нужные колонки
        usecols = ["карта", "статус", "инфо", "дата/время создания", "телефон", "выделено"]
        dtype_optimized = {
            'карта': 'string',
            'статус': 'category',  # ← Категориальные данные экономят память!
            'инфо': 'string',
            'телефон': 'string',
            'выделено': 'string',
            'дата/время создания': 'string'  # Конвертируем позже
        }

        df_payout = pd.read_excel(payout_file, usecols=usecols, dtype=dtype_optimized)

        # Конвертируем дату после чтения
        df_payout["дата/время создания"] = pd.to_datetime(
            df_payout["дата/время создания"], format="%d.%m.%Y %H:%M:%S", errors="coerce"
        )

        logger.info(f"[payout] Загружен payout-файл: {os.path.basename(payout_file)} ({len(df_payout)} строк)")
        log_memory("После загрузки payout")
    except Exception as e:
        logger.exception(f"[payout] ❌ Ошибка чтения payout-файла: {e}")
        return {}

    # === 2️⃣ ОПТИМИЗИРОВАННАЯ Загрузка cd-файла ===
    try:
        # Только нужные колонки для cd-файла
        cd_usecols = ["карта", "направление"]
        cd_dtype = {'карта': 'string', 'направление': 'category'}

        df_cd = pd.read_excel(cd_file, usecols=cd_usecols, dtype=cd_dtype)
        df_cd.columns = df_cd.columns.str.strip()
        df_cd.rename(columns=lambda c: c.strip().lower(), inplace=True)

        logger.info(f"[payout] Загружен cd-файл: {os.path.basename(cd_file)} ({len(df_cd)} строк)")
        log_memory("После загрузки cd файла")
    except Exception as e:
        logger.exception(f"[payout] ❌ Ошибка чтения cd-файла: {e}")
        return {}

    # === 3️⃣ Нормализация столбцов payout ===
    df_payout.columns = df_payout.columns.str.strip()
    df_payout.rename(columns=lambda c: c.strip().lower(), inplace=True)

    # убираем лишние пробелы в "карта"
    if "карта" in df_payout.columns:
        df_payout["карта"] = df_payout["карта"].astype(str).str.strip()

    # убираем лишние пробелы в "карта" cd-файла
    if "карта" in df_cd.columns:
        df_cd["карта"] = df_cd["карта"].astype(str).str.strip()

    logger.info("[payout] ✅ Данные загружены и нормализованы.")

    # === 4️⃣ Исключаем карты с направлением IN ===
    if "направление" in df_cd.columns and "карта" in df_cd.columns:
        # Используем set для быстрого поиска
        in_cards = set(df_cd.loc[df_cd["направление"].str.upper() == "IN", "карта"])
        before = len(df_payout)
        df_payout = df_payout[~df_payout["карта"].isin(in_cards)].copy()
        logger.info(f"[payout] 🧭 Исключено {before - len(df_payout)} строк по направлению IN.")

        # Очищаем память от временных объектов
        del in_cards
        gc.collect()
    else:
        logger.warning("[payout] ⚠️ В cd-файле отсутствуют нужные колонки: 'Карта' и 'Направление'.")

    # === 5️⃣ Загружаем настройки из YAML ===
    payouts_errors = CONFIG.get("PayoutsErrors", {}) or {}
    ignore_errors = CONFIG.get("IgnoreErrors", []) or []

    # Преобразуем в нижний регистр для удобства поиска
    payouts_errors_norm = {k.lower().strip(): v for k, v in payouts_errors.items()}
    ignore_errors_norm = [x.lower().strip() for x in ignore_errors]

    logger.info(
        f"[payout] 📘 Загружено {len(payouts_errors_norm)} известных ошибок и {len(ignore_errors_norm)} игнорируемых.")

    # === 6️⃣ Фильтрация строк payout-файла ===
    # интересуют только статусы "ошибка" и "оплачен"
    valid_statuses = {"ошибка", "оплачен"}
    df_payout["status"] = df_payout["status"].astype(str).str.lower().str.strip()

    before = len(df_payout)
    df_payout = df_payout[df_payout["status"].isin(valid_statuses)].copy()
    logger.info(f"[payout] 🧮 Оставлено {len(df_payout)} строк (из {before}), где статус ошибка/оплачен.")
    log_memory("После фильтрации статусов")

    # === 7️⃣ ОПТИМИЗИРОВАННЫЙ Анализ последовательностей ошибок ===
    problem_cards = []  # для списка "Перевести в IN"
    check_cards = set()  # для списка "Проверить"

    # Проверяем, что нужные столбцы есть
    required_cols = {"карта", "status", "info", "дата/время создания", "телефон", "выделено"}
    if not required_cols.issubset(df_payout.columns):
        logger.warning(
            f"[payout] ⚠️ В payout-файле отсутствуют нужные колонки: {required_cols - set(df_payout.columns)}")
        return {}

    # 🔧 ОПТИМИЗАЦИЯ: Обработка карт батчами по 1000
    unique_cards = df_payout["карта"].unique()
    total_cards = len(unique_cards)
    BATCH_SIZE = 1000

    logger.info(f"[payout] 🔄 Обработка {total_cards} уникальных карт батчами по {BATCH_SIZE}")

    for i in range(0, total_cards, BATCH_SIZE):
        batch_cards = unique_cards[i:i + BATCH_SIZE]
        batch_df = df_payout[df_payout["карта"].isin(batch_cards)].copy()

        # Группируем по карте внутри батча
        for card, group in batch_df.groupby("карта"):
            group = group.sort_values("дата/время создания", ascending=False)

            consecutive = 0
            last_info = None
            last_datetime = None
            phone = str(group["телефон"].iloc[0]) if "телефон" in group.columns else ""
            pool = str(group["выделено"].iloc[0]) if "выделено" in group.columns else ""

            for _, row in group.iterrows():
                status = str(row["status"]).lower().strip()
                info = str(row["info"]).lower().strip()

                # Если статус "оплачен" → серия обрывается
                if status == "оплачен":
                    consecutive = 0
                    last_info = None
                    continue

                # Если статус "ошибка"
                if status == "ошибка":
                    matched_error = next((err for err in payouts_errors_norm if err in info), None)

                    if matched_error:
                        # ✅ приоритет у известных ошибок
                        threshold = payouts_errors_norm[matched_error].get("threshold", 1)
                        if last_info == matched_error:
                            consecutive += 1
                        else:
                            consecutive = 1
                            last_info = matched_error
                        last_datetime = row["дата/время создания"]

                        if consecutive >= threshold:
                            problem_cards.append({
                                "Карта": card,
                                "Телефон": phone,
                                "Выделено": pool,
                                "Info": matched_error,
                                "Количество подряд ошибок": consecutive,
                                "Последняя дата ошибки": (
                                    pd.to_datetime(last_datetime).strftime("%d.%m.%Y %H:%M:%S")
                                    if pd.notna(last_datetime) else ""
                                ),
                            })
                            break  # карта уже попала в список — дальше не анализируем

                    else:
                        # 🧩 только если ошибка НЕизвестная, проверяем игнор
                        if any(phrase in info for phrase in ignore_errors_norm):
                            continue  # просто пропускаем эту строку (игнорируем)

                        # неизвестная ошибка → в список "Проверить"
                        err_dt = row.get("дата/время создания")
                        if pd.notna(err_dt):
                            try:
                                err_dt = pd.to_datetime(err_dt).strftime("%d.%m.%Y %H:%M:%S")
                            except Exception:
                                err_dt = str(err_dt)
                        else:
                            err_dt = ""

                        check_cards.add((card, phone, pool, info, err_dt))

        # 🔧 ОЧИСТКА ПАМЯТИ после каждого батча
        del batch_df
        gc.collect()

        if i % 5000 == 0:  # Логируем прогресс каждые 5000 карт
            log_memory(f"Обработано {i}/{total_cards} карт")
            logger.info(f"[payout] 📊 Прогресс: обработано {i}/{total_cards} карт")

    # === 8️⃣ Формирование итоговых таблиц ===
    # Преобразуем списки в DataFrame
    df_problem = pd.DataFrame(problem_cards)
    df_check = pd.DataFrame(
        list(check_cards),
        columns=["Карта", "Телефон", "Выделено", "Info", "Дата ошибки"]
    )

    # Сортировка по дате (от новых к старым)
    if not df_check.empty and "Дата ошибки" in df_check.columns:
        df_check["_sort_key"] = pd.to_datetime(df_check["Дата ошибки"], errors="coerce")
        df_check.sort_values("_sort_key", ascending=False, na_position="last", inplace=True)
        df_check.drop(columns=["_sort_key"], inplace=True)

    if not df_problem.empty and "Последняя дата ошибки" in df_problem.columns:
        df_problem["_sort_key"] = pd.to_datetime(
            df_problem["Последняя дата ошибки"], format="%d.%m.%Y %H:%M:%S", errors="coerce"
        )
        df_problem.sort_values("_sort_key", ascending=False, inplace=True)
        df_problem.drop(columns=["_sort_key"], inplace=True)

    # Удаляем дубликаты по карте (оставляем самую свежую)
    if not df_problem.empty:
        df_problem.drop_duplicates(subset=["Карта"], inplace=True)
    if not df_check.empty:
        df_check.drop_duplicates(subset=["Карта"], inplace=True)

    # Страхуем формат дат перед сохранением
    if "Последняя дата ошибки" in df_problem.columns:
        df_problem["Последняя дата ошибки"] = df_problem["Последняя дата ошибки"].astype(str)
    if "Дата ошибки" in df_check.columns:
        df_check["Дата ошибки"] = df_check["Дата ошибки"].astype(str)

    log_memory("После формирования таблиц")

    # === 9️⃣ Формирование Excel-отчёта ===
    wb = Workbook()
    wb.remove(wb.active)

    if not df_problem.empty:
        write_df_to_sheet(wb, "Перевести в IN", df_problem)
    else:
        ws = wb.create_sheet("Перевести в IN")
        ws.append(["Нет карт для перевода в IN"])

    if not df_check.empty:
        write_df_to_sheet(wb, "Проверить", df_check)
    else:
        ws = wb.create_sheet("Проверить")
        ws.append(["Нет карт для проверки"])

    # Сохраняем отчёт во временный файл
    tmp_name = f"report_{os.path.basename(payout_file)}"
    tmp_path = os.path.join(os.getenv("TMP", "/tmp"), tmp_name)
    wb.save(tmp_path)
    logger.info(f"[payout] 📁 Отчёт сохранён: {tmp_path}")

    # === 🔟 Отправляем файл в Telegram и перемещаем ===
    from integrations.telegram_bot import send_message_sync, send_file_sync
    from integrations.dropbox_watcher import move_file

    try:
        # 1️⃣ Отправляем отчёт в Telegram
        if os.path.exists(tmp_path):
            caption = (
                f"📊 Отчёт по {os.path.basename(payout_file)}\n"
                f"Карт на перевод в in: {len(df_problem)}\n"
                f"Карт на проверку: {len(df_check)}"
            )
            send_file_sync(tmp_path, caption=caption)
            logger.info(f"[payout] 📤 Отчёт отправлен в Telegram: {tmp_path}")

        # 2️⃣ Перемещаем исходный payout-файл в /processed с добавлением даты
        DROPBOX_PROCESSED_PATH = os.getenv("DROPBOX_PROCESSED_PATH")
        today_str = datetime.now().strftime("(%d.%m.%Y)")
        base_name = os.path.basename(payout_file)
        new_name = (
            base_name[:-5] + f"_{today_str}.xlsx"
            if base_name.lower().endswith(".xlsx")
            else base_name + f"_{today_str}.xlsx"
        )

        src_dropbox_path = os.path.join(os.getenv("DROPBOX_INPUT_PATH"), base_name)
        dest_dropbox_path = os.path.join(DROPBOX_PROCESSED_PATH, new_name)

        move_file(src_dropbox_path, dest_dropbox_path)
        logger.info(f"[payout] ✅ Файл {base_name} перемещён в /processed как {new_name}")

    except Exception as e:
        logger.exception(f"[payout] ⚠️ Ошибка при отправке/перемещении файла: {e}")

    # === 11️⃣ ФИНАЛЬНАЯ ОЧИСТКА ПАМЯТИ ===
    del df_payout, df_cd, df_problem, df_check
    gc.collect()
    log_memory("Конец payout анализа")

    # === 12️⃣ Возвращаем сводку ===
    logger.info(
        f"[payout] ✅ Анализ завершён: перевод={len(df_problem)}, проверка={len(df_check)}"
    )
    return {
        "workbook": wb,
        "report_path": tmp_path,
        "summary": {
            "Карты на перевод": len(df_problem),
            "Карты на проверку": len(df_check)
        }
    }