#!/usr/bin/env python3
"""Local smoke test for Wallet Editor lifecycle registry (no Dropbox)."""
from __future__ import annotations

import json
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from integrations.wallet_editor_registry_lifecycle import (  # noqa: E402
    ALL_RESULTS_COLUMNS,
    HOLD_COLUMNS,
    HOLD_MARK,
    MISSING_OTLEZKA_DATE_TEXT,
    MISSING_OTLEZKA_STATUS,
    OTLEZKA_COLUMNS,
    RUNS_COLUMNS,
    SHEET_ALL_RESULTS,
    SHEET_HOLD,
    SHEET_OTLEZKA,
    SHEET_RUNS,
    STATUS_K_VKLUCHENIYU,
    STATUS_OZHIDAET,
    STATUS_PROSROCHENO,
    WARN_MESSAGE_TEMPLATE,
    apply_missing_otlezka_red_fill,
    build_runs_row,
    load_warned_partners,
    normalize_sheet,
    partners_to_warn,
    recalculate_all_results,
    save_warned_partners,
    sync_warned_partners_after_otlezka,
)
from integrations.wallet_editor_registry import (  # noqa: E402
    _save_workbook,
    _process_missing_otlezka_warnings,
)
from automation.audit import Stats  # noqa: E402
from automation.runtime import WalletEditorTask  # noqa: E402

MSK = ZoneInfo("Europe/Moscow")
TODAY = date(2026, 6, 3)


def _dt(d: date, h: int = 12) -> str:
    return datetime(d.year, d.month, d.day, h, 0, 0, tzinfo=MSK).strftime("%d.%m.%Y %H:%M:%S")


def _row(
    *,
    card: str,
    partner: str,
    disable: str,
    status: str = "OK",
    action: str = "remove_partner",
) -> dict[str, str]:
    return {
        "Дата отключения": disable,
        "Дата включения": "",
        "Статус включения": "",
        "Включено": "",
        "Комментарий включения": "",
        "card": card,
        "partner": "",
        "action": action,
        "value": partner,
        "status": status,
        "comment": "",
        "hold": "",
    }


def main() -> int:
    results: list[tuple[str, str, str]] = []
    workbook_path: Path | None = None

    def record(scenario: str, ok: bool, detail: str) -> None:
        results.append((scenario, "PASS" if ok else "FAIL", detail))

    hold_df = normalize_sheet(pd.DataFrame(), HOLD_COLUMNS)
    otlezka_df = normalize_sheet(pd.DataFrame(), OTLEZKA_COLUMNS)

    # --- Scenario 1: Missing Отлёжка ---
    all_df = pd.DataFrame([_row(card="1111111111111111", partner="TEST_PARTNER_A", disable=_dt(TODAY))])
    out, missing = recalculate_all_results(all_df, hold_df, otlezka_df, today=TODAY)
    r = out.iloc[0]
    s1_ok = (
        r["Дата включения"] == MISSING_OTLEZKA_DATE_TEXT
        and r["Статус включения"] == MISSING_OTLEZKA_STATUS
        and "TEST_PARTNER_A" in missing
    )
    record("S1 Missing Отлёжка", s1_ok, f"date={r['Дата включения']!r} status={r['Статус включения']!r}")

    messages: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        state = Path(tmp) / "state"
        with patch.dict("os.environ", {"STATE_DIR": str(state)}):
            save_warned_partners(set())
            task = WalletEditorTask(
                file_path="/tmp/in.xlsx",
                chat_id=-999,
                telegram_user_id=1,
                operator_profile="DENIS",
                source_file_name="smoke.xlsx",
                login="l",
                password="p",
                auth_state_path="/tmp/a.json",
                run_id="smoke-warn-1",
            )
            with patch(
                "integrations.wallet_editor_registry.send_message_sync",
                side_effect=lambda text, chat_id=None, **_kw: messages.append(text),
            ):
                _process_missing_otlezka_warnings(task, missing, otlezka_df)
    warn_ok = len(messages) == 1 and "TEST_PARTNER_A" in messages[0] and "Не настроена отлёжка" in messages[0]
    record("S1 Telegram warning", warn_ok, f"messages={len(messages)}")

    # --- Scenario 2: Отлёжка 3 days -> ОЖИДАЕТ ---
    otlezka_df = normalize_sheet(
        pd.DataFrame([{"partner": "TEST_PARTNER_A", "Полные дни": 3, "comment": ""}]),
        OTLEZKA_COLUMNS,
    )
    out2, _ = recalculate_all_results(all_df, hold_df, otlezka_df, today=TODAY)
    r2 = out2.iloc[0]
    expected_reenable = (TODAY + timedelta(days=3)).strftime("%d.%m.%Y")
    s2_ok = r2["Дата включения"] == expected_reenable and r2["Статус включения"] == STATUS_OZHIDAET
    record(
        "S2 Отлёжка 3 дня",
        s2_ok,
        f"date={r2['Дата включения']!r} status={r2['Статус включения']!r} expected={expected_reenable}",
    )

    # --- Scenario 3: ПРОСРОЧЕНО ---
    old_disable = _dt(TODAY - timedelta(days=10))
    all_s3 = pd.DataFrame([_row(card="2222", partner="TEST_PARTNER_A", disable=old_disable)])
    out3, _ = recalculate_all_results(all_s3, hold_df, otlezka_df, today=TODAY)
    r3 = out3.iloc[0]
    s3_ok = r3["Статус включения"] == STATUS_PROSROCHENO
    record("S3 ПРОСРОЧЕНО", s3_ok, f"disable={old_disable} reenable={r3['Дата включения']!r} status={r3['Статус включения']!r}")

    # --- Scenario 4: К ВКЛЮЧЕНИЮ ---
    disable_s4 = _dt(TODAY - timedelta(days=3))
    all_s4 = pd.DataFrame([_row(card="3333", partner="TEST_PARTNER_A", disable=disable_s4)])
    out4, _ = recalculate_all_results(all_s4, hold_df, otlezka_df, today=TODAY)
    r4 = out4.iloc[0]
    s4_ok = r4["Статус включения"] == STATUS_K_VKLUCHENIYU
    record("S4 К ВКЛЮЧЕНИЮ", s4_ok, f"reenable={r4['Дата включения']!r} status={r4['Статус включения']!r}")

    # --- Scenario 5: HOLD ---
    hold_df = normalize_sheet(
        pd.DataFrame([{"Дата добавления": _dt(TODAY), "card": "4444", "partner": "TEST_PARTNER_A", "comment": ""}]),
        HOLD_COLUMNS,
    )
    all_s5 = pd.DataFrame([_row(card="4444", partner="TEST_PARTNER_A", disable=_dt(TODAY))])
    out5, _ = recalculate_all_results(all_s5, hold_df, otlezka_df, today=TODAY)
    r5 = out5.iloc[0]
    s5_ok = r5["hold"] == HOLD_MARK and r5["Дата включения"] == "" and r5["Статус включения"] == HOLD_MARK
    record("S5 HOLD", s5_ok, f"hold={r5['hold']!r} date={r5['Дата включения']!r} status={r5['Статус включения']!r}")

    # --- Scenario 6: Recalculate all (B then add Отлёжка) ---
    hold_df = normalize_sheet(pd.DataFrame(), HOLD_COLUMNS)
    all_s6 = pd.DataFrame(
        [
            _row(card="5555", partner="TEST_PARTNER_B", disable=_dt(TODAY - timedelta(days=1))),
            _row(card="6666", partner="TEST_PARTNER_A", disable=_dt(TODAY)),
        ]
    )
    out6a, miss6 = recalculate_all_results(all_s6, hold_df, pd.DataFrame(), today=TODAY)
    b_row = out6a[out6a["partner"] == "TEST_PARTNER_B"].iloc[0]
    s6a_ok = b_row["Статус включения"] == MISSING_OTLEZKA_STATUS
    otlezka_s6 = normalize_sheet(
        pd.DataFrame(
            [
                {"partner": "TEST_PARTNER_A", "Полные дни": 3, "comment": ""},
                {"partner": "TEST_PARTNER_B", "Полные дни": 5, "comment": ""},
            ]
        ),
        OTLEZKA_COLUMNS,
    )
    out6b, _ = recalculate_all_results(out6a, hold_df, otlezka_s6, today=TODAY)
    b_after = out6b[out6b["partner"] == "TEST_PARTNER_B"].iloc[0]
    expected_b = (TODAY - timedelta(days=1) + timedelta(days=5)).strftime("%d.%m.%Y")
    s6b_ok = (
        len(out6b) == 2
        and b_after["Дата включения"] == expected_b
        and b_after["Статус включения"] != MISSING_OTLEZKA_STATUS
    )
    record("S6a B missing", s6a_ok, f"B status={b_row['Статус включения']!r}")
    record("S6b B after Отлёжка", s6b_ok, f"rows={len(out6b)} B date={b_after['Дата включения']!r}")

    # --- Scenario 7 & 8: workbook structure ---
    runs = build_runs_row(
        started_at=datetime(2026, 6, 3, 9, 0, tzinfo=MSK),
        finished_at=datetime(2026, 6, 3, 9, 5, tzinfo=MSK),
        input_rows=2,
        stats_ok=2,
        stats_fail=0,
        stats_skip=0,
        output_file="smoke_result.xlsx",
    )
    s7_ok = list(runs.columns) == RUNS_COLUMNS
    record("S7 runs columns", s7_ok, str(list(runs.columns)))

    final_all = out6b
    s8_ok = list(final_all.columns) == ALL_RESULTS_COLUMNS
    record("S8 all_results columns", s8_ok, str(list(final_all.columns)))

    with tempfile.TemporaryDirectory() as tmp:
        wb = Path(tmp) / "wallet_editor_smoke.xlsx"
        _save_workbook(wb, final_all, runs, hold_df, otlezka_s6)
        apply_missing_otlezka_red_fill(wb)
        workbook_path = wb
        with pd.ExcelFile(wb) as book:
            sheets = book.sheet_names
            red_check = "ok"
            if SHEET_ALL_RESULTS in sheets:
                df_check = pd.read_excel(book, SHEET_ALL_RESULTS)
                if (df_check["Дата включения"] == MISSING_OTLEZKA_DATE_TEXT).any():
                    from openpyxl import load_workbook as lw

                    ws = lw(wb)[SHEET_ALL_RESULTS]
                    headers = [c.value for c in ws[1]]
                    col = headers.index("Дата включения") + 1
                    has_red = any(
                        ws.cell(row=i, column=col).fill.start_color.rgb in ("FFFFC7CE", "00FFC7CE")
                        for i in range(2, ws.max_row + 1)
                        if ws.cell(row=i, column=col).value == MISSING_OTLEZKA_DATE_TEXT
                    )
                    red_check = "red_fill=yes" if has_red else "red_fill=no"
        record("Workbook sheets", set(sheets) == {SHEET_ALL_RESULTS, SHEET_RUNS, SHEET_HOLD, SHEET_OTLEZKA}, str(sheets))
        record("Red fill missing otlezka", "red_fill=yes" in red_check, red_check)

    print("=== Smoke results (today=%s) ===" % TODAY)
    print(f"{'Scenario':<28} {'Result':<6} Detail")
    print("-" * 70)
    failed = 0
    for name, status, detail in results:
        print(f"{name:<28} {status:<6} {detail}")
        if status == "FAIL":
            failed += 1
    print("-" * 70)
    print(f"Workbook: {workbook_path}")
    if workbook_path and workbook_path.exists():
        print("\nall_results snapshot:")
        print(pd.read_excel(workbook_path, SHEET_ALL_RESULTS).to_string(index=False))
        print("\nОтлёжка:")
        print(pd.read_excel(workbook_path, SHEET_OTLEZKA).to_string(index=False))
    verdict = "WALLET_EDITOR_LIFECYCLE_SMOKE_PASSED" if failed == 0 else "SMOKE_FAILED"
    print(f"\nVerdict: {verdict}")
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
