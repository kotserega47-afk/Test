# utils/excel_utils.py

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

# -----------------------------
# Стилизация листов
# -----------------------------
def style_worksheet(ws) -> None:
    """Стилизация Excel-листа: жирные заголовки и автоширина колонок"""
    if ws.max_row < 1:
        return

    # Заголовки
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Автоширина
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2


# -----------------------------
# Преобразование списков в строки
# -----------------------------
def flatten_lists_in_df(df: pd.DataFrame) -> pd.DataFrame:
    """Преобразует списки/кортежи в строки перед записью в Excel, возвращает новый df"""
    df_copy = df.copy()

    def _cell_to_str(x):
        if isinstance(x, (list, tuple)):
            return ", ".join(map(str, x))
        if pd.isna(x):
            return ""
        return x

    for col in df_copy.columns:
        df_copy[col] = df_copy[col].apply(_cell_to_str)
    return df_copy


# -----------------------------
# Запись DataFrame в лист Excel
# -----------------------------
def write_df_to_sheet(wb: Workbook, sheet_name: str, df: pd.DataFrame) -> None:
    """Создаёт лист в wb и записывает DataFrame"""
    ws = wb.create_sheet(sheet_name)
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    style_worksheet(ws)


# -----------------------------
# Создание столбчатого графика
# -----------------------------
def add_bar_chart(ws, title: str, categories_col: int, data_cols: list[int], start_row: int = 2, start_col: int = 1) -> None:
    """
    Добавляет BarChart в лист ws.
    :param ws: openpyxl worksheet
    :param title: заголовок графика
    :param categories_col: индекс колонки с категориями (1-based)
    :param data_cols: список колонок с данными (1-based)
    :param start_row: строка, с которой начинается график
    :param start_col: колонка, куда вставить график
    """
    max_row = ws.max_row
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = "Количество"
    chart.x_axis.title = ws.cell(row=1, column=categories_col).value

    cats = Reference(ws, min_col=categories_col, min_row=2, max_row=max_row)
    for col in data_cols:
        data = Reference(ws, min_col=col, min_row=1, max_row=max_row)
        chart.add_data(data, titles_from_data=True)

    chart.set_categories(cats)
    chart.shape = 4
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True

    ws.add_chart(chart, f"{get_column_letter(start_col)}{start_row}")
