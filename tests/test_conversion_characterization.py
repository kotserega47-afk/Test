"""Characterization tests for analyzers/conversion.py (Phase 1 — no production changes)."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

import pandas as pd
import pytest
from openpyxl import load_workbook

from analyzers import conversion
from analyzers.conversion_analyzer import ConversionAnalyzer, normalize_name
from analyzers.conversion_dto import ConversionAnalysisResult, SpecialCardsState
from core.rules_v2.accessors import ConversionRulesAccessor
from core.rules_v2.models import MetaInfo, RulesSnapshotV2
from reporters.conversion_reporter import render_excel, render_telegram
from core.rules_v2.models import (
    ExclusionRule,
    JobParam,
    MetaInfo,
    PartnerDef,
    RulesSnapshotV2,
    ThresholdRule,
)
from core.rules_v2.normalizers import build_partner_key
from tests.fixtures.conversion.fixture_data import (
    COL_MAPPING,
    CONV_FILENAME,
    FIXED_NOW,
    RAW_PARTNER_ALPHA,
    RAW_PARTNER_BETA,
    write_card_fixture,
    write_conversion_fixture,
    write_special_cards_fixture,
)

MSK = ZoneInfo("Europe/Moscow")
_GOLDEN_DIR = Path(__file__).resolve().parent / "fixtures" / "conversion" / "golden"
_PK_ALPHA = build_partner_key(RAW_PARTNER_ALPHA)
_PK_BETA = build_partner_key(RAW_PARTNER_BETA)


def _load_golden(name: str) -> dict:
    with (_GOLDEN_DIR / name).open(encoding="utf-8") as f:
        return json.load(f)


def _build_snapshot(*, with_exclusion: bool = False) -> RulesSnapshotV2:
    partners = {
        _PK_ALPHA: PartnerDef(
            partner_key=_PK_ALPHA,
            partner_code="101",
            source_name=RAW_PARTNER_ALPHA,
            display_name=RAW_PARTNER_ALPHA,
            short_name=RAW_PARTNER_ALPHA,
            partner_type="legacy_partner",
            enabled=True,
        ),
        _PK_BETA: PartnerDef(
            partner_key=_PK_BETA,
            partner_code="202",
            source_name=RAW_PARTNER_BETA,
            display_name=RAW_PARTNER_BETA,
            short_name=RAW_PARTNER_BETA,
            partner_type="legacy_partner",
            enabled=True,
        ),
    }
    job_params = [
        JobParam(
            job_key="conversion",
            scope_type="job",
            scope_key="*",
            param_key="valid_status",
            value_type="str",
            value="Готов к работе",
            enabled=True,
        ),
        JobParam(
            job_key="conversion",
            scope_type="job",
            scope_key="*",
            param_key="valid_status",
            value_type="str",
            value="Активный вход",
            enabled=True,
        ),
    ]
    threshold_rules = [
        ThresholdRule(
            rule_key="conv_thr_alpha",
            job_key="conversion",
            scope_type="partner",
            scope_key=_PK_ALPHA,
            metric_key="error_streak",
            threshold_max=3.0,
            enabled=True,
        ),
    ]
    exclusion_rules: list[ExclusionRule] = []
    if with_exclusion:
        exclusion_rules.append(
            ExclusionRule(
                exclusion_key="conv_excl_beta",
                job_key="conversion",
                scope_type="partner",
                scope_key=_PK_BETA,
                start_dt=datetime(2026, 1, 10, 0, 0, 0, tzinfo=MSK),
                end_dt=datetime(2026, 1, 10, 23, 59, 59, tzinfo=MSK),
                reason="test exclusion",
                enabled=True,
            )
        )
    return RulesSnapshotV2(
        meta=MetaInfo(
            ruleset_version="test-conv-char",
            updated_at=FIXED_NOW,
            updated_by="pytest",
        ),
        partners=partners,
        job_params=job_params,
        threshold_rules=threshold_rules,
        exclusion_rules=exclusion_rules,
    )


@pytest.fixture
def conv_paths(tmp_path):
    conv_path = write_conversion_fixture(tmp_path / CONV_FILENAME)
    card_path = write_card_fixture(tmp_path / "card_test.xlsx")
    return conv_path, card_path


@pytest.fixture
def conversion_mocks(monkeypatch):
    """Patch external deps; enable Telegram via module-level CHAT_ID."""
    monkeypatch.setattr(conversion, "CHAT_ID", "999888777")
    monkeypatch.setattr(conversion, "now_msk", lambda: FIXED_NOW)

    messages: list[tuple] = []
    files: list[tuple] = []

    def _record_msg(text, chat_id=None, **_kwargs):
        messages.append((text, chat_id))

    def _record_file(path, caption=None, chat_id=None, **_kwargs):
        files.append((path, caption, chat_id))

    with patch.object(conversion, "get_snapshot_v2", return_value=_build_snapshot()), patch.object(
        conversion, "download_file", return_value=False
    ), patch.object(conversion, "send_message_sync", side_effect=_record_msg), patch.object(
        conversion, "send_file_sync", side_effect=_record_file
    ):
        yield {"messages": messages, "files": files}


def _run_conversion(conv_path: str, card_path: str, **kwargs):
    defaults = {
        "generate_excel": True,
        "send_telegram": True,
        "rules_force_sync": False,
    }
    defaults.update(kwargs)
    return conversion.run(
        conv_path,
        [card_path],
        COL_MAPPING,
        **defaults,
    )


def _build_analysis(conv_path: str, card_path: str, *, with_exclusion: bool = False):
    analyzer = ConversionAnalyzer()
    return analyzer.analyze(
        conv_file=conv_path,
        card_files=[card_path],
        col_mapping=COL_MAPPING,
        snapshot=_build_snapshot(with_exclusion=with_exclusion),
        special_state=SpecialCardsState.empty(),
    )


class TestConversionRulesAccessor:
    def test_get_valid_statuses_reads_job_params(self):
        rules = ConversionRulesAccessor.from_snapshot(_build_snapshot())
        assert rules.get_valid_statuses() == {"готов к работе", "активный вход"}

    def test_get_valid_statuses_fallback_when_empty(self):
        snapshot = RulesSnapshotV2(
            meta=MetaInfo(
                ruleset_version="empty",
                updated_at=FIXED_NOW,
                updated_by="pytest",
            )
        )
        rules = ConversionRulesAccessor.from_snapshot(snapshot)
        assert rules.get_valid_statuses() == {"готов к работе", "активный вход"}

    def test_get_error_streak_threshold_map_filters_conversion_metric(self):
        rules = ConversionRulesAccessor.from_snapshot(_build_snapshot())
        threshold_map = rules.get_error_streak_threshold_map(normalize_name)
        assert threshold_map == {"test partner": 3}

    def test_get_default_threshold_is_four(self):
        rules = ConversionRulesAccessor.from_snapshot(_build_snapshot())
        assert rules.get_default_threshold() == 4

    def test_get_partner_exclusion_rules_filters_by_job_and_scope(self):
        snapshot = _build_snapshot(with_exclusion=True)
        snapshot.exclusion_rules.append(
            ExclusionRule(
                exclusion_key="wallet_excl",
                job_key="wallet",
                scope_type="partner",
                scope_key=_PK_BETA,
                start_dt=datetime(2026, 1, 10, 0, 0, 0, tzinfo=MSK),
                end_dt=datetime(2026, 1, 10, 23, 59, 59, tzinfo=MSK),
                reason="other job",
                enabled=True,
            )
        )
        rules = ConversionRulesAccessor.from_snapshot(snapshot)
        exclusion_rules = rules.get_partner_exclusion_rules()
        assert len(exclusion_rules) == 1
        assert exclusion_rules[0].job_key == "conversion"

    def test_resolve_exclusion_partner_norm_matches_current_behavior(self):
        rules = ConversionRulesAccessor.from_snapshot(_build_snapshot(with_exclusion=True))
        rule = rules.get_partner_exclusion_rules()[0]
        assert rules.resolve_exclusion_partner_norm(rule, normalize_name) == "test beta"


class TestConversionReporter:
    def test_conversion_reporter_produces_same_excel_contract(self, conv_paths, tmp_path):
        conv_path, card_path = conv_paths
        analysis = _build_analysis(conv_path, card_path)
        contract = _load_golden("expected_excel_contract.json")

        wb = render_excel(analysis)
        report_path = tmp_path / "reporter_contract.xlsx"
        wb.save(report_path)

        loaded = load_workbook(report_path, read_only=True)
        try:
            assert loaded.sheetnames == contract["sheet_names"]
        finally:
            loaded.close()

        off = contract["sheets"]["Отключить"]
        df_off = pd.read_excel(report_path, sheet_name="Отключить")
        for col in off["required_columns"]:
            assert col in df_off.columns
        assert off["problem_card"] in df_off["card"].astype(str).tolist()

    def test_conversion_reporter_produces_same_telegram_fragments(self, conv_paths):
        conv_path, card_path = conv_paths
        analysis = _build_analysis(conv_path, card_path)
        telegram = render_telegram(analysis, conv_path)
        fragments = _load_golden("expected_telegram_fragments.json")

        problem_text = "\n".join(telegram.problem_messages)
        for frag in fragments["problem_message_fragments"]:
            assert frag in problem_text, f"missing problem fragment: {frag!r}"

        for frag in fragments["summary_message_fragments"]:
            assert frag in telegram.summary_message, f"missing summary fragment: {frag!r}"

    def test_render_excel_empty_sections_saves_with_placeholder_sheet(self, tmp_path):
        analysis = ConversionAnalysisResult(
            summary={
                "Карт в работе по партнёрам": {},
                "Карт в работе по пулам": {},
                "Max ошибки": 0,
                "Карты на отключение": 0,
            },
            problem_cards=pd.DataFrame(),
            cards_in_work_by_partner=pd.Series(dtype=int),
            cards_in_work_by_pool=pd.Series(dtype=int),
            special_cards_state=SpecialCardsState.empty(),
        )

        wb = render_excel(analysis)
        report_path = tmp_path / "empty_conversion_report.xlsx"
        wb.save(report_path)

        loaded = load_workbook(report_path, read_only=True)
        try:
            assert loaded.sheetnames == ["Нет данных"]
            assert loaded["Нет данных"]["A1"].value == "Нет данных для отчёта"
            assert loaded["Нет данных"]["A2"].value == "Все секции отчёта пустые"
        finally:
            loaded.close()


class TestOriginalPartner:
    def test_original_partner_column_present(self, conv_paths, conversion_mocks):
        conv_path, card_path = conv_paths
        result = _run_conversion(conv_path, card_path)
        contract = _load_golden("expected_excel_contract.json")

        df_off = pd.read_excel(result["report_path"], sheet_name="Отключить")
        for col in contract["sheets"]["Отключить"]["required_columns"]:
            assert col in df_off.columns

        cols = list(df_off.columns)
        assert "original_partner" in cols
        assert cols.index("original_partner") == cols.index("partner") + 1

    def test_original_partner_preserves_raw_value(self, conv_paths, conversion_mocks):
        conv_path, card_path = conv_paths
        result = _run_conversion(conv_path, card_path)
        problem = result["problem_cards"]

        assert problem.iloc[0]["partner"] == normalize_name(RAW_PARTNER_ALPHA)
        assert problem.iloc[0]["original_partner"] == RAW_PARTNER_ALPHA

    def test_original_partner_preserves_explicit_raw_label(self, tmp_path):
        raw_partner = "Partner A (123)"
        conv_path = tmp_path / "conv_raw_partner.xlsx"
        card_path = tmp_path / "card_raw_partner.xlsx"

        pd.DataFrame(
            [
                {
                    "Дата/Время создания": "10.01.2026 12:00:00",
                    "Карта": "CARDX",
                    "Партнёр": raw_partner,
                    "Статус": "Ошибка",
                },
                {
                    "Дата/Время создания": "10.01.2026 11:00:00",
                    "Карта": "CARDX",
                    "Партнёр": raw_partner,
                    "Статус": "Ошибка",
                },
                {
                    "Дата/Время создания": "10.01.2026 10:00:00",
                    "Карта": "CARDX",
                    "Партнёр": raw_partner,
                    "Статус": "Ошибка",
                },
            ]
        ).to_excel(conv_path, index=False)

        pd.DataFrame(
            [
                {
                    "Карта": "CARDX",
                    "Партнёр": raw_partner,
                    "Статус": "Готов к работе",
                    "Пул": "PoolA",
                },
                {
                    "Карта": "CARDY",
                    "Партнёр": raw_partner,
                    "Статус": "Активный вход",
                    "Пул": "PoolA",
                },
            ]
        ).to_excel(card_path, index=False)

        pk_raw = build_partner_key(raw_partner)
        snapshot = RulesSnapshotV2(
            meta=MetaInfo(
                ruleset_version="raw-partner-test",
                updated_at=FIXED_NOW,
                updated_by="pytest",
            ),
            partners={
                pk_raw: PartnerDef(
                    partner_key=pk_raw,
                    partner_code="123",
                    source_name=raw_partner,
                    display_name=raw_partner,
                    short_name=raw_partner,
                    partner_type="legacy_partner",
                    enabled=True,
                ),
            },
            job_params=_build_snapshot().job_params,
            threshold_rules=[
                ThresholdRule(
                    rule_key="conv_thr_raw",
                    job_key="conversion",
                    scope_type="partner",
                    scope_key=pk_raw,
                    metric_key="error_streak",
                    threshold_max=3.0,
                    enabled=True,
                ),
            ],
        )

        analysis = ConversionAnalyzer().analyze(
            conv_file=str(conv_path),
            card_files=[str(card_path)],
            col_mapping=COL_MAPPING,
            snapshot=snapshot,
            special_state=SpecialCardsState.empty(),
        )

        problem = analysis.problem_cards
        assert not problem.empty
        assert problem.iloc[0]["partner"] == "partner a"
        assert problem.iloc[0]["original_partner"] == raw_partner

    def test_thresholds_still_use_normalized_partner(self, conv_paths, conversion_mocks):
        conv_path, card_path = conv_paths
        result = _run_conversion(conv_path, card_path)
        problem = result["problem_cards"]
        rules = ConversionRulesAccessor.from_snapshot(_build_snapshot())
        threshold_map = rules.get_error_streak_threshold_map(normalize_name)

        row = problem.iloc[0]
        assert row["partner"] == normalize_name(RAW_PARTNER_ALPHA)
        assert row["original_partner"] == RAW_PARTNER_ALPHA
        assert row["partner"] in threshold_map
        assert row["original_partner"] not in threshold_map
        assert int(row["threshold"]) == threshold_map[row["partner"]]

    def test_exclusions_still_use_normalized_partner(self, conv_paths, monkeypatch):
        monkeypatch.setattr(conversion, "CHAT_ID", "")

        with patch.object(
            conversion, "get_snapshot_v2", return_value=_build_snapshot(with_exclusion=True)
        ), patch.object(conversion, "download_file", return_value=False), patch.object(
            conversion, "send_message_sync"
        ), patch.object(conversion, "send_file_sync"):
            conv_path, card_path = conv_paths
            result = _run_conversion(conv_path, card_path, send_telegram=False)

        problem = result["problem_cards"]
        rules = ConversionRulesAccessor.from_snapshot(_build_snapshot(with_exclusion=True))
        excluded_norm = rules.resolve_exclusion_partner_norm(
            rules.get_partner_exclusion_rules()[0],
            normalize_name,
        )

        assert excluded_norm == normalize_name(RAW_PARTNER_BETA)
        assert result["summary"]["Карты на отключение"] == 1
        assert problem.iloc[0]["partner"] == normalize_name(RAW_PARTNER_ALPHA)
        assert problem.iloc[0]["original_partner"] == RAW_PARTNER_ALPHA

    def test_telegram_contract_unchanged(self, conv_paths, conversion_mocks):
        conv_path, card_path = conv_paths
        mocks = conversion_mocks
        _run_conversion(conv_path, card_path)

        all_text = "\n".join(t[0] for t in mocks["messages"])
        fragments = _load_golden("expected_telegram_fragments.json")

        for frag in fragments["problem_message_fragments"]:
            assert frag in all_text, f"missing problem fragment: {frag!r}"

        summary_msgs = [t[0] for t in mocks["messages"] if "Анализ" in t[0] and "завершён" in t[0]]
        assert summary_msgs
        for frag in fragments["summary_message_fragments"]:
            assert frag in summary_msgs[0], f"missing summary fragment: {frag!r}"


class TestConversionFacadeDelegation:
    def test_conversion_run_delegates_to_analyzer_without_contract_drift(
        self, conv_paths, conversion_mocks
    ):
        conv_path, card_path = conv_paths
        with patch.object(conversion._analyzer, "analyze", wraps=conversion._analyzer.analyze) as analyze_mock:
            result = _run_conversion(conv_path, card_path)

        analyze_mock.assert_called_once()
        assert set(result.keys()) == {"summary", "workbook", "problem_cards", "report_path"}
        assert result["summary"] == _load_golden("expected_summary.json")


class TestConversionReturnContract:
    def test_run_return_dict_keys_and_types(self, conv_paths, conversion_mocks):
        conv_path, card_path = conv_paths
        result = _run_conversion(conv_path, card_path)

        assert set(result.keys()) == {"summary", "workbook", "problem_cards", "report_path"}
        assert isinstance(result["summary"], dict)
        assert result["workbook"] is not None
        assert isinstance(result["problem_cards"], pd.DataFrame)
        assert result["report_path"]
        assert Path(result["report_path"]).exists()
        assert result["report_path"].endswith("conversion_test_(15.01.2026).xlsx")

    def test_summary_matches_golden(self, conv_paths, conversion_mocks):
        conv_path, card_path = conv_paths
        result = _run_conversion(conv_path, card_path)
        expected = _load_golden("expected_summary.json")
        assert result["summary"] == expected

    def test_problem_cards_fixture_expectations(self, conv_paths, conversion_mocks):
        conv_path, card_path = conv_paths
        result = _run_conversion(conv_path, card_path)
        problem = result["problem_cards"]
        assert not problem.empty
        assert set(problem["card"]) == {"card001"}
        assert int(problem.iloc[0]["max_consecutive_errors"]) >= 3
        assert problem.iloc[0]["partner"] == "test partner"


class TestConversionExcelContract:
    def test_excel_sheets_and_columns(self, conv_paths, conversion_mocks):
        conv_path, card_path = conv_paths
        result = _run_conversion(conv_path, card_path)
        contract = _load_golden("expected_excel_contract.json")

        wb = load_workbook(result["report_path"], read_only=True)
        try:
            assert wb.sheetnames == contract["sheet_names"]

            off = contract["sheets"]["Отключить"]
            df_off = pd.read_excel(result["report_path"], sheet_name="Отключить")
            for col in off["required_columns"]:
                assert col in df_off.columns
            assert off["problem_card"] in df_off["card"].astype(str).tolist()
            assert df_off["max_consecutive_errors"].max() >= off["min_streak"]

            df_work = pd.read_excel(result["report_path"], sheet_name="Карт в работе")
            for col in contract["sheets"]["Карт в работе"]["required_columns"]:
                assert col in df_work.columns

            df_pool = pd.read_excel(result["report_path"], sheet_name="Карт в работе (Пулы)")
            for col in contract["sheets"]["Карт в работе (Пулы)"]["required_columns"]:
                assert col in df_pool.columns
        finally:
            wb.close()


class TestConversionTelegramContract:
    def test_telegram_messages_and_file(self, conv_paths, conversion_mocks):
        conv_path, card_path = conv_paths
        mocks = conversion_mocks
        _run_conversion(conv_path, card_path)

        assert mocks["messages"]
        assert mocks["files"]

        all_text = "\n".join(t[0] for t in mocks["messages"])
        fragments = _load_golden("expected_telegram_fragments.json")

        for frag in fragments["problem_message_fragments"]:
            assert frag in all_text, f"missing problem fragment: {frag!r}"

        summary_msgs = [t[0] for t in mocks["messages"] if "Анализ" in t[0] and "завершён" in t[0]]
        assert summary_msgs
        summary_text = summary_msgs[0]
        for frag in fragments["summary_message_fragments"]:
            assert frag in summary_text, f"missing summary fragment: {frag!r}"

        file_path, caption, chat_id = mocks["files"][0]
        assert Path(file_path).exists()
        assert "conversion_test.xlsx" in (caption or "")
        assert chat_id == "999888777"

    def test_special_cards_missing_notification(self, conv_paths, conversion_mocks):
        conv_path, card_path = conv_paths
        mocks = conversion_mocks
        _run_conversion(conv_path, card_path)

        all_text = "\n".join(t[0] for t in mocks["messages"])
        frag = _load_golden("expected_telegram_fragments.json")["special_cards_missing_fragment"]
        assert frag in all_text


class TestConversionNoTelegram:
    def test_missing_chat_id_fail_soft(self, conv_paths, monkeypatch):
        monkeypatch.setattr(conversion, "CHAT_ID", "")
        monkeypatch.setattr(conversion, "now_msk", lambda: FIXED_NOW)

        with patch.object(conversion, "get_snapshot_v2", return_value=_build_snapshot()), patch.object(
            conversion, "download_file", return_value=False
        ), patch.object(conversion, "send_message_sync") as msg_mock, patch.object(
            conversion, "send_file_sync"
        ) as file_mock:
            conv_path, card_path = conv_paths
            result = _run_conversion(conv_path, card_path)

        msg_mock.assert_not_called()
        file_mock.assert_not_called()
        assert result["workbook"] is not None
        assert result["report_path"]
        assert Path(result["report_path"]).exists()


class TestConversionSpecialCardsLoaded:
    def test_special_cards_download_success(self, conv_paths, monkeypatch, tmp_path):
        monkeypatch.setattr(conversion, "CHAT_ID", "")
        monkeypatch.setattr(conversion, "now_msk", lambda: FIXED_NOW)
        special_src = write_special_cards_fixture(tmp_path / "special_src.xlsx")

        def _fake_download(_src, dst):
            pd.read_excel(special_src).to_excel(dst, index=False)
            return True

        with patch.object(conversion, "get_snapshot_v2", return_value=_build_snapshot()), patch.object(
            conversion, "download_file", side_effect=_fake_download
        ), patch.object(conversion, "send_message_sync") as msg_mock, patch.object(
            conversion, "send_file_sync"
        ):
            conv_path, card_path = conv_paths
            result = _run_conversion(conv_path, card_path, send_telegram=False)

        assert result["summary"]["Карты на отключение"] == 1
        msg_mock.assert_not_called()


class TestConversionExclusionRules:
    def test_exclusion_removes_partner_rows(self, conv_paths, monkeypatch):
        monkeypatch.setattr(conversion, "CHAT_ID", "")
        monkeypatch.setattr(conversion, "now_msk", lambda: FIXED_NOW)

        with patch.object(
            conversion, "get_snapshot_v2", return_value=_build_snapshot(with_exclusion=True)
        ), patch.object(conversion, "download_file", return_value=False), patch.object(
            conversion, "send_message_sync"
        ), patch.object(
            conversion, "send_file_sync"
        ):
            conv_path, card_path = conv_paths
            result = _run_conversion(conv_path, card_path, send_telegram=False)

        assert result["summary"]["Карты на отключение"] == 1
        assert "card002" not in set(result["problem_cards"]["card"].astype(str))


class TestConversionHelpers:
    def test_normalize_name(self):
        assert conversion.normalize_name("  Test (101) ") == "test"
        assert conversion.normalize_name("Амобайл X") == "а-мобайл x"
        assert conversion.normalize_name("") == ""
        assert conversion.normalize_name(None) == ""

    def test_normalize_partners_list(self):
        out = conversion.normalize_partners_list("Foo (1), Bar (2)")
        assert out == ["foo", "bar"]

    def test_count_last_error_streak(self):
        df = pd.DataFrame(
            [
                {"card": "C1", "partner_norm": "p1", "datetime": pd.Timestamp("2026-01-10 12:00:00"), "status": "ошибка"},
                {"card": "C1", "partner_norm": "p1", "datetime": pd.Timestamp("2026-01-10 11:00:00"), "status": "ошибка"},
                {"card": "C1", "partner_norm": "p1", "datetime": pd.Timestamp("2026-01-10 10:00:00"), "status": "оплачен"},
                {"card": "C2", "partner_norm": "p2", "datetime": pd.Timestamp("2026-01-10 12:00:00"), "status": "оплачен"},
                {"card": "C2", "partner_norm": "p2", "datetime": pd.Timestamp("2026-01-10 11:00:00"), "status": "ошибка"},
            ]
        )
        out = conversion.count_last_error_streak(df)
        by_card = {row["card"]: int(row["max_consecutive_errors"]) for _, row in out.iterrows()}
        assert by_card["C1"] == 2
        assert by_card["C2"] == 0

    def test_load_data_mapping_and_datetime_drop(self, tmp_path):
        path = tmp_path / "mini.xlsx"
        pd.DataFrame(
            [
                {
                    "Карта": "X1",
                    "Партнёр": "P",
                    "Статус": "Ошибка",
                    "Дата/Время создания": "10.01.2026 10:00:00",
                },
                {
                    "Карта": "X2",
                    "Партнёр": "P",
                    "Статус": "Ошибка",
                    "Дата/Время создания": "bad-date",
                },
            ]
        ).to_excel(path, index=False)

        df = conversion.load_data(str(path), COL_MAPPING)
        assert set(df.columns) == {"card", "partner", "original_partner", "status", "datetime"}
        assert len(df) == 1
        assert df.iloc[0]["card"] == "x1"
        assert df.iloc[0]["partner"] == "p"
        assert df.iloc[0]["original_partner"] == "P"
        assert pd.notna(df.iloc[0]["datetime"])


class TestMainProcessFileConsumer:
    def test_process_file_uses_only_analyzer_result_summary(self, tmp_path, monkeypatch):
        """
        run_conversion_pipeline logs summary only for business output;
        report_path is read for observability payload; workbook/problem_cards are ignored.
        """
        conv_name = "conversion_test.xlsx"
        card_name = "card_test.xlsx"
        conv_local = tmp_path / conv_name
        card_local = tmp_path / card_name
        write_conversion_fixture(conv_local)
        write_card_fixture(card_local)

        monkeypatch.setenv("DROPBOX_INPUT_PATH", "/dropbox/in")
        monkeypatch.setenv("DROPBOX_PROCESSED_PATH", "/dropbox/out")
        monkeypatch.setattr("main.LOCAL_TMP_PATH", str(tmp_path))
        monkeypatch.setattr("integrations.conversion_pipeline.now_msk", lambda: FIXED_NOW)

        import main
        from integrations import conversion_pipeline

        tracked = {"summary": 0, "workbook": 0, "problem_cards": 0, "report_path": 0}

        class TrackingDict(dict):
            def get(self, key, default=None):
                if key in tracked:
                    tracked[key] += 1
                return super().get(key, default)

        def fake_run(*args, **kwargs):
            return TrackingDict(
                {
                    "summary": {"Карты на отключение": 1},
                    "workbook": object(),
                    "problem_cards": pd.DataFrame({"card": ["CARD001"]}),
                    "report_path": str(tmp_path / "report.xlsx"),
                }
            )

        fake_run.__module__ = "analyzers.conversion"

        with patch.object(conversion_pipeline, "download_file", return_value=True), patch.object(
            conversion_pipeline, "move_file", return_value=True
        ), patch.object(conversion_pipeline, "append_event"), patch.object(
            conversion_pipeline, "_safe_state_update"
        ), patch.object(conversion, "run", fake_run):
            ok = main.process_file(conv_name, aux_filename=card_name)

        assert ok is True
        assert tracked["summary"] == 1
        assert tracked["workbook"] == 0
        assert tracked["problem_cards"] == 0
        assert tracked["report_path"] == 1


class TestConversionCardsInWorkPairs:
    def test_analyze_completes_when_all_active_cards_are_problems(self, tmp_path):
        conv_path = write_conversion_fixture(tmp_path / CONV_FILENAME)
        card_df = pd.DataFrame(
            [
                {
                    "Карта": "CARD001",
                    "Партнёр": RAW_PARTNER_ALPHA,
                    "Статус": "Готов к работе",
                    "Пул": "PoolAlpha",
                },
            ]
        )
        card_path = tmp_path / "card_only_problem.xlsx"
        card_df.to_excel(card_path, index=False)

        analysis = ConversionAnalyzer().analyze(
            conv_file=conv_path,
            card_files=[str(card_path)],
            col_mapping=COL_MAPPING,
            snapshot=_build_snapshot(),
            special_state=SpecialCardsState.empty(),
        )

        assert analysis.summary["Карт в работе по партнёрам"] == {}
        assert analysis.cards_in_work_by_partner.empty
        assert analysis.summary["Карты на отключение"] == 1

        wb = render_excel(analysis)
        assert "Карт в работе" not in wb.sheetnames

    def test_cards_in_work_counts_comma_separated_partners(self, tmp_path):
        conv_path = tmp_path / "conv_multi_partner.xlsx"
        pd.DataFrame(
            [
                {
                    "Дата/Время создания": "10.01.2026 12:00:00",
                    "Карта": "CARD005",
                    "Партнёр": RAW_PARTNER_ALPHA,
                    "Статус": "Оплачен",
                },
            ]
        ).to_excel(conv_path, index=False)

        card_path = tmp_path / "card_multi_partner.xlsx"
        pd.DataFrame(
            [
                {
                    "Карта": "CARD005",
                    "Партнёр": f"{RAW_PARTNER_ALPHA}, {RAW_PARTNER_BETA}",
                    "Статус": "Готов к работе",
                    "Пул": "PoolMulti",
                },
            ]
        ).to_excel(card_path, index=False)

        analysis = ConversionAnalyzer().analyze(
            conv_file=str(conv_path),
            card_files=[str(card_path)],
            col_mapping=COL_MAPPING,
            snapshot=_build_snapshot(),
            special_state=SpecialCardsState.empty(),
        )

        work = analysis.summary["Карт в работе по партнёрам"]
        assert work["test partner (101)"] == 1
        assert work["test beta (202)"] == 1

    def test_cards_in_work_dedupes_duplicate_partner_card_rows(self, tmp_path):
        conv_path = tmp_path / "conv_dedup.xlsx"
        pd.DataFrame(
            [
                {
                    "Дата/Время создания": "10.01.2026 12:00:00",
                    "Карта": "CARD006",
                    "Партнёр": RAW_PARTNER_BETA,
                    "Статус": "Оплачен",
                },
            ]
        ).to_excel(conv_path, index=False)

        card_path = tmp_path / "card_dedup.xlsx"
        pd.DataFrame(
            [
                {
                    "Карта": "CARD006",
                    "Партнёр": RAW_PARTNER_BETA,
                    "Статус": "Готов к работе",
                    "Пул": "PoolBeta",
                },
                {
                    "Карта": "CARD006",
                    "Партнёр": RAW_PARTNER_BETA,
                    "Статус": "Готов к работе",
                    "Пул": "PoolBeta",
                },
            ]
        ).to_excel(card_path, index=False)

        analysis = ConversionAnalyzer().analyze(
            conv_file=str(conv_path),
            card_files=[str(card_path)],
            col_mapping=COL_MAPPING,
            snapshot=_build_snapshot(),
            special_state=SpecialCardsState.empty(),
        )

        assert analysis.summary["Карт в работе по партнёрам"] == {"test beta (202)": 1}
