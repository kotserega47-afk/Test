"""Wallet Editor Dropbox registry — lifecycle sheets and rules."""
from __future__ import annotations

import io
import json
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

import pandas as pd
import pytest

from automation.audit import Stats
from automation.runtime import WalletEditorTask
from integrations.wallet_editor_registry import (
    TIMEOUT_MESSAGE,
    append_run_to_dropbox_registry,
    resolve_source,
)
from integrations.wallet_editor_registry_settings import RegistrySettings
from integrations.wallet_editor_registry_xlsx import create_styled_registry_workbook, save_registry_workbook
from integrations.wallet_editor_auto_enable_eligibility import select_auto_enable_candidates
from integrations.wallet_editor_registry_lifecycle import (
    ALL_RESULTS_COLUMNS,
    DISABLE_DATE_COLUMN,
    HOLD_COLUMNS,
    HOLD_MARK,
    MISSING_OTLEZKA_DATE_TEXT,
    MISSING_OTLEZKA_STATUS,
    OPERATION_DATE_COLUMN,
    OPERATION_DATE_NUMBER_FORMAT,
    OTLEZKA_COLUMNS,
    RUNS_COLUMNS,
    SHEET_ALL_RESULTS,
    SHEET_HOLD,
    SHEET_OTLEZKA,
    SHEET_RUNS,
    STATUS_K_VKLUCHENIYU,
    STATUS_OZHIDAET,
    STATUS_PROSROCHENO,
    STATUS_VKLUCHENO,
    WARN_MESSAGE_TEMPLATE,
    load_warned_partners,
    partner_from_row,
    recalculate_all_results,
    result_row_dates,
    rows_from_result_excel,
    warned_partners_path,
)

MSK = ZoneInfo("Europe/Moscow")
TODAY = date(2026, 6, 3)
RUN_STARTED = datetime(2026, 6, 3, 9, 0, 0, tzinfo=MSK)
RUN_FINISHED = datetime(2026, 6, 3, 9, 5, 0, tzinfo=MSK)
DROPBOX_PATH = "/Ostin/platform/Tests/wallet_editor.xlsx"


def _make_task(
    *,
    profile: str = "DENIS",
    chat_id: int = -1001,
    run_id: str = "run-test-001",
) -> WalletEditorTask:
    return WalletEditorTask(
        file_path="/tmp/wallet_editor/in.xlsx",
        chat_id=chat_id,
        telegram_user_id=111,
        operator_profile=profile,
        source_file_name="batch.xlsx",
        login="login",
        password="pass",
        auth_state_path="/tmp/auth.json",
        run_id=run_id,
    )


def _write_result_xlsx(
    path: Path,
    *,
    rows: int = 1,
    partner: str = "Ostin",
    disable_at: str = "03.06.2026 09:00:00",
    status: str = "OK",
    action: str = "remove_partner",
) -> None:
    processed = datetime.strptime(disable_at, "%d.%m.%Y %H:%M:%S").replace(tzinfo=MSK)
    operation_date, disable_date = result_row_dates(action, status, processed)
    pd.DataFrame(
        {
            OPERATION_DATE_COLUMN: [operation_date] * rows,
            DISABLE_DATE_COLUMN: [disable_date] * rows,
            "card": [f"411111111111111{i}" for i in range(rows)],
            "action": [action] * rows,
            "value": [partner] * rows,
            "status": [status] * rows,
            "comment": ["removed"] * rows,
        }
    ).to_excel(path, index=False)


def _read_registry(data: bytes) -> dict[str, pd.DataFrame]:
    with pd.ExcelFile(io.BytesIO(data), engine="openpyxl") as book:
        out = {
            "all_results": pd.read_excel(book, SHEET_ALL_RESULTS),
            "runs": pd.read_excel(book, SHEET_RUNS),
        }
        if SHEET_HOLD in book.sheet_names:
            out["hold"] = pd.read_excel(book, SHEET_HOLD)
        else:
            out["hold"] = pd.DataFrame(columns=HOLD_COLUMNS)
        if SHEET_OTLEZKA in book.sheet_names:
            out[SHEET_OTLEZKA] = pd.read_excel(book, SHEET_OTLEZKA)
        else:
            out[SHEET_OTLEZKA] = pd.DataFrame(columns=OTLEZKA_COLUMNS)
        return out


@pytest.fixture
def registry_env(monkeypatch, tmp_path):
    store: dict[str, bytes] = {}
    revs: dict[str, str] = {}

    def fake_download_with_rev(dropbox_path: str, local_path: str) -> tuple[str, str | None]:
        content = store.get(dropbox_path)
        if content is None:
            return "not_found", None
        Path(local_path).write_bytes(content)
        return "ok", revs.get(dropbox_path, "rev-initial")

    def fake_get_rev(dropbox_path: str) -> str | None:
        if dropbox_path not in store:
            return None
        return revs.get(dropbox_path, "rev-initial")

    def fake_upload_if_rev(
        local_path: str, dropbox_path: str, expected_rev: str | None
    ) -> str:
        from integrations import dropbox_watcher

        if expected_rev is not None:
            current = dropbox_watcher.get_dropbox_file_rev(dropbox_path)
            if current != expected_rev:
                return "rev_conflict"
        store[dropbox_path] = Path(local_path).read_bytes()
        revs[dropbox_path] = f"rev-after-{len(store)}"
        return "uploaded"

    state_root = tmp_path / "state"
    monkeypatch.setenv("STATE_DIR", str(state_root))
    monkeypatch.setenv("DROPBOX_WALLET_EDITOR_PATH", DROPBOX_PATH)
    with patch(
        "integrations.wallet_editor_registry.download_file_with_rev",
        side_effect=fake_download_with_rev,
    ):
        with patch(
            "integrations.wallet_editor_registry.upload_file_if_rev",
            side_effect=fake_upload_if_rev,
        ):
            with patch(
                "integrations.dropbox_watcher.get_dropbox_file_rev",
                side_effect=fake_get_rev,
            ):
                with patch(
                    "integrations.wallet_editor_registry_lifecycle.now_msk",
                    return_value=datetime(2026, 6, 3, 12, 0, 0, tzinfo=MSK),
                ):
                    yield store, tmp_path, revs


def test_registry_all_results_columns_order(registry_env, tmp_path):
    store, _, _ = registry_env
    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path)
    append_run_to_dropbox_registry(
        _make_task(run_id="run-cols"),
        str(result_path),
        Stats(ok=1, fail=0, skip=0),
        run_started_at=RUN_STARTED,
        run_finished_at=RUN_FINISHED,
    )
    sheets = _read_registry(store[DROPBOX_PATH])
    assert list(sheets["all_results"].columns) == ALL_RESULTS_COLUMNS


def test_registry_runs_columns_order(registry_env, tmp_path):
    store, _, _ = registry_env
    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path)
    append_run_to_dropbox_registry(
        _make_task(run_id="run-runs-cols"),
        str(result_path),
        Stats(ok=1, fail=0, skip=0),
        run_started_at=RUN_STARTED,
        run_finished_at=RUN_FINISHED,
    )
    sheets = _read_registry(store[DROPBOX_PATH])
    assert list(sheets["runs"].columns) == RUNS_COLUMNS


def test_registry_creates_hold_sheet(registry_env, tmp_path):
    store, _, _ = registry_env
    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path)
    append_run_to_dropbox_registry(
        _make_task(run_id="run-hold-sheet"),
        str(result_path),
        Stats(ok=1, fail=0, skip=0),
        run_started_at=RUN_STARTED,
        run_finished_at=RUN_FINISHED,
    )
    sheets = _read_registry(store[DROPBOX_PATH])
    assert list(sheets["hold"].columns) == HOLD_COLUMNS


def test_registry_creates_otlezka_sheet(registry_env, tmp_path):
    store, _, _ = registry_env
    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path)
    append_run_to_dropbox_registry(
        _make_task(run_id="run-otlezka-sheet"),
        str(result_path),
        Stats(ok=1, fail=0, skip=0),
        run_started_at=RUN_STARTED,
        run_finished_at=RUN_FINISHED,
    )
    sheets = _read_registry(store[DROPBOX_PATH])
    assert list(sheets[SHEET_OTLEZKA].columns) == OTLEZKA_COLUMNS


def test_registry_partner_from_remove_partner_value():
    assert partner_from_row("remove_partner", "Ostin") == "Ostin"
    assert partner_from_row("add_partner", "Partner A") == "Partner A"
    assert partner_from_row("add_partner", "") == ""
    assert partner_from_row("set_status", "Ostin") == ""


def test_rows_from_result_excel_add_partner_populates_partner():
    result_df = pd.DataFrame(
        [
            {
                "Дата отключения": "01.06.2026 10:00:00",
                "card": "4111",
                "action": "add_partner",
                "value": "Partner A",
                "status": "OK",
                "comment": "added Partner A",
            }
        ]
    )
    rows = rows_from_result_excel(result_df)
    assert rows.iloc[0]["partner"] == "Partner A"
    assert rows.iloc[0]["action"] == "add_partner"
    assert rows.iloc[0]["status"] == "OK"
    assert rows.iloc[0]["comment"] == "added Partner A"


def test_recalculate_add_partner_does_not_set_lifecycle_dates():
    all_df = pd.DataFrame(
        [
            {
                "Дата отключения": "01.06.2026 10:00:00",
                "Дата включения": "",
                "Статус включения": "",
                "Включено": "",
                "Комментарий включения": "",
                "card": "4111",
                "partner": "Partner A",
                "action": "add_partner",
                "status": "OK",
                "comment": "added Partner A",
                "hold": "",
            }
        ]
    )
    otlezka = pd.DataFrame([{"partner": "Partner A", "Полные дни": 3, "comment": ""}])
    out, missing = recalculate_all_results(all_df, pd.DataFrame(), otlezka, today=TODAY)
    row = out.iloc[0]
    assert row["partner"] == "Partner A"
    assert row["Дата включения"] == ""
    assert row["Статус включения"] == ""
    assert row["hold"] == ""
    assert missing == set()

    eligible = select_auto_enable_candidates(out, include_overdue=True)
    assert eligible.selected == ()


def test_recalculate_add_partner_via_rows_from_result_preserves_partner():
    result_df = pd.DataFrame(
        [
            {
                "Дата отключения": "01.06.2026 10:00:00",
                "card": "4111",
                "action": "add_partner",
                "value": "Partner A",
                "status": "OK",
                "comment": "added Partner A",
            }
        ]
    )
    rows = rows_from_result_excel(result_df)
    otlezka = pd.DataFrame([{"partner": "Partner A", "Полные дни": 3, "comment": ""}])
    out, _ = recalculate_all_results(rows, pd.DataFrame(), otlezka, today=TODAY)
    assert out.iloc[0]["partner"] == "Partner A"


def test_registry_reenable_date_uses_partner_days():
    all_df = pd.DataFrame(
        [
            {
                "Дата отключения": "01.06.2026 10:00:00",
                "Дата включения": "",
                "Статус включения": "",
                "Включено": "",
                "Комментарий включения": "",
                "card": "4111",
                "partner": "",
                "action": "remove_partner",
                "value": "Ostin",
                "status": "OK",
                "comment": "",
                "hold": "",
            }
        ]
    )
    otlezka = pd.DataFrame([{"partner": "Ostin", "Полные дни": 3, "comment": ""}])
    out, missing = recalculate_all_results(all_df, pd.DataFrame(), otlezka, today=TODAY)
    assert out.iloc[0]["Дата включения"] == "04.06.2026"
    assert missing == set()


def test_registry_missing_otlezka_sets_text_and_status():
    all_df = pd.DataFrame(
        [
            {
                "Дата отключения": "03.06.2026 09:00:00",
                "Дата включения": "",
                "Статус включения": "",
                "Включено": "",
                "Комментарий включения": "",
                "card": "4111",
                "partner": "",
                "action": "remove_partner",
                "value": "UnknownPartner",
                "status": "OK",
                "comment": "",
                "hold": "",
            }
        ]
    )
    out, missing = recalculate_all_results(all_df, pd.DataFrame(), pd.DataFrame(), today=TODAY)
    assert out.iloc[0]["Дата включения"] == MISSING_OTLEZKA_DATE_TEXT
    assert out.iloc[0]["Статус включения"] == MISSING_OTLEZKA_STATUS
    assert "UnknownPartner" in missing


def test_registry_hold_blocks_reenable():
    all_df = pd.DataFrame(
        [
            {
                "Дата отключения": "01.06.2026 10:00:00",
                "Дата включения": "",
                "Статус включения": "",
                "Включено": "",
                "Комментарий включения": "",
                "card": "4111",
                "partner": "",
                "action": "remove_partner",
                "value": "Ostin",
                "status": "OK",
                "comment": "",
                "hold": "",
            }
        ]
    )
    hold = pd.DataFrame(
        [{"Дата добавления": "01.06.2026", "card": "4111", "partner": "Ostin", "comment": ""}]
    )
    otlezka = pd.DataFrame([{"partner": "Ostin", "Полные дни": 3, "comment": ""}])
    out, _ = recalculate_all_results(all_df, hold, otlezka, today=TODAY)
    assert out.iloc[0]["hold"] == HOLD_MARK
    assert out.iloc[0]["Дата включения"] == ""
    assert out.iloc[0]["Статус включения"] == HOLD_MARK


def test_registry_recalculates_existing_rows_after_otlezka_added(registry_env, tmp_path):
    store, _, _ = registry_env
    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path, partner="Ostin", disable_at="01.06.2026 10:00:00")

    append_run_to_dropbox_registry(
        _make_task(run_id="run-before-otlezka"),
        str(result_path),
        Stats(ok=1, fail=0, skip=0),
        run_started_at=RUN_STARTED,
        run_finished_at=RUN_FINISHED,
    )
    sheets = _read_registry(store[DROPBOX_PATH])
    assert sheets["all_results"].iloc[0]["Дата включения"] == MISSING_OTLEZKA_DATE_TEXT

    local = tmp_path / "with_otlezka.xlsx"
    with pd.ExcelWriter(local, engine="openpyxl") as writer:
        sheets["all_results"].to_excel(writer, sheet_name=SHEET_ALL_RESULTS, index=False)
        sheets["runs"].to_excel(writer, sheet_name=SHEET_RUNS, index=False)
        sheets["hold"].to_excel(writer, sheet_name=SHEET_HOLD, index=False)
        pd.DataFrame([{"partner": "Ostin", "Полные дни": 3, "comment": ""}]).to_excel(
            writer, sheet_name=SHEET_OTLEZKA, index=False
        )
    store[DROPBOX_PATH] = local.read_bytes()

    append_run_to_dropbox_registry(
        _make_task(run_id="run-after-otlezka"),
        str(result_path),
        Stats(ok=1, fail=0, skip=0),
        run_started_at=RUN_STARTED,
        run_finished_at=RUN_FINISHED,
    )
    sheets2 = _read_registry(store[DROPBOX_PATH])
    assert sheets2["all_results"].iloc[0]["Дата включения"] == "04.06.2026"
    assert len(sheets2["all_results"]) == 2


def test_registry_lifecycle_status_waiting_today_overdue():
    base = {
        "Дата отключения": "01.06.2026 10:00:00",
        "Дата включения": "",
        "Статус включения": "",
        "Включено": "",
        "Комментарий включения": "",
        "card": "1",
        "partner": "",
        "action": "remove_partner",
        "value": "Ostin",
        "status": "OK",
        "comment": "",
        "hold": "",
    }
    otlezka = pd.DataFrame([{"partner": "Ostin", "Полные дни": 5, "comment": ""}])

    waiting, _ = recalculate_all_results(
        pd.DataFrame([{**base}]), pd.DataFrame(), otlezka, today=date(2026, 6, 2)
    )
    assert waiting.iloc[0]["Статус включения"] == STATUS_OZHIDAET

    today_row, _ = recalculate_all_results(
        pd.DataFrame([{**base}]), pd.DataFrame(), otlezka, today=date(2026, 6, 6)
    )
    assert today_row.iloc[0]["Статус включения"] == STATUS_K_VKLUCHENIYU

    overdue, _ = recalculate_all_results(
        pd.DataFrame([{**base}]), pd.DataFrame(), otlezka, today=date(2026, 6, 10)
    )
    assert overdue.iloc[0]["Статус включения"] == STATUS_PROSROCHENO


def test_registry_included_status_overrides_lifecycle():
    all_df = pd.DataFrame(
        [
            {
                "Дата отключения": "01.06.2026 10:00:00",
                "Дата включения": "10.06.2026",
                "Статус включения": "",
                "Включено": "OK",
                "Комментарий включения": "done",
                "card": "1",
                "partner": "",
                "action": "remove_partner",
                "value": "Ostin",
                "status": "OK",
                "comment": "",
                "hold": "",
            }
        ]
    )
    otlezka = pd.DataFrame([{"partner": "Ostin", "Полные дни": 3, "comment": ""}])
    out, _ = recalculate_all_results(all_df, pd.DataFrame(), otlezka, today=date(2026, 6, 2))
    assert out.iloc[0]["Статус включения"] == STATUS_VKLUCHENO


def test_missing_otlezka_warning_sent_once_per_partner(registry_env, tmp_path):
    store, _, _ = registry_env
    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path, partner="NoOtlezka")
    messages: list[str] = []

    with patch(
        "integrations.wallet_editor_registry.send_message_sync",
        side_effect=lambda text, chat_id=None, **_kw: messages.append(text),
    ):
        append_run_to_dropbox_registry(
            _make_task(run_id="warn-1", chat_id=-555),
            str(result_path),
            Stats(ok=1, fail=0, skip=0),
            run_started_at=RUN_STARTED,
            run_finished_at=RUN_FINISHED,
        )
        append_run_to_dropbox_registry(
            _make_task(run_id="warn-2", chat_id=-555),
            str(result_path),
            Stats(ok=1, fail=0, skip=0),
            run_started_at=RUN_STARTED,
            run_finished_at=RUN_FINISHED,
        )

    warn_msgs = [m for m in messages if "Не настроена отлёжка" in m]
    assert len(warn_msgs) == 1
    assert "NoOtlezka" in warn_msgs[0]


def test_missing_otlezka_warning_state_cleared_after_fix(registry_env, tmp_path):
    store, _, _ = registry_env
    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path, partner="Teon")

    with patch("integrations.wallet_editor_registry.send_message_sync"):
        append_run_to_dropbox_registry(
            _make_task(run_id="teon-1"),
            str(result_path),
            Stats(ok=1, fail=0, skip=0),
            run_started_at=RUN_STARTED,
            run_finished_at=RUN_FINISHED,
        )
    assert "teon" in load_warned_partners()

    local = tmp_path / "add_teon.xlsx"
    sheets = _read_registry(store[DROPBOX_PATH])
    with pd.ExcelWriter(local, engine="openpyxl") as writer:
        sheets["all_results"].to_excel(writer, sheet_name=SHEET_ALL_RESULTS, index=False)
        sheets["runs"].to_excel(writer, sheet_name=SHEET_RUNS, index=False)
        sheets["hold"].to_excel(writer, sheet_name=SHEET_HOLD, index=False)
        pd.DataFrame([{"partner": "Teon", "Полные дни": 2, "comment": ""}]).to_excel(
            writer, sheet_name=SHEET_OTLEZKA, index=False
        )
    store[DROPBOX_PATH] = local.read_bytes()

    with patch("integrations.wallet_editor_registry.send_message_sync"):
        append_run_to_dropbox_registry(
            _make_task(run_id="teon-2"),
            str(result_path),
            Stats(ok=1, fail=0, skip=0),
            run_started_at=RUN_STARTED,
            run_finished_at=RUN_FINISHED,
        )

    warned = load_warned_partners()
    assert "teon" not in warned
    sheets2 = _read_registry(store[DROPBOX_PATH])
    assert sheets2["all_results"].iloc[0]["Дата включения"] == "05.06.2026"


def test_registry_best_effort_failure_still_does_not_break_worker(registry_env, tmp_path):
    import time

    import automation.worker as worker_mod

    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path, rows=1)

    with worker_mod._registry_lock:
        worker_mod._profile_workers.clear()

    with patch(
        "integrations.wallet_editor_registry.download_file_with_rev",
        return_value=("error", None),
    ):
        with patch("automation.worker.run") as mock_run:
            mock_run.return_value = (str(result_path), Stats(ok=1, fail=0, skip=0))
            with patch("automation.worker.send_text") as send_text:
                with patch("automation.worker.send_document") as send_document:
                    worker_mod.add_task(_make_task(run_id="run-worker-fail"))
                    deadline = time.time() + 2
                    while worker_mod._profile_workers["DENIS"].queue.unfinished_tasks > 0:
                        if time.time() > deadline:
                            break
                        time.sleep(0.02)

    send_text.assert_called()
    send_document.assert_called()


def test_registry_idempotent_same_run(registry_env, tmp_path):
    store, _, _ = registry_env
    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path, rows=2)
    task = _make_task(run_id="run-dup")

    append_run_to_dropbox_registry(
        task, str(result_path), Stats(ok=2, fail=0, skip=0),
        run_started_at=RUN_STARTED, run_finished_at=RUN_FINISHED,
    )
    append_run_to_dropbox_registry(
        task, str(result_path), Stats(ok=2, fail=0, skip=0),
        run_started_at=RUN_STARTED, run_finished_at=RUN_FINISHED,
    )

    sheets = _read_registry(store[DROPBOX_PATH])
    assert len(sheets["all_results"]) == 2
    assert len(sheets["runs"]) == 1


def test_registry_source_conversion_auto():
    assert resolve_source("CONVERSION_AUTO") == "conversion_auto"


def _append_once(registry_env, tmp_path, *, run_id: str = "fmt-run"):
    store, _, _ = registry_env
    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path)
    append_run_to_dropbox_registry(
        _make_task(run_id=run_id),
        str(result_path),
        Stats(ok=1, fail=0, skip=0),
        run_started_at=RUN_STARTED,
        run_finished_at=RUN_FINISHED,
    )
    return store


def test_registry_preserves_column_widths(registry_env, tmp_path):
    store, _, revs = registry_env
    styled = tmp_path / "styled.xlsx"
    create_styled_registry_workbook(styled)
    store[DROPBOX_PATH] = styled.read_bytes()
    revs[DROPBOX_PATH] = "rev-styled"

    _append_once(registry_env, tmp_path, run_id="preserve-width")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    ws = wb[SHEET_ALL_RESULTS]
    assert ws.column_dimensions["A"].width == 22.5
    assert ws.column_dimensions["G"].width == 18.0
    wb.close()


def test_registry_preserves_freeze_panes(registry_env, tmp_path):
    store, _, revs = registry_env
    styled = tmp_path / "styled.xlsx"
    create_styled_registry_workbook(styled)
    store[DROPBOX_PATH] = styled.read_bytes()
    revs[DROPBOX_PATH] = "rev-freeze"

    _append_once(registry_env, tmp_path, run_id="preserve-freeze")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    assert wb[SHEET_ALL_RESULTS].freeze_panes == "A2"
    wb.close()


def test_registry_preserves_header_fill(registry_env, tmp_path):
    store, _, revs = registry_env
    styled = tmp_path / "styled.xlsx"
    create_styled_registry_workbook(styled)
    store[DROPBOX_PATH] = styled.read_bytes()
    revs[DROPBOX_PATH] = "rev-fill"

    _append_once(registry_env, tmp_path, run_id="preserve-fill")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    cell = wb[SHEET_ALL_RESULTS].cell(row=1, column=1)
    assert cell.fill.start_color.rgb in ("004472C4", "4472C4")
    wb.close()


def test_registry_card_written_as_text(registry_env, tmp_path):
    store, _, revs = registry_env
    styled = tmp_path / "styled.xlsx"
    create_styled_registry_workbook(styled)
    store[DROPBOX_PATH] = styled.read_bytes()
    revs[DROPBOX_PATH] = "rev-card"

    result_path = tmp_path / "result.xlsx"
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    headers = ["Дата отключения", "card", "action", "value", "status", "comment"]
    for col_idx, name in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=name)
    ws.cell(row=2, column=1, value="03.06.2026 09:00:00")
    card_cell = ws.cell(row=2, column=2, value="0041111111111111")
    card_cell.number_format = "@"
    ws.cell(row=2, column=3, value="remove_partner")
    ws.cell(row=2, column=4, value="Ostin")
    ws.cell(row=2, column=5, value="OK")
    ws.cell(row=2, column=6, value="removed")
    wb.save(result_path)
    wb.close()

    append_run_to_dropbox_registry(
        _make_task(run_id="card-text"),
        str(result_path),
        Stats(ok=1, fail=0, skip=0),
        run_started_at=RUN_STARTED,
        run_finished_at=RUN_FINISHED,
    )

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    ws = wb[SHEET_ALL_RESULTS]
    headers = [c.value for c in ws[1]]
    card_col = headers.index("card") + 1
    cell = ws.cell(row=2, column=card_col)
    assert cell.number_format == "@"
    assert str(cell.value) == "0041111111111111"
    wb.close()


def test_registry_does_not_write_value_column_to_all_results(registry_env, tmp_path):
    store, _, _ = registry_env
    _append_once(registry_env, tmp_path, run_id="no-value-col")
    sheets = _read_registry(store[DROPBOX_PATH])
    assert "value" not in sheets["all_results"].columns


def test_registry_migrates_partner_from_value(registry_env, tmp_path):
    store, _, revs = registry_env
    legacy_path = tmp_path / "legacy.xlsx"
    legacy_cols = list(ALL_RESULTS_COLUMNS)
    legacy_cols.insert(legacy_cols.index("partner") + 1, "value")
    row = {col: "" for col in legacy_cols}
    row.update(
        {
            "Дата отключения": "01.06.2026 10:00:00",
            "card": "4111",
            "action": "remove_partner",
            "value": "LegacyPartner",
            "status": "OK",
        }
    )
    with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
        pd.DataFrame([row], columns=legacy_cols).to_excel(
            writer, sheet_name=SHEET_ALL_RESULTS, index=False
        )
        pd.DataFrame(columns=RUNS_COLUMNS).to_excel(writer, sheet_name=SHEET_RUNS, index=False)
    store[DROPBOX_PATH] = legacy_path.read_bytes()
    revs[DROPBOX_PATH] = "rev-legacy"

    _append_once(registry_env, tmp_path, run_id="migrate-value")

    sheets = _read_registry(store[DROPBOX_PATH])
    assert "value" not in sheets["all_results"].columns
    assert sheets["all_results"].iloc[0]["partner"] == "LegacyPartner"


def test_legacy_workbook_inserts_operation_date_column(registry_env, tmp_path):
    store, _, revs = registry_env
    legacy_path = tmp_path / "legacy_no_op_date.xlsx"
    legacy_cols = [col for col in ALL_RESULTS_COLUMNS if col != OPERATION_DATE_COLUMN]
    row = {col: "" for col in legacy_cols}
    row.update(
        {
            DISABLE_DATE_COLUMN: "01.06.2026 10:00:00",
            "card": "4111",
            "action": "remove_partner",
            "partner": "Ostin",
            "status": "OK",
        }
    )
    with pd.ExcelWriter(legacy_path, engine="openpyxl") as writer:
        pd.DataFrame([row], columns=legacy_cols).to_excel(
            writer, sheet_name=SHEET_ALL_RESULTS, index=False
        )
        pd.DataFrame(columns=RUNS_COLUMNS).to_excel(writer, sheet_name=SHEET_RUNS, index=False)
    store[DROPBOX_PATH] = legacy_path.read_bytes()
    revs[DROPBOX_PATH] = "rev-no-op-date"

    _append_once(registry_env, tmp_path, run_id="legacy-op-date")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    ws = wb[SHEET_ALL_RESULTS]
    headers = [ws.cell(row=1, column=col_idx).value for col_idx in range(1, len(ALL_RESULTS_COLUMNS) + 1)]
    assert headers[0] == OPERATION_DATE_COLUMN
    assert headers[1] == DISABLE_DATE_COLUMN
    assert str(ws.cell(row=2, column=2).value).startswith("01.06.2026")
    sheets = _read_registry(store[DROPBOX_PATH])
    assert list(sheets["all_results"].columns) == ALL_RESULTS_COLUMNS
    wb.close()


def test_registry_does_not_overwrite_hold_sheet(registry_env, tmp_path):
    store, _, revs = registry_env
    wb_path = tmp_path / "with_hold.xlsx"
    with pd.ExcelWriter(wb_path, engine="openpyxl") as writer:
        pd.DataFrame(columns=ALL_RESULTS_COLUMNS).to_excel(
            writer, sheet_name=SHEET_ALL_RESULTS, index=False
        )
        pd.DataFrame(columns=RUNS_COLUMNS).to_excel(writer, sheet_name=SHEET_RUNS, index=False)
        pd.DataFrame(
            [{"Дата добавления": "01.01.2020", "card": "USERCARD", "partner": "UserP", "comment": "keep"}]
        ).to_excel(writer, sheet_name=SHEET_HOLD, index=False)
    store[DROPBOX_PATH] = wb_path.read_bytes()
    revs[DROPBOX_PATH] = "rev-hold"

    _append_once(registry_env, tmp_path, run_id="hold-keep")

    sheets = _read_registry(store[DROPBOX_PATH])
    assert sheets["hold"].iloc[0]["card"] == "USERCARD"
    assert sheets["hold"].iloc[0]["comment"] == "keep"


def test_registry_does_not_overwrite_otlezka_sheet(registry_env, tmp_path):
    store, _, revs = registry_env
    wb_path = tmp_path / "with_otlezka.xlsx"
    with pd.ExcelWriter(wb_path, engine="openpyxl") as writer:
        pd.DataFrame(columns=ALL_RESULTS_COLUMNS).to_excel(
            writer, sheet_name=SHEET_ALL_RESULTS, index=False
        )
        pd.DataFrame(columns=RUNS_COLUMNS).to_excel(writer, sheet_name=SHEET_RUNS, index=False)
        pd.DataFrame(
            [{"partner": "UserPartner", "Полные дни": 99, "comment": "manual"}]
        ).to_excel(writer, sheet_name=SHEET_OTLEZKA, index=False)
    store[DROPBOX_PATH] = wb_path.read_bytes()
    revs[DROPBOX_PATH] = "rev-otlezka"

    _append_once(registry_env, tmp_path, run_id="otlezka-keep")

    sheets = _read_registry(store[DROPBOX_PATH])
    assert sheets[SHEET_OTLEZKA].iloc[0]["partner"] == "UserPartner"
    assert int(sheets[SHEET_OTLEZKA].iloc[0]["Полные дни"]) == 99
    assert sheets[SHEET_OTLEZKA].iloc[0]["comment"] == "manual"


def test_registry_rev_conflict_skips_upload(registry_env, tmp_path):
    store, _, revs = registry_env
    styled = tmp_path / "styled.xlsx"
    create_styled_registry_workbook(styled)
    store[DROPBOX_PATH] = styled.read_bytes()
    before = store[DROPBOX_PATH]
    revs[DROPBOX_PATH] = "rev-at-download"
    fast = RegistrySettings(
        registry_warning_seconds=30,
        registry_timeout_seconds=2,
        registry_retry_interval_seconds=1,
    )

    with patch(
        "integrations.dropbox_watcher.get_dropbox_file_rev",
        return_value="rev-changed-by-user",
    ):
        with patch(
            "integrations.wallet_editor_registry.load_registry_settings",
            return_value=fast,
        ):
            _append_once(registry_env, tmp_path, run_id="rev-skip")

    assert store[DROPBOX_PATH] == before


def test_registry_rev_conflict_warns_telegram(registry_env, tmp_path):
    store, _, revs = registry_env
    styled = tmp_path / "styled.xlsx"
    create_styled_registry_workbook(styled)
    store[DROPBOX_PATH] = styled.read_bytes()
    revs[DROPBOX_PATH] = "rev-at-download"
    result_path = tmp_path / "result.xlsx"
    _write_result_xlsx(result_path)
    messages: list[str] = []

    fast = RegistrySettings(
        registry_warning_seconds=30,
        registry_timeout_seconds=2,
        registry_retry_interval_seconds=1,
    )

    with patch(
        "integrations.dropbox_watcher.get_dropbox_file_rev",
        return_value="rev-changed-by-user",
    ):
        with patch(
            "integrations.wallet_editor_registry.send_message_sync",
            side_effect=lambda text, chat_id=None, **_kw: messages.append(text),
        ):
            append_run_to_dropbox_registry(
                _make_task(run_id="rev-warn", chat_id=-999),
                str(result_path),
                Stats(ok=1, fail=0, skip=0),
                run_started_at=RUN_STARTED,
                run_finished_at=RUN_FINISHED,
                settings=fast,
            )

    assert any(TIMEOUT_MESSAGE in m for m in messages)


def test_registry_open_without_rev_change_allows_upload(registry_env, tmp_path):
    store, _, revs = registry_env
    styled = tmp_path / "styled.xlsx"
    create_styled_registry_workbook(styled)
    store[DROPBOX_PATH] = styled.read_bytes()
    revs[DROPBOX_PATH] = "rev-stable"

    _append_once(registry_env, tmp_path, run_id="rev-ok")

    sheets = _read_registry(store[DROPBOX_PATH])
    assert len(sheets["all_results"]) == 1


def _style_all_results_data_row(ws, row_idx: int) -> None:
    from openpyxl.styles import Alignment, Border, Side

    card_col = ALL_RESULTS_COLUMNS.index("card") + 1
    status_col = ALL_RESULTS_COLUMNS.index("status") + 1
    comment_col = ALL_RESULTS_COLUMNS.index("comment") + 1
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(ALL_RESULTS_COLUMNS) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.border = border
        if col_idx in {card_col, status_col, comment_col}:
            cell.alignment = center
    ws.cell(row=row_idx, column=card_col).number_format = "@"


def _style_runs_data_row(ws, row_idx: int) -> None:
    from openpyxl.styles import Border, Side

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col_idx in range(1, len(RUNS_COLUMNS) + 1):
        ws.cell(row=row_idx, column=col_idx).border = border


def _store_workbook_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def test_registry_append_copies_all_results_row_style(registry_env, tmp_path):
    store, _, revs = registry_env
    styled = tmp_path / "styled.xlsx"
    create_styled_registry_workbook(styled)
    store[DROPBOX_PATH] = styled.read_bytes()
    revs[DROPBOX_PATH] = "rev-style-append"

    _append_once(registry_env, tmp_path, run_id="style-seed-row")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    ws = wb[SHEET_ALL_RESULTS]
    _style_all_results_data_row(ws, 2)
    store[DROPBOX_PATH] = _store_workbook_bytes(wb)

    _append_once(registry_env, tmp_path, run_id="style-second-row")

    wb2 = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    ws2 = wb2[SHEET_ALL_RESULTS]
    card_col = ALL_RESULTS_COLUMNS.index("card") + 1
    status_col = ALL_RESULTS_COLUMNS.index("status") + 1
    new_row = 3

    assert ws2.cell(row=new_row, column=card_col).border.left.style == "thin"
    assert ws2.cell(row=new_row, column=status_col).alignment.horizontal == "center"
    assert ws2.cell(row=new_row, column=card_col).number_format == "@"
    assert ws2.column_dimensions["A"].width == 22.5
    assert ws2.column_dimensions["G"].width == 18.0
    assert ws2.freeze_panes == "A2"
    wb2.close()


def test_registry_append_copies_runs_row_style(registry_env, tmp_path):
    store, _, revs = registry_env
    styled = tmp_path / "styled.xlsx"
    create_styled_registry_workbook(styled)
    store[DROPBOX_PATH] = styled.read_bytes()
    revs[DROPBOX_PATH] = "rev-runs-style"

    _append_once(registry_env, tmp_path, run_id="runs-style-seed")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    _style_runs_data_row(wb[SHEET_RUNS], 2)
    store[DROPBOX_PATH] = _store_workbook_bytes(wb)

    _append_once(registry_env, tmp_path, run_id="runs-style-second")

    wb2 = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    ws_runs = wb2[SHEET_RUNS]
    assert ws_runs.cell(row=3, column=1).border.left.style == "thin"
    assert ws_runs.cell(row=3, column=len(RUNS_COLUMNS)).border.left.style == "thin"
    wb2.close()


def _operation_date_cell(ws, row_idx: int = 2):
    col = ALL_RESULTS_COLUMNS.index(OPERATION_DATE_COLUMN) + 1
    return ws.cell(row=row_idx, column=col)


def _disable_date_cell(ws, row_idx: int = 2):
    col = ALL_RESULTS_COLUMNS.index(DISABLE_DATE_COLUMN) + 1
    return ws.cell(row=row_idx, column=col)


def _all_results_row(**overrides) -> dict:
    row = {col: "" for col in ALL_RESULTS_COLUMNS}
    row.update(overrides)
    return row


def test_operation_date_written_as_excel_date(registry_env, tmp_path):
    store, _, _ = registry_env
    _append_once(registry_env, tmp_path, run_id="op-date-excel")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    ws = wb[SHEET_ALL_RESULTS]
    cell = _operation_date_cell(ws)
    assert isinstance(cell.value, (date, datetime))
    assert not isinstance(cell.value, str)
    assert cell.number_format == OPERATION_DATE_NUMBER_FORMAT
    wb.close()


def test_operation_date_preserved_on_append(registry_env, tmp_path):
    store, _, revs = registry_env
    _append_once(registry_env, tmp_path, run_id="op-date-first")
    _append_once(registry_env, tmp_path, run_id="op-date-second")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    ws = wb[SHEET_ALL_RESULTS]
    for row_idx in (2, 3):
        cell = _operation_date_cell(ws, row_idx)
        assert isinstance(cell.value, (date, datetime))
        assert cell.number_format == OPERATION_DATE_NUMBER_FORMAT
    wb.close()


@pytest.mark.parametrize(
    "legacy_value",
    ["08.06.2026", "08.06.2026 00:00:00"],
)
def test_legacy_string_operation_date_normalized_on_save(tmp_path, legacy_value: str):
    from openpyxl import Workbook

    wb_path = tmp_path / "legacy_op_date.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = SHEET_ALL_RESULTS
    for col_idx, name in enumerate(ALL_RESULTS_COLUMNS, 1):
        ws.cell(row=1, column=col_idx, value=name)
    ws.cell(row=2, column=1, value=legacy_value)
    ws.cell(row=2, column=ALL_RESULTS_COLUMNS.index(DISABLE_DATE_COLUMN) + 1, value="01.06.2026 10:00:00")
    ws.cell(row=2, column=ALL_RESULTS_COLUMNS.index("card") + 1, value="4111")
    ws.cell(row=2, column=ALL_RESULTS_COLUMNS.index("action") + 1, value="remove_partner")
    ws.cell(row=2, column=ALL_RESULTS_COLUMNS.index("partner") + 1, value="Ostin")
    ws.cell(row=2, column=ALL_RESULTS_COLUMNS.index("status") + 1, value="OK")
    wb.create_sheet(SHEET_RUNS)
    wb.save(wb_path)
    wb.close()

    all_df = pd.DataFrame(
        [
            _all_results_row(
                **{
                    OPERATION_DATE_COLUMN: legacy_value,
                    DISABLE_DATE_COLUMN: "01.06.2026 10:00:00",
                    "card": "4111",
                    "partner": "Ostin",
                    "action": "remove_partner",
                    "status": "OK",
                }
            )
        ]
    )
    save_registry_workbook(
        wb_path,
        all_results=all_df,
        runs=pd.DataFrame(columns=RUNS_COLUMNS),
        hold_exists=True,
        otlezka_exists=True,
        is_new_file=False,
    )

    from openpyxl import load_workbook

    wb2 = load_workbook(wb_path)
    ws2 = wb2[SHEET_ALL_RESULTS]
    cell = _operation_date_cell(ws2)
    assert isinstance(cell.value, (date, datetime))
    assert not isinstance(cell.value, str)
    assert cell.number_format == OPERATION_DATE_NUMBER_FORMAT
    if isinstance(cell.value, datetime):
        assert cell.value.time() == datetime.min.time()
    wb2.close()


def test_remove_partner_ok_disable_date_remains_timestamp(registry_env, tmp_path):
    store, _, _ = registry_env
    _append_once(registry_env, tmp_path, run_id="disable-ts")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    ws = wb[SHEET_ALL_RESULTS]
    cell = _disable_date_cell(ws)
    assert cell.value is not None
    rendered = (
        cell.value.strftime("%d.%m.%Y %H:%M:%S")
        if isinstance(cell.value, datetime)
        else str(cell.value)
    )
    assert "09:00:00" in rendered
    assert cell.number_format != OPERATION_DATE_NUMBER_FORMAT
    wb.close()


def test_add_partner_result_has_empty_disable_date_in_registry(registry_env, tmp_path):
    store, _, _ = registry_env
    result_path = tmp_path / "add_partner.xlsx"
    processed = datetime(2026, 6, 8, 12, 0, 0, tzinfo=MSK)
    operation_date, disable_date = result_row_dates("add_partner", "OK", processed)
    assert disable_date == ""
    pd.DataFrame(
        {
            OPERATION_DATE_COLUMN: [operation_date],
            DISABLE_DATE_COLUMN: [disable_date],
            "card": ["4111111111111111"],
            "action": ["add_partner"],
            "value": ["Ostin"],
            "status": ["OK"],
            "comment": ["added"],
        }
    ).to_excel(result_path, index=False)

    append_run_to_dropbox_registry(
        _make_task(run_id="add-no-disable"),
        str(result_path),
        Stats(ok=1, fail=0, skip=0),
        run_started_at=RUN_STARTED,
        run_finished_at=RUN_FINISHED,
    )

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(store[DROPBOX_PATH]))
    ws = wb[SHEET_ALL_RESULTS]
    row_idx = ws.max_row
    op_cell = _operation_date_cell(ws, row_idx)
    disable_cell = _disable_date_cell(ws, row_idx)
    assert isinstance(op_cell.value, (date, datetime))
    assert op_cell.number_format == OPERATION_DATE_NUMBER_FORMAT
    assert disable_cell.value in (None, "")
    wb.close()
