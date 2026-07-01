"""Unit tests for RegistryExportBuilder (Phase 4)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

import pandas as pd
import pytest
from openpyxl import load_workbook

from integrations.wallet_editor_registry_db.manual_snapshot import HoldSnapshotRow, OtlezkaSnapshotRow
from integrations.wallet_editor_registry_db.manual_store import InMemoryManualSyncStore
from integrations.wallet_editor_registry_db.manual_sync_state import (
    META_LAST_SNAPSHOT_HASH,
    META_LAST_SYNC_AT,
    META_LAST_SYNC_STATUS,
)
from integrations.wallet_editor_registry_db.config import LegacyRegistrySourceError
from integrations.wallet_editor_registry_db.registry_export_builder import (
    RegistryExportBuilder,
    format_registry_export_summary,
)
from integrations.wallet_editor_registry_lifecycle import (
    ALL_RESULTS_COLUMNS,
    HOLD_COLUMNS,
    OTLEZKA_COLUMNS,
    RUNS_COLUMNS,
    SHEET_ALL_RESULTS,
    SHEET_HOLD,
    SHEET_OTLEZKA,
    SHEET_RUNS,
)
from integrations.wallet_editor_registry_xlsx import SHEET_README, SHEET_SYNC_STATUS

MSK = ZoneInfo("Europe/Moscow")


def _empty_frames():
    return (
        pd.DataFrame(columns=ALL_RESULTS_COLUMNS),
        pd.DataFrame(columns=RUNS_COLUMNS),
    )


def _seed_store() -> InMemoryManualSyncStore:
    store = InMemoryManualSyncStore()
    store.set_meta(META_LAST_SYNC_STATUS, "success")
    store.set_meta(META_LAST_SNAPSHOT_HASH, "abc123def456789")
    store.set_meta(META_LAST_SYNC_AT, "2026-07-01T12:00:00+03:00")
    store.hold[("4111", "ostin")] = {
        "active": True,
        "row": HoldSnapshotRow(
            card="4111",
            partner="Ostin",
            card_norm="4111",
            partner_norm="ostin",
            added_at="01.07.2026",
            comment="",
            source_row_index=1,
        ),
    }
    store.otlezka["ostin"] = {
        "active": True,
        "row": OtlezkaSnapshotRow(
            partner="Ostin",
            partner_norm="ostin",
            full_days=30,
            comment="",
            source_row_index=1,
        ),
    }
    return store


class TestRegistryExportBuilder:
    def test_builds_workbook_from_pg_only(self, tmp_path, monkeypatch):
        monkeypatch.setenv("WALLET_EDITOR_REGISTRY_SOURCE", "postgres")
        store = _seed_store()

        with (
            patch(
                "integrations.wallet_editor_registry_db.registry_export_builder.load_registry_frames_from_postgres",
                return_value=_empty_frames(),
            ),
            patch(
                "integrations.dropbox_watcher.download_file_with_rev",
                side_effect=AssertionError("Dropbox download must not run"),
            ),
            patch(
                "integrations.dropbox_watcher.upload_file_if_rev",
                side_effect=AssertionError("Dropbox upload must not run"),
            ),
        ):
            artifact = RegistryExportBuilder(store=store).build(
                output_path=tmp_path,
                generated_at=datetime(2026, 7, 1, 15, 30, 0, tzinfo=MSK),
            )

        assert artifact.path.is_file()
        assert artifact.filename == "wallet_editor_export_20260701_153000_MSK.xlsx"

        wb = load_workbook(artifact.path, read_only=True)
        assert wb.sheetnames == [
            SHEET_ALL_RESULTS,
            SHEET_RUNS,
            SHEET_HOLD,
            SHEET_OTLEZKA,
            SHEET_README,
            SHEET_SYNC_STATUS,
        ]
        readme = [wb[SHEET_README].cell(row=i, column=1).value for i in range(1, 20)]
        readme_text = "\n".join(line for line in readme if line)
        assert "WalletEditor Registry Export" in readme_text
        assert "Source of truth:" in readme_text
        assert "PostgreSQL" in readme_text
        assert "Do not edit this workbook." in readme_text
        assert "operator workbook" in readme_text
        assert "/registry_export" in readme_text
        assert "Manual snapshot hash: abc123def456" in readme_text
        wb.close()

    def test_hold_sheet_uses_hold_columns(self, tmp_path, monkeypatch):
        monkeypatch.setenv("WALLET_EDITOR_REGISTRY_SOURCE", "postgres")
        store = _seed_store()

        with patch(
            "integrations.wallet_editor_registry_db.registry_export_builder.load_registry_frames_from_postgres",
            return_value=_empty_frames(),
        ):
            artifact = RegistryExportBuilder(store=store).build(output_path=tmp_path)

        wb = load_workbook(artifact.path, read_only=True)
        hold_ws = wb[SHEET_HOLD]
        assert hold_ws.cell(row=1, column=1).value == HOLD_COLUMNS[0]
        assert hold_ws.cell(row=2, column=2).value == "4111"
        wb.close()

    def test_missing_manual_sync_degraded_not_blocked(self, tmp_path, monkeypatch):
        monkeypatch.setenv("WALLET_EDITOR_REGISTRY_SOURCE", "postgres")
        store = InMemoryManualSyncStore()

        with patch(
            "integrations.wallet_editor_registry_db.registry_export_builder.load_registry_frames_from_postgres",
            return_value=_empty_frames(),
        ):
            artifact = RegistryExportBuilder(store=store).build(output_path=tmp_path)

        assert artifact.summary.manual_sync_degraded is True
        text = format_registry_export_summary(artifact.summary)
        assert "DEGRADED: no successful manual sync" in text

    def test_db_read_failure_raises(self, monkeypatch):
        monkeypatch.setenv("WALLET_EDITOR_REGISTRY_SOURCE", "postgres")

        with patch(
            "integrations.wallet_editor_registry_db.registry_export_builder.load_registry_frames_from_postgres",
            side_effect=RuntimeError("db down"),
        ):
            with pytest.raises(RuntimeError, match="db down"):
                RegistryExportBuilder(store=_seed_store()).build()

    def test_excel_registry_source_rejected(self, monkeypatch):
        monkeypatch.setenv("WALLET_EDITOR_REGISTRY_SOURCE", "excel")

        with pytest.raises(LegacyRegistrySourceError, match="no longer supported"):
            RegistryExportBuilder(store=_seed_store()).load()

    def test_summary_includes_row_counts(self, tmp_path, monkeypatch):
        monkeypatch.setenv("WALLET_EDITOR_REGISTRY_SOURCE", "postgres")
        store = _seed_store()
        all_df = pd.DataFrame([{col: "" for col in ALL_RESULTS_COLUMNS}])
        runs_df = pd.DataFrame([{col: "" for col in RUNS_COLUMNS}])

        with patch(
            "integrations.wallet_editor_registry_db.registry_export_builder.load_registry_frames_from_postgres",
            return_value=(all_df, runs_df),
        ):
            artifact = RegistryExportBuilder(store=store).build(output_path=tmp_path)

        text = format_registry_export_summary(artifact.summary)
        assert "all_results rows: 1" in text
        assert "runs rows: 1" in text
        assert "active hold rows: 1" in text
        assert "active Отлёжка rows: 1" in text
        assert "Read-only export" in text
        assert "/registry_export" in text
