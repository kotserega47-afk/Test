"""HOLD enforcement for add_partner in manual WalletEditor and auto-enable."""

from __future__ import annotations

from unittest.mock import MagicMock, patch

import pandas as pd
import pytest

from automation.audit import Stats
from automation.engine import (
    _apply_add_partner_hold_precheck,
    _group_actions,
    run,
)
from automation.runtime import RunConfig
from integrations.wallet_editor_auto_enable_eligibility import CandidateRow
from integrations.wallet_editor_auto_enable_executor import (
    ERROR_HOLD,
    ERROR_HOLD_CHECK_FAILED,
    REGISTRY_FAIL,
    REGISTRY_OK,
    REGISTRY_SKIP,
    build_batch_execution_report,
    process_enable_candidate,
)
from integrations.wallet_editor_hold import (
    HOLD_CHECK_FAILED_AUTO_COMMENT,
    HOLD_CHECK_FAILED_MANUAL_COMMENT,
    HOLD_SKIP_COMMENT,
    HoldPairsSnapshot,
    is_card_partner_on_hold,
    load_hold_pairs_from_dropbox,
    normalize_hold_card,
    normalize_hold_partner,
)
from integrations.wallet_editor_registry import EnableRegistryUpdate, apply_enable_updates_to_all_results
from integrations.wallet_editor_registry_lifecycle import (
    ACTION_REMOVE_PARTNER,
    ALL_RESULTS_COLUMNS,
    DISABLE_DATE_COLUMN,
    OPERATION_DATE_COLUMN,
    recalculate_all_results,
)
from integrations.wallet_editor_auto_enable_settings import AutoEnableSettings


def _settings(**overrides) -> AutoEnableSettings:
    base = dict(
        enabled=True,
        dry_run=False,
        approval_required=False,
        max_rows_per_batch=200,
        max_rows_per_run=0,
        seconds_per_card_timeout=10,
        batch_timeout_buffer_seconds=300,
        working_statuses=("готов к работе", "активный вход", "активный выход"),
        auto_return_statuses=("не готов. плановый прозвон",),
        auto_return_target_status="Готов к работе",
        allowed_statuses_for_enable=(),
        deprecated_working_statuses_fallback=False,
        include_overdue=True,
        telegram_route_report="wallet_editor_auto_enable",
        telegram_route_alert="wallet_editor_auto_enable_alert",
    )
    base.update(overrides)
    return AutoEnableSettings(**base)


def _candidate(**overrides) -> CandidateRow:
    base = dict(
        card="4111",
        partner="Ostin",
        disable_at="01.06.2026 10:00:00",
        enable_status="К ВКЛЮЧЕНИЮ",
        vklyucheno="",
        source_row_index=0,
    )
    base.update(overrides)
    return CandidateRow(**base)


def _cfg() -> RunConfig:
    return RunConfig(login="u", password="p", auth_state_path="/tmp/auth.json")


@pytest.fixture
def page() -> MagicMock:
    return MagicMock()


def _hold_snapshot(*pairs: tuple[str, str]) -> HoldPairsSnapshot:
    normalized = frozenset(
        (normalize_hold_card(card), normalize_hold_partner(partner)) for card, partner in pairs
    )
    return HoldPairsSnapshot(normalized, True)


@pytest.fixture
def mock_wallet_editor_playwright(monkeypatch):
    page = MagicMock()
    context = MagicMock()
    browser = MagicMock()
    browser.new_context.return_value = context
    context.new_page.return_value = page

    playwright = MagicMock()
    playwright.chromium.launch.return_value = browser

    context_manager = MagicMock()
    context_manager.__enter__ = MagicMock(return_value=playwright)
    context_manager.__exit__ = MagicMock(return_value=False)
    monkeypatch.setattr("automation.engine.sync_playwright", lambda: context_manager)
    monkeypatch.setattr("automation.engine._ensure_logged_in", MagicMock())
    open_card = MagicMock()
    monkeypatch.setattr("automation.engine.open_card", open_card)
    monkeypatch.setattr("automation.engine.ensure_partner_removed", MagicMock(return_value="removed"))
    monkeypatch.setattr("automation.engine.ensure_partner_added", MagicMock(return_value="added partner"))
    monkeypatch.setattr("automation.engine.retry", lambda fn, *args, **kwargs: fn())
    monkeypatch.setattr(
        "automation.engine._apply_auto_no_partners_status_after_actions",
        MagicMock(return_value=False),
    )
    monkeypatch.setattr("automation.engine.save", MagicMock(return_value="saved"))
    monkeypatch.setattr(
        "automation.engine._verify_card_enable_after_save",
        MagicMock(return_value=None),
    )
    return open_card


# --- A) helper ---


def test_normalize_hold_card_and_partner_casefold():
    assert normalize_hold_card("4111") == "4111"
    assert normalize_hold_card("  Ostin  ") == "ostin"
    assert normalize_hold_partner("Ostin") == "ostin"
    assert normalize_hold_partner("  PARTNER  ") == "partner"


def test_is_card_partner_on_hold_true():
    pairs = {(normalize_hold_card("4111"), normalize_hold_partner("Ostin"))}
    assert is_card_partner_on_hold("4111", "Ostin", pairs) is True
    assert is_card_partner_on_hold("4111", "ostin", pairs) is True


def test_is_card_partner_on_hold_false():
    pairs = {(normalize_hold_card("4111"), normalize_hold_partner("Ostin"))}
    assert is_card_partner_on_hold("4111", "Other", pairs) is False
    assert is_card_partner_on_hold("9999", "Ostin", pairs) is False


def test_is_card_partner_on_hold_empty_pairs():
    assert is_card_partner_on_hold("4111", "Ostin", set()) is False
    assert is_card_partner_on_hold("4111", "Ostin", frozenset()) is False


def test_load_hold_pairs_from_dropbox_empty_when_no_path(monkeypatch):
    monkeypatch.setenv("DROPBOX_WALLET_EDITOR_PATH", "")
    snapshot = load_hold_pairs_from_dropbox()
    assert snapshot.available is False
    assert snapshot.pairs == frozenset()
    assert snapshot.error


def test_load_hold_pairs_from_dropbox_reads_hold_sheet(monkeypatch, tmp_path):
    import openpyxl

    xlsx = tmp_path / "registry.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "all_results"
    wb.create_sheet("runs")
    hold = wb.create_sheet("hold")
    hold.append(["card", "partner", "comment"])
    hold.append(["4111", "Ostin", "blocked"])
    wb.save(xlsx)

    monkeypatch.setenv("DROPBOX_WALLET_EDITOR_PATH", "/fake/registry.xlsx")

    def fake_download(_path, local):
        import shutil

        shutil.copy(xlsx, local)
        return "ok", "rev-1"

    monkeypatch.setattr(
        "integrations.wallet_editor_hold.download_file_with_rev",
        fake_download,
    )

    snapshot = load_hold_pairs_from_dropbox()
    assert snapshot.available is True
    assert is_card_partner_on_hold("4111", "Ostin", snapshot.pairs) is True


# --- B) manual engine ---


def test_group_actions_skips_fail_and_skip_rows():
    df = pd.DataFrame(
        {
            "card": ["c1", "c1", "c2", "c3"],
            "action": ["add_partner", "remove_partner", "add_partner", "set_status"],
            "value": ["Ostin", "Ostin", "Ostin", "Active"],
            "status": ["SKIP", "OK", "FAIL", ""],
        }
    )
    grouped = _group_actions(df)
    assert grouped == {"c1": [(1, "remove_partner", "Ostin")], "c3": [(3, "set_status", "Active")]}


def test_apply_add_partner_hold_precheck_marks_held_rows():
    df = pd.DataFrame(
        {
            "card": ["4111"],
            "action": ["add_partner"],
            "value": ["Ostin"],
            "status": [""],
            "comment": [""],
            OPERATION_DATE_COLUMN: [""],
            DISABLE_DATE_COLUMN: [""],
        }
    )
    stats = Stats()
    _apply_add_partner_hold_precheck(
        df,
        _hold_snapshot(("4111", "Ostin")),
        stats,
    )
    assert df.iloc[0]["status"] == "SKIP"
    assert df.iloc[0]["comment"] == HOLD_SKIP_COMMENT
    assert str(df.iloc[0]["Дата операции"]).strip()
    assert not str(df.iloc[0]["Дата отключения"]).strip()
    assert stats.skip == 1


def test_apply_add_partner_hold_precheck_fail_closed():
    df = pd.DataFrame(
        {
            "card": ["4111"],
            "action": ["add_partner"],
            "value": ["Ostin"],
            "status": [""],
            "comment": [""],
            OPERATION_DATE_COLUMN: [""],
            DISABLE_DATE_COLUMN: [""],
        }
    )
    stats = Stats()
    _apply_add_partner_hold_precheck(
        df,
        HoldPairsSnapshot.unavailable("download failed"),
        stats,
    )
    assert df.iloc[0]["status"] == "SKIP"
    assert df.iloc[0]["comment"] == HOLD_CHECK_FAILED_MANUAL_COMMENT
    assert stats.skip == 1


def test_add_partner_held_skip_no_open_card(tmp_path, monkeypatch, mock_wallet_editor_playwright):
    open_card = mock_wallet_editor_playwright
    monkeypatch.setattr(
        "automation.engine.load_hold_pairs_snapshot",
        lambda: _hold_snapshot(("4111111111111111", "Ostin")),
    )

    input_path = tmp_path / "input.xlsx"
    result_path = tmp_path / "result.xlsx"
    pd.DataFrame(
        {"card": ["4111111111111111"], "action": ["add_partner"], "value": ["Ostin"]}
    ).to_excel(input_path, index=False)

    _, stats = run(
        str(input_path),
        RunConfig(
            login="test-login",
            password="test-password",
            result_file_path=str(result_path),
        ),
    )

    open_card.assert_not_called()
    df = pd.read_excel(result_path)
    assert df.iloc[0]["status"] == "SKIP"
    assert df.iloc[0]["comment"] == HOLD_SKIP_COMMENT
    assert str(df.iloc[0]["Дата операции"]).strip()
    assert pd.isna(df.iloc[0]["Дата отключения"]) or not str(df.iloc[0]["Дата отключения"]).strip()
    assert stats.skip == 1


def test_add_partner_not_held_runs_normally(tmp_path, monkeypatch, mock_wallet_editor_playwright):
    open_card = mock_wallet_editor_playwright
    monkeypatch.setattr(
        "automation.engine.load_hold_pairs_snapshot",
        lambda: _hold_snapshot(("9999", "Other")),
    )

    input_path = tmp_path / "input.xlsx"
    result_path = tmp_path / "result.xlsx"
    pd.DataFrame(
        {"card": ["4111111111111111"], "action": ["add_partner"], "value": ["Ostin"]}
    ).to_excel(input_path, index=False)

    _, stats = run(
        str(input_path),
        RunConfig(
            login="test-login",
            password="test-password",
            result_file_path=str(result_path),
        ),
    )

    open_card.assert_called_once()
    df = pd.read_excel(result_path)
    assert df.iloc[0]["status"] == "OK"
    assert stats.ok >= 1


def test_held_add_partner_plus_remove_partner_same_card(
    tmp_path, monkeypatch, mock_wallet_editor_playwright
):
    open_card = mock_wallet_editor_playwright
    monkeypatch.setattr(
        "automation.engine.load_hold_pairs_snapshot",
        lambda: _hold_snapshot(("4111111111111111", "Ostin")),
    )

    input_path = tmp_path / "input.xlsx"
    result_path = tmp_path / "result.xlsx"
    pd.DataFrame(
        {
            "card": ["4111111111111111", "4111111111111111"],
            "action": ["add_partner", "remove_partner"],
            "value": ["Ostin", "Other"],
        }
    ).to_excel(input_path, index=False)

    run(
        str(input_path),
        RunConfig(
            login="test-login",
            password="test-password",
            result_file_path=str(result_path),
        ),
    )

    open_card.assert_called_once()
    df = pd.read_excel(result_path)
    add_row = df[df["action"] == "add_partner"].iloc[0]
    remove_row = df[df["action"] == "remove_partner"].iloc[0]
    assert add_row["status"] == "SKIP"
    assert add_row["comment"] == HOLD_SKIP_COMMENT
    assert remove_row["status"] == "OK"


def test_all_rows_held_skip_no_open_card(tmp_path, monkeypatch, mock_wallet_editor_playwright):
    open_card = mock_wallet_editor_playwright
    monkeypatch.setattr(
        "automation.engine.load_hold_pairs_snapshot",
        lambda: _hold_snapshot(("4111111111111111", "Ostin")),
    )

    input_path = tmp_path / "input.xlsx"
    result_path = tmp_path / "result.xlsx"
    pd.DataFrame(
        {"card": ["4111111111111111"], "action": ["add_partner"], "value": ["Ostin"]}
    ).to_excel(input_path, index=False)

    run(
        str(input_path),
        RunConfig(
            login="test-login",
            password="test-password",
            result_file_path=str(result_path),
        ),
    )

    open_card.assert_not_called()


def test_remove_partner_unaffected_by_hold(tmp_path, monkeypatch, mock_wallet_editor_playwright):
    open_card = mock_wallet_editor_playwright
    monkeypatch.setattr(
        "automation.engine.load_hold_pairs_snapshot",
        lambda: _hold_snapshot(("4111111111111111", "Ostin")),
    )

    input_path = tmp_path / "input.xlsx"
    result_path = tmp_path / "result.xlsx"
    pd.DataFrame(
        {"card": ["4111111111111111"], "action": ["remove_partner"], "value": ["Ostin"]}
    ).to_excel(input_path, index=False)

    run(
        str(input_path),
        RunConfig(
            login="test-login",
            password="test-password",
            result_file_path=str(result_path),
        ),
    )

    open_card.assert_called_once()
    df = pd.read_excel(result_path)
    assert df.iloc[0]["status"] == "OK"


def test_set_status_unaffected_by_hold(tmp_path, monkeypatch, mock_wallet_editor_playwright):
    open_card = mock_wallet_editor_playwright
    ensure_status = MagicMock(return_value="set status ok")
    monkeypatch.setattr("automation.engine.ensure_status_set", ensure_status)
    monkeypatch.setattr(
        "automation.engine.load_hold_pairs_snapshot",
        lambda: _hold_snapshot(("4111111111111111", "Ostin")),
    )

    input_path = tmp_path / "input.xlsx"
    result_path = tmp_path / "result.xlsx"
    pd.DataFrame(
        {"card": ["4111111111111111"], "action": ["set_status"], "value": ["Active"]}
    ).to_excel(input_path, index=False)

    run(
        str(input_path),
        RunConfig(
            login="test-login",
            password="test-password",
            result_file_path=str(result_path),
        ),
    )

    open_card.assert_called_once()
    ensure_status.assert_called_once()


# --- C) auto-enable ---


def test_auto_enable_candidate_held_skip_before_open_card(page):
    open_calls: list[str] = []

    outcome = process_enable_candidate(
        page,
        _candidate(),
        settings=_settings(),
        cfg=_cfg(),
        hold_snapshot=_hold_snapshot(("4111", "Ostin")),
        open_card_fn=lambda _p, card: open_calls.append(card),
        get_status_fn=lambda _p: pytest.fail("status read should not run"),
        get_chips_fn=lambda _p: pytest.fail("chips read should not run"),
        add_partner_fn=lambda *_a, **_k: pytest.fail("add_partner should not run"),
        save_fn=lambda *_a, **_k: pytest.fail("save should not run"),
    )

    assert outcome.registry_value == REGISTRY_SKIP
    assert outcome.registry_comment == HOLD_SKIP_COMMENT
    assert outcome.error_code == ERROR_HOLD
    assert outcome.mutated is False
    assert outcome.saved is False
    assert open_calls == []


def test_auto_enable_hold_check_failed_fail(page):
    outcome = process_enable_candidate(
        page,
        _candidate(),
        settings=_settings(),
        cfg=_cfg(),
        hold_snapshot=HoldPairsSnapshot.unavailable("registry down"),
        open_card_fn=lambda *_a, **_k: pytest.fail("open_card should not run"),
    )

    assert outcome.registry_value == REGISTRY_FAIL
    assert outcome.registry_comment == HOLD_CHECK_FAILED_AUTO_COMMENT
    assert outcome.error_code == ERROR_HOLD_CHECK_FAILED
    assert outcome.mutated is False
    assert outcome.saved is False


def test_auto_enable_non_held_candidate_unchanged(page):
    open_calls: list[str] = []
    chips: list[str] = []

    def add_partner(_page, partner, _cfg):
        chips.append(partner)
        return f"added {partner}"

    outcome = process_enable_candidate(
        page,
        _candidate(),
        settings=_settings(),
        cfg=_cfg(),
        hold_snapshot=_hold_snapshot(("9999", "Other")),
        open_card_fn=lambda _p, card: open_calls.append(card),
        get_status_fn=lambda _p: "Готов к работе",
        get_chips_fn=lambda _p: list(chips),
        add_partner_fn=add_partner,
        save_fn=lambda _p, _cfg: "saved",
    )

    assert outcome.registry_value == REGISTRY_OK
    assert open_calls == ["4111", "4111"]
    assert outcome.mutated is True
    assert outcome.saved is True


def test_batch_report_includes_hold_blocked(page):
    held = process_enable_candidate(
        page,
        _candidate(),
        settings=_settings(),
        cfg=_cfg(),
        hold_snapshot=_hold_snapshot(("4111", "Ostin")),
        open_card_fn=lambda *_a, **_k: None,
    )
    report = build_batch_execution_report([held], batch_index=1, batch_total=1, settings=_settings())
    assert "hold_blocked: 1" in report
    assert "SKIP: 1" in report


def test_registry_patch_receives_skip_for_held_outcome():
    row = {
        "Дата отключения": "01.06.2026 10:00:00",
        "Дата включения": "06.06.2026",
        "Статус включения": "К ВКЛЮЧЕНИЮ",
        "Включено": "",
        "Комментарий включения": "",
        "card": "4111",
        "partner": "Ostin",
        "action": ACTION_REMOVE_PARTNER,
        "status": "OK",
        "comment": "",
        "hold": "",
    }
    all_df = pd.DataFrame([row], columns=ALL_RESULTS_COLUMNS)
    otlezka = pd.DataFrame([{"partner": "Ostin", "Полные дни": 5, "comment": ""}])
    recalculated, _ = recalculate_all_results(all_df, pd.DataFrame(), otlezka, today=__import__("datetime").date(2026, 6, 6))

    updated, _ = apply_enable_updates_to_all_results(
        recalculated,
        [
            EnableRegistryUpdate(
                card="4111",
                partner="Ostin",
                disable_date="01.06.2026 10:00:00",
                vklyucheno="SKIP",
                comment=HOLD_SKIP_COMMENT,
            )
        ],
    )

    assert updated.at[0, "Включено"] == "SKIP"
    assert updated.at[0, "Комментарий включения"] == HOLD_SKIP_COMMENT
