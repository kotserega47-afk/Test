"""Unit tests for manual workbook snapshot parsing and hashing."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from integrations.wallet_editor_registry_db.manual_snapshot import (
    ManualSyncValidationError,
    compute_snapshot_hash,
    parse_manual_workbook,
)
from integrations.wallet_editor_registry_lifecycle import (
    ALL_RESULTS_COLUMNS,
    HOLD_COLUMNS,
    OTLEZKA_COLUMNS,
    RUNS_COLUMNS,
    SHEET_ALL_RESULTS,
    SHEET_HOLD,
    SHEET_OTLEZKA,
    SHEET_RUNS,
)


def _write_manual_workbook(
    path: Path,
    *,
    hold_rows: list[dict] | None = None,
    otlezka_rows: list[dict] | None = None,
    include_legacy: bool = False,
    include_otlezka: bool = True,
) -> None:
    hold_rows = hold_rows or [
        {
            "Дата добавления": "01.01.2026",
            "card": "4111111111111111",
            "partner": "Ostin",
            "comment": "note",
        }
    ]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        if include_legacy:
            pd.DataFrame(columns=ALL_RESULTS_COLUMNS).to_excel(
                writer, sheet_name=SHEET_ALL_RESULTS, index=False
            )
            pd.DataFrame(columns=RUNS_COLUMNS).to_excel(
                writer, sheet_name=SHEET_RUNS, index=False
            )
        pd.DataFrame(hold_rows, columns=HOLD_COLUMNS).to_excel(
            writer, sheet_name=SHEET_HOLD, index=False
        )
        if include_otlezka:
            otlezka_rows = otlezka_rows or [
                {"partner": "Ostin", "Полные дни": 30, "comment": "cooldown"}
            ]
            pd.DataFrame(otlezka_rows, columns=OTLEZKA_COLUMNS).to_excel(
                writer, sheet_name=SHEET_OTLEZKA, index=False
            )


def _snapshot_from_rows(
    hold_rows: list[dict],
    otlezka_rows: list[dict] | None = None,
    *,
    include_legacy: bool = False,
) -> tuple[object, str]:
    path = Path.cwd() / ".pytest_manual_wb.xlsx"
    try:
        _write_manual_workbook(
            path,
            hold_rows=hold_rows,
            otlezka_rows=otlezka_rows,
            include_legacy=include_legacy,
        )
        snapshot = parse_manual_workbook(path, "ok")
        return snapshot, compute_snapshot_hash(snapshot)
    finally:
        path.unlink(missing_ok=True)


class TestManualSnapshotHash:
    def test_same_normalized_data_same_hash(self):
        rows = [
            {
                "Дата добавления": "01.01.2026",
                "card": "4111111111111111",
                "partner": "Ostin",
                "comment": "a",
            }
        ]
        _, h1 = _snapshot_from_rows(rows)
        _, h2 = _snapshot_from_rows(rows)
        assert h1 == h2

    def test_row_reorder_same_hash(self, tmp_path):
        rows_a = [
            {
                "Дата добавления": "01.01.2026",
                "card": "4111",
                "partner": "A",
                "comment": "",
            },
            {
                "Дата добавления": "02.01.2026",
                "card": "4222",
                "partner": "B",
                "comment": "",
            },
        ]
        rows_b = list(reversed(rows_a))
        path_a = tmp_path / "a.xlsx"
        path_b = tmp_path / "b.xlsx"
        _write_manual_workbook(path_a, hold_rows=rows_a, otlezka_rows=[])
        _write_manual_workbook(path_b, hold_rows=rows_b, otlezka_rows=[])
        snap_a = parse_manual_workbook(path_a, "ok")
        snap_b = parse_manual_workbook(path_b, "ok")
        assert compute_snapshot_hash(snap_a) == compute_snapshot_hash(snap_b)

    def test_trim_and_casefold_same_hash(self):
        rows_a = [
            {
                "Дата добавления": "",
                "card": " 4111 ",
                "partner": "Ostin",
                "comment": "",
            }
        ]
        rows_b = [
            {
                "Дата добавления": "",
                "card": "4111",
                "partner": "OSTIN",
                "comment": "",
            }
        ]
        _, h1 = _snapshot_from_rows(rows_a, otlezka_rows=[])
        _, h2 = _snapshot_from_rows(rows_b, otlezka_rows=[])
        assert h1 == h2

    def test_comment_change_changes_hash(self):
        base = [
            {
                "Дата добавления": "",
                "card": "4111",
                "partner": "Ostin",
                "comment": "one",
            }
        ]
        changed = [
            {
                "Дата добавления": "",
                "card": "4111",
                "partner": "Ostin",
                "comment": "two",
            }
        ]
        _, h1 = _snapshot_from_rows(base, otlezka_rows=[])
        _, h2 = _snapshot_from_rows(changed, otlezka_rows=[])
        assert h1 != h2

    def test_added_at_change_changes_hash(self):
        base = [
            {
                "Дата добавления": "01.01.2026",
                "card": "4111",
                "partner": "Ostin",
                "comment": "",
            }
        ]
        changed = [
            {
                "Дата добавления": "02.01.2026",
                "card": "4111",
                "partner": "Ostin",
                "comment": "",
            }
        ]
        _, h1 = _snapshot_from_rows(base, otlezka_rows=[])
        _, h2 = _snapshot_from_rows(changed, otlezka_rows=[])
        assert h1 != h2

    def test_full_days_change_changes_hash(self):
        hold = [
            {
                "Дата добавления": "",
                "card": "4111",
                "partner": "Ostin",
                "comment": "",
            }
        ]
        _, h1 = _snapshot_from_rows(
            hold,
            otlezka_rows=[{"partner": "Ostin", "Полные дни": 10, "comment": ""}],
        )
        _, h2 = _snapshot_from_rows(
            hold,
            otlezka_rows=[{"partner": "Ostin", "Полные дни": 11, "comment": ""}],
        )
        assert h1 != h2

    def test_legacy_sheets_ignored_for_hash(self, tmp_path):
        hold = [
            {
                "Дата добавления": "",
                "card": "4111",
                "partner": "Ostin",
                "comment": "",
            }
        ]
        path_plain = tmp_path / "plain.xlsx"
        path_legacy = tmp_path / "legacy.xlsx"
        _write_manual_workbook(path_plain, hold_rows=hold, otlezka_rows=[], include_legacy=False)
        _write_manual_workbook(path_legacy, hold_rows=hold, otlezka_rows=[], include_legacy=True)
        h_plain = compute_snapshot_hash(parse_manual_workbook(path_plain, "ok"))
        h_legacy = compute_snapshot_hash(parse_manual_workbook(path_legacy, "ok"))
        assert h_plain == h_legacy

    def test_legacy_sheet_edits_do_not_change_hash(self, tmp_path):
        hold = [
            {
                "Дата добавления": "",
                "card": "4111",
                "partner": "Ostin",
                "comment": "",
            }
        ]
        path_v1 = tmp_path / "v1.xlsx"
        path_v2 = tmp_path / "v2.xlsx"
        _write_manual_workbook(path_v1, hold_rows=hold, otlezka_rows=[], include_legacy=True)
        _write_manual_workbook(path_v2, hold_rows=hold, otlezka_rows=[], include_legacy=True)

        from openpyxl import load_workbook

        wb = load_workbook(path_v2)
        ws = wb[SHEET_ALL_RESULTS]
        ws.cell(row=2, column=1, value="edited")
        wb.save(path_v2)
        wb.close()

        h1 = compute_snapshot_hash(parse_manual_workbook(path_v1, "ok"))
        h2 = compute_snapshot_hash(parse_manual_workbook(path_v2, "ok"))
        assert h1 == h2


class TestManualSnapshotValidation:
    def test_duplicate_hold_key_raises(self):
        rows = [
            {
                "Дата добавления": "",
                "card": "4111",
                "partner": "Ostin",
                "comment": "",
            },
            {
                "Дата добавления": "",
                "card": "4111",
                "partner": "Ostin",
                "comment": "dup",
            },
        ]
        path = Path.cwd() / ".pytest_dup_hold.xlsx"
        try:
            _write_manual_workbook(path, hold_rows=rows, otlezka_rows=[])
            with pytest.raises(ManualSyncValidationError, match="duplicate hold"):
                parse_manual_workbook(path, "ok")
        finally:
            path.unlink(missing_ok=True)

    def test_duplicate_otlezka_partner_raises(self):
        hold = [
            {
                "Дата добавления": "",
                "card": "4111",
                "partner": "Ostin",
                "comment": "",
            }
        ]
        otlezka = [
            {"partner": "Ostin", "Полные дни": 10, "comment": ""},
            {"partner": "Ostin", "Полные дни": 20, "comment": ""},
        ]
        path = Path.cwd() / ".pytest_dup_ot.xlsx"
        try:
            _write_manual_workbook(path, hold_rows=hold, otlezka_rows=otlezka)
            with pytest.raises(ManualSyncValidationError, match="duplicate"):
                parse_manual_workbook(path, "ok")
        finally:
            path.unlink(missing_ok=True)

    def test_missing_hold_sheet_blocks(self, tmp_path):
        path = tmp_path / "no_hold.xlsx"
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            pd.DataFrame(
                [{"partner": "Ostin", "Полные дни": 1, "comment": ""}],
                columns=OTLEZKA_COLUMNS,
            ).to_excel(writer, sheet_name=SHEET_OTLEZKA, index=False)
        with pytest.raises(ManualSyncValidationError, match="hold"):
            parse_manual_workbook(path, "ok")

    def test_missing_otlezka_sheet_warns_only(self, tmp_path):
        path = tmp_path / "no_ot.xlsx"
        _write_manual_workbook(path, include_otlezka=False)
        snapshot = parse_manual_workbook(path, "ok")
        assert snapshot.otlezka_rows == ()
        assert any("Отлёжка" in w for w in snapshot.warnings)
