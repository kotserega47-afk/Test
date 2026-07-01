"""Unit tests for registry export workbook formatting."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from integrations.wallet_editor_registry_db.registry_export_format import (
    EXPORT_DATETIME_DISPLAY,
    EXPORT_SHEET_README,
    EXPORT_SHEET_RESULTS,
    EXPORT_SHEET_RUNS,
    EXPORT_SHEET_STATS,
    MAX_COLUMN_WIDTH,
    STATS_ADD_TITLE,
    STATS_REMOVE_TITLE,
    STATS_TOTAL_HEADER,
    STATS_TOTAL_ROW_LABEL,
    build_statistics_sheet_rows,
    format_export_datetime_value,
    prepare_export_frames,
    sort_all_results_for_export,
    sort_runs_for_export,
    write_registry_export_workbook,
)
from integrations.wallet_editor_registry_lifecycle import (
    ACTION_ADD_PARTNER,
    ACTION_REMOVE_PARTNER,
    ALL_RESULTS_COLUMNS,
    HOLD_COLUMNS,
    OPERATION_DATE_COLUMN,
    OTLEZKA_COLUMNS,
    RUNS_COLUMNS,
)


def _all_results_row(**overrides) -> dict:
    row = {col: "" for col in ALL_RESULTS_COLUMNS}
    row.update(overrides)
    return row


def _sample_frames():
    all_results = pd.DataFrame(
        [
            _all_results_row(
                **{
                    OPERATION_DATE_COLUMN: "02.06.2026 10:00:00",
                    "partner": "Beta",
                    "action": ACTION_REMOVE_PARTNER,
                }
            ),
            _all_results_row(
                **{
                    OPERATION_DATE_COLUMN: "01.06.2026 09:00:00",
                    "partner": "Alpha",
                    "action": ACTION_REMOVE_PARTNER,
                }
            ),
            _all_results_row(
                **{
                    OPERATION_DATE_COLUMN: "01.06.2026 11:00:00",
                    "partner": "Alpha",
                    "action": ACTION_ADD_PARTNER,
                }
            ),
            _all_results_row(
                **{
                    OPERATION_DATE_COLUMN: "03.06.2026 12:00:00",
                    "partner": "Gamma",
                    "action": ACTION_ADD_PARTNER,
                }
            ),
        ]
    )
    runs = pd.DataFrame(
        [
            {col: "" for col in RUNS_COLUMNS}
            | {
                "started_at": "02.06.2026 08:00:00",
                "finished_at": "02.06.2026 08:05:00",
            },
            {col: "" for col in RUNS_COLUMNS}
            | {
                "started_at": "01.06.2026 07:00:00",
                "finished_at": "01.06.2026 07:05:00",
            },
        ]
    )
    hold = pd.DataFrame(
        [
            {
                "Дата добавления": "02.06.2026",
                "card": "4111",
                "partner": "Beta",
                "comment": "",
            },
            {
                "Дата добавления": "01.06.2026",
                "card": "4222",
                "partner": "Alpha",
                "comment": "",
            },
        ]
    )
    otlezka = pd.DataFrame(columns=OTLEZKA_COLUMNS)
    return all_results, runs, hold, otlezka


class TestExportSortingAndFormatting:
    def test_sort_all_results_by_operation_date_ascending(self):
        all_results, _, _, _ = _sample_frames()
        sorted_df = sort_all_results_for_export(all_results)
        dates = sorted_df[OPERATION_DATE_COLUMN].tolist()
        assert dates[0].startswith("01.06.2026")
        assert dates[-1].startswith("03.06.2026")

    def test_sort_runs_by_started_at_ascending(self):
        _, runs, _, _ = _sample_frames()
        sorted_df = sort_runs_for_export(runs)
        assert sorted_df.iloc[0]["started_at"].startswith("01.06.2026")
        assert sorted_df.iloc[1]["started_at"].startswith("02.06.2026")

    def test_format_export_datetime_value(self):
        assert format_export_datetime_value("2026-06-03T14:30:45+03:00") == "03.06.2026 14:30:45"
        assert format_export_datetime_value("03.06.2026") == "03.06.2026 00:00:00"
        assert (
            format_export_datetime_value(datetime(2026, 6, 3, 14, 30, 45))
            == "03.06.2026 14:30:45"
        )

    def test_prepare_export_frames_formats_datetime_columns(self):
        all_results, runs, hold, _ = _sample_frames()
        prepared_all, prepared_runs, prepared_hold, _ = prepare_export_frames(
            all_results,
            runs,
            hold,
            pd.DataFrame(columns=OTLEZKA_COLUMNS),
        )
        assert prepared_all.iloc[0][OPERATION_DATE_COLUMN] == "01.06.2026 09:00:00"
        assert prepared_runs.iloc[0]["started_at"] == "01.06.2026 07:00:00"
        assert prepared_hold.iloc[0]["Дата добавления"] == "01.06.2026 00:00:00"


class TestStatisticsSheet:
    def test_build_remove_and_add_partner_blocks(self):
        all_results, _, _, _ = _sample_frames()
        rows = build_statistics_sheet_rows(all_results)
        text = "\n".join(" | ".join(str(cell) for cell in row) for row in rows if row)
        assert STATS_REMOVE_TITLE in text
        assert STATS_ADD_TITLE in text
        assert "Alpha" in text
        assert STATS_TOTAL_ROW_LABEL in text
        assert STATS_TOTAL_HEADER in text

    def test_statistics_totals_and_vsego_column(self):
        all_results, _, _, _ = _sample_frames()
        rows = build_statistics_sheet_rows(all_results)
        add_idx = next(i for i, row in enumerate(rows) if row and row[0] == STATS_ADD_TITLE)
        remove_header = next(row for row in rows if row and row[0] == "Партнёр")
        remove_total = next(
            row
            for row in rows[:add_idx]
            if row and row[0] == STATS_TOTAL_ROW_LABEL
        )
        assert remove_header[-1] == STATS_TOTAL_HEADER
        assert remove_total[-1] == 2


class TestExportWorkbookPresentation:
    def test_workbook_has_russian_sheets_and_formatting(self, tmp_path):
        all_results, runs, hold, otlezka = _sample_frames()
        output = tmp_path / "export.xlsx"
        write_registry_export_workbook(
            output,
            all_results=all_results,
            runs=runs,
            hold=hold,
            otlezka=otlezka,
            readme_lines=["Реестр WalletEditor", "PostgreSQL", "/registry_export"],
        )

        wb = load_workbook(output)
        assert wb.sheetnames == [
            EXPORT_SHEET_RESULTS,
            EXPORT_SHEET_RUNS,
            "Hold",
            "Отлёжка",
            EXPORT_SHEET_STATS,
            EXPORT_SHEET_README,
        ]

        results_ws = wb[EXPORT_SHEET_RESULTS]
        assert results_ws.freeze_panes == "A2"
        assert results_ws.auto_filter.ref is not None
        assert results_ws.cell(row=1, column=1).font.bold is True
        assert results_ws.cell(row=2, column=1).value.startswith("01.06.2026")
        assert results_ws.cell(row=3, column=1).value.startswith("01.06.2026")
        assert results_ws.cell(row=5, column=1).value.startswith("03.06.2026")
        assert results_ws.cell(row=2, column=1).border.left.style == "thin"

        runs_ws = wb[EXPORT_SHEET_RUNS]
        assert runs_ws.cell(row=2, column=1).value == "01.06.2026 07:00:00"
        assert runs_ws.auto_filter.ref is not None

        stats_ws = wb[EXPORT_SHEET_STATS]
        assert stats_ws.cell(row=1, column=1).value == STATS_REMOVE_TITLE
        assert EXPORT_SHEET_STATS in wb.sheetnames

        readme_ws = wb[EXPORT_SHEET_README]
        readme_text = readme_ws.cell(row=1, column=1).value
        assert "Реестр WalletEditor" in readme_text
        assert readme_ws.auto_filter.ref is None

        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    assert cell.data_type != "f"
            for col_idx in range(1, ws.max_column + 1):
                width = ws.column_dimensions[get_column_letter(col_idx)].width
                if width is not None:
                    assert width <= MAX_COLUMN_WIDTH

        wb.close()

    def test_column_width_capped_at_50(self, tmp_path):
        all_results, runs, hold, otlezka = _sample_frames()
        all_results.at[0, "comment"] = "x" * 100
        write_registry_export_workbook(
            tmp_path / "wide.xlsx",
            all_results=all_results,
            runs=runs,
            hold=hold,
            otlezka=otlezka,
            readme_lines=["README"],
        )
        wb = load_workbook(tmp_path / "wide.xlsx")
        comment_col = ALL_RESULTS_COLUMNS.index("comment") + 1
        width = wb[EXPORT_SHEET_RESULTS].column_dimensions[
            get_column_letter(comment_col)
        ].width
        assert width <= MAX_COLUMN_WIDTH
        wb.close()
