"""Stage C — shadow-write hooks after Excel SUCCESS (I-REG-08)."""

from __future__ import annotations

import shutil
import threading
import time
from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock, patch
from zoneinfo import ZoneInfo

import pandas as pd
import pytest

from automation.audit import Stats
from automation.runtime import WalletEditorTask
from integrations.wallet_editor_registry import (
    EnableRegistryUpdate,
    _AppendOutcome,
    append_run_to_dropbox_registry,
    build_registry_health_report,
    format_registry_health_report,
    patch_enable_results_in_dropbox_registry,
)
from integrations.wallet_editor_registry_db.config import ENV_MIRROR_ENABLED
from integrations.wallet_editor_registry_db.mirror import (
    MirrorResultBatch,
    schedule_mirror_batch,
    write_mirror_batch,
)
from integrations.wallet_editor_registry_db.mirror_state import reset_mirror_health_for_tests
from integrations.wallet_editor_registry_db.models import RegistryResultRow, RegistryRunRow
from integrations.wallet_editor_registry_lifecycle import (
    ALL_RESULTS_COLUMNS,
    DISABLE_DATE_COLUMN,
    OPERATION_DATE_COLUMN,
    result_row_dates,
)
from integrations.wallet_editor_registry_refresh import refresh_wallet_editor_registry_lifecycle
from integrations.wallet_editor_registry_xlsx import save_registry_workbook

MSK = ZoneInfo("Europe/Moscow")
DROPBOX_PATH = "/Ostin/platform/Tests/wallet_editor.xlsx"
RUN_STARTED = datetime(2026, 6, 22, 9, 0, 0, tzinfo=MSK)
RUN_FINISHED = datetime(2026, 6, 22, 9, 5, 0, tzinfo=MSK)


def _make_task(*, run_id: str = "mirror-run-1") -> WalletEditorTask:
    return WalletEditorTask(
        file_path="/tmp/wallet_editor/in.xlsx",
        chat_id=-1001,
        telegram_user_id=111,
        operator_profile="DENIS",
        source_file_name="batch.xlsx",
        login="login",
        password="pass",
        auth_state_path="/tmp/auth.json",
        run_id=run_id,
    )


def _write_result(path: Path) -> None:
    processed = datetime(2026, 6, 22, 9, 0, 0, tzinfo=MSK)
    operation_date, disable_date = result_row_dates("remove_partner", "OK", processed)
    pd.DataFrame(
        {
            OPERATION_DATE_COLUMN: [operation_date],
            DISABLE_DATE_COLUMN: [disable_date],
            "card": ["4111111111111111"],
            "action": ["remove_partner"],
            "value": ["Ostin"],
            "status": ["OK"],
            "comment": [""],
        }
    ).to_excel(path, index=False)


@pytest.fixture(autouse=True)
def _reset_mirror_health():
    reset_mirror_health_for_tests()
    yield
    reset_mirror_health_for_tests()


@pytest.fixture
def fast_registry_settings():
    settings = MagicMock(
        registry_timeout_seconds=1,
        registry_warning_seconds=60,
        registry_retry_interval_seconds=0,
    )
    with patch(
        "integrations.wallet_editor_registry.load_registry_settings",
        return_value=settings,
    ):
        yield settings


@pytest.fixture
def registry_env(monkeypatch, tmp_path):
    monkeypatch.setenv("DROPBOX_WALLET_EDITOR_PATH", DROPBOX_PATH)
    monkeypatch.setenv("STATE_DIR", str(tmp_path / "state"))
    result_path = tmp_path / "result.xlsx"
    _write_result(result_path)
    return result_path


class TestScheduleMirrorBatch:
    def test_flag_off_does_not_write(self, monkeypatch):
        monkeypatch.setenv(ENV_MIRROR_ENABLED, "0")
        batch = MirrorResultBatch(
            runs=(
                RegistryRunRow(
                    run_id="r1",
                    started_at="22.06.2026 09:00:00",
                    finished_at="22.06.2026 09:05:00",
                    input_rows=1,
                    success_rows=1,
                    failed_rows=0,
                    skipped_rows=0,
                ),
            ),
        )
        with patch(
            "integrations.wallet_editor_registry_db.mirror.write_mirror_batch"
        ) as write_mock:
            schedule_mirror_batch(batch, operation="append")
            time.sleep(0.05)
        write_mock.assert_not_called()

    def test_flag_on_writes_async(self, monkeypatch):
        monkeypatch.setenv(ENV_MIRROR_ENABLED, "1")
        called = threading.Event()
        batch = MirrorResultBatch(
            runs=(
                RegistryRunRow(
                    run_id="r1",
                    started_at="22.06.2026 09:00:00",
                    finished_at="22.06.2026 09:05:00",
                    input_rows=1,
                    success_rows=1,
                    failed_rows=0,
                    skipped_rows=0,
                ),
            ),
        )

        def _fake_write(_batch: MirrorResultBatch) -> None:
            called.set()

        with patch(
            "integrations.wallet_editor_registry_db.mirror.write_mirror_batch",
            side_effect=_fake_write,
        ):
            schedule_mirror_batch(batch, operation="append")
            assert called.wait(timeout=2.0)


class TestAppendMirrorHook:
    def test_mirror_scheduled_after_excel_success(
        self, registry_env, monkeypatch, fast_registry_settings
    ):
        monkeypatch.setenv(ENV_MIRROR_ENABLED, "1")
        task = _make_task()
        stats = Stats()
        stats.ok = 1

        with (
            patch(
                "integrations.wallet_editor_registry.download_file_with_rev",
                return_value=("not_found", None),
            ),
            patch(
                "integrations.wallet_editor_registry.upload_file_if_rev",
                return_value="uploaded",
            ),
            patch(
                "integrations.wallet_editor_registry.schedule_mirror_batch"
            ) as mirror_mock,
        ):
            append_run_to_dropbox_registry(
                task,
                str(registry_env),
                stats,
                run_started_at=RUN_STARTED,
                run_finished_at=RUN_FINISHED,
            )

        mirror_mock.assert_called_once()
        batch = mirror_mock.call_args.args[0]
        assert batch is not None
        assert len(batch.runs) == 1
        assert len(batch.results) == 1
        assert mirror_mock.call_args.kwargs["operation"] == "append"

    def test_mirror_not_called_on_excel_upload_failure(
        self, registry_env, monkeypatch, fast_registry_settings
    ):
        monkeypatch.setenv(ENV_MIRROR_ENABLED, "1")
        task = _make_task()
        stats = Stats()
        stats.ok = 1

        with (
            patch(
                "integrations.wallet_editor_registry.download_file_with_rev",
                return_value=("ok", "rev1"),
            ),
            patch(
                "integrations.wallet_editor_registry.upload_file_if_rev",
                return_value="rev_conflict",
            ),
            patch("integrations.wallet_editor_registry.schedule_mirror_batch") as mirror_mock,
        ):
            append_run_to_dropbox_registry(
                task,
                str(registry_env),
                stats,
                run_started_at=RUN_STARTED,
                run_finished_at=RUN_FINISHED,
            )

        mirror_mock.assert_not_called()

    def test_db_failure_does_not_fail_excel_append(
        self, registry_env, monkeypatch, fast_registry_settings
    ):
        monkeypatch.setenv(ENV_MIRROR_ENABLED, "1")
        task = _make_task(run_id="mirror-db-fail")
        stats = Stats()
        stats.ok = 1

        from integrations.wallet_editor_registry_async import (
            OUTBOX_STATUS_SYNCED,
            create_outbox_record,
            get_outbox_record,
            persist_durable_result_copy,
        )

        durable = persist_durable_result_copy(task.run_id, str(registry_env))
        create_outbox_record(
            task,
            durable_path=durable,
            stats=stats,
            run_started_at=RUN_STARTED,
            run_finished_at=RUN_FINISHED,
            output_file="result.xlsx",
        )

        with (
            patch(
                "integrations.wallet_editor_registry.download_file_with_rev",
                return_value=("not_found", None),
            ),
            patch(
                "integrations.wallet_editor_registry.upload_file_if_rev",
                return_value="uploaded",
            ),
            patch(
                "integrations.wallet_editor_registry_db.mirror.write_mirror_batch",
                side_effect=RuntimeError("db down"),
            ),
        ):
            append_run_to_dropbox_registry(
                task,
                str(registry_env),
                stats,
                run_started_at=RUN_STARTED,
                run_finished_at=RUN_FINISHED,
            )
            time.sleep(0.15)

        record = get_outbox_record(task.run_id)
        assert record is not None
        assert record.status == OUTBOX_STATUS_SYNCED


class TestPatchMirrorHook:
    def _registry_file(self, tmp_path: Path) -> Path:
        processed = datetime(2026, 6, 22, 9, 0, 0, tzinfo=MSK)
        operation_date, disable_date = result_row_dates("remove_partner", "OK", processed)
        all_results = pd.DataFrame(
            [
                {
                    OPERATION_DATE_COLUMN: operation_date,
                    DISABLE_DATE_COLUMN: disable_date,
                    "Дата включения": "25.06.2026",
                    "Статус включения": "К ВКЛЮЧЕНИЮ",
                    "Включено": "",
                    "Комментарий включения": "",
                    "card": "4111111111111111",
                    "partner": "Ostin",
                    "action": "remove_partner",
                    "status": "OK",
                    "comment": "",
                    "hold": "",
                }
            ],
            columns=ALL_RESULTS_COLUMNS,
        )
        runs = pd.DataFrame(columns=["started_at", "finished_at", "input_rows", "success_rows", "failed_rows", "skipped_rows", "output_file"])
        path = tmp_path / "wallet_editor.xlsx"
        save_registry_workbook(
            path,
            all_results=all_results,
            runs=runs,
            hold_exists=False,
            otlezka_exists=False,
            is_new_file=True,
        )
        return path

    def test_patch_mirror_after_success(self, tmp_path, monkeypatch, fast_registry_settings):
        monkeypatch.setenv("DROPBOX_WALLET_EDITOR_PATH", DROPBOX_PATH)
        monkeypatch.setenv(ENV_MIRROR_ENABLED, "1")
        local = self._registry_file(tmp_path)

        updates = [
            EnableRegistryUpdate(
                card="4111111111111111",
                partner="Ostin",
                disable_date="22.06.2026 09:00:00",
                vklyucheno="OK",
                comment="enabled",
                source_row_index=0,
            )
        ]

        with (
            patch(
                "integrations.wallet_editor_registry.download_file_with_rev",
                side_effect=lambda _path, dest: (shutil.copy(local, dest), ("ok", "rev1"))[1],
            ) as download_mock,
            patch(
                "integrations.wallet_editor_registry.upload_file_if_rev",
                return_value="uploaded",
            ),
            patch(
                "integrations.wallet_editor_registry.schedule_mirror_batch"
            ) as mirror_mock,
        ):
            result = patch_enable_results_in_dropbox_registry(updates)

        assert result.success is True
        mirror_mock.assert_called_once()
        batch = mirror_mock.call_args.args[0]
        assert len(batch.results) == 1
        assert batch.results[0].vklyucheno == "OK"
        download_mock.assert_called()


class TestRefreshMirrorHook:
    def test_refresh_mirror_only_after_upload_success(
        self, tmp_path, monkeypatch, fast_registry_settings
    ):
        monkeypatch.setenv("DROPBOX_WALLET_EDITOR_PATH", DROPBOX_PATH)
        monkeypatch.setenv(ENV_MIRROR_ENABLED, "1")

        from integrations.wallet_editor_registry_lifecycle import SHEET_OTLEZKA

        processed = datetime(2026, 6, 20, 9, 0, 0, tzinfo=MSK)
        operation_date, disable_date = result_row_dates("remove_partner", "OK", processed)
        all_results = pd.DataFrame(
            [
                {
                    OPERATION_DATE_COLUMN: operation_date,
                    DISABLE_DATE_COLUMN: disable_date,
                    "Дата включения": "23.06.2026",
                    "Статус включения": "ОЖИДАЕТ",
                    "Включено": "",
                    "Комментарий включения": "",
                    "card": "4111111111111111",
                    "partner": "Ostin",
                    "action": "remove_partner",
                    "status": "OK",
                    "comment": "",
                    "hold": "",
                }
            ],
            columns=ALL_RESULTS_COLUMNS,
        )
        runs = pd.DataFrame(
            [
                {
                    "started_at": "20.06.2026 09:00:00",
                    "finished_at": "20.06.2026 09:05:00",
                    "input_rows": 1,
                    "success_rows": 1,
                    "failed_rows": 0,
                    "skipped_rows": 0,
                    "output_file": "r.xlsx",
                }
            ]
        )
        local = tmp_path / "wallet_editor.xlsx"
        save_registry_workbook(
            local,
            all_results=all_results,
            runs=runs,
            hold_exists=False,
            otlezka_exists=False,
            is_new_file=True,
        )
        from openpyxl import load_workbook

        from integrations.wallet_editor_registry_lifecycle import SHEET_OTLEZKA as OTLEZKA_SHEET

        wb = load_workbook(local)
        ws = wb[OTLEZKA_SHEET]
        ws.cell(row=2, column=1, value="Ostin")
        ws.cell(row=2, column=2, value=3)
        wb.save(local)
        wb.close()

        with (
            patch(
                "integrations.wallet_editor_registry_refresh.load_registry_settings",
                return_value=fast_registry_settings,
            ),
            patch(
                "integrations.wallet_editor_registry_refresh.download_file_with_rev",
                side_effect=lambda _path, dest: (shutil.copy(local, dest), ("ok", "rev1"))[1],
            ),
            patch(
                "integrations.wallet_editor_registry_refresh.upload_file_if_rev",
                return_value="uploaded",
            ),
            patch(
                "integrations.wallet_editor_registry_refresh.schedule_mirror_batch"
            ) as mirror_mock,
        ):
            result = refresh_wallet_editor_registry_lifecycle(today=datetime(2026, 6, 23).date())

        assert result.success is True
        assert result.uploaded is True
        mirror_mock.assert_called_once()


class TestHealthMirrorFields:
    def test_health_includes_mirror_section(self, monkeypatch):
        monkeypatch.setenv(ENV_MIRROR_ENABLED, "0")
        report = build_registry_health_report()
        text = format_registry_health_report(report)
        assert "mirror enabled: False" in text
        assert "mirror recent failures:" in text


class TestNoDbReadsInRuntime:
    def test_registry_module_has_no_db_connect(self):
        source = Path("integrations/wallet_editor_registry.py").read_text(encoding="utf-8")
        assert "connect(for_mirror=False)" not in source
        assert "connect(for_mirror=True)" not in source
        assert "PostgresRegistryStore" not in source

    def test_refresh_module_has_no_db_connect(self):
        source = Path("integrations/wallet_editor_registry_refresh.py").read_text(
            encoding="utf-8"
        )
        assert "connect(" not in source

    def test_auto_enable_reads_postgres_only_when_source_postgres(self):
        source = Path("integrations/wallet_editor_auto_enable.py").read_text(encoding="utf-8")
        assert "registry_source_is_postgres" in source
        assert "load_registry_frames_from_postgres" in source
